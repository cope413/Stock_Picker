"""Offline tests for landry.ledger -- the guarded, atomic way entries get
written to a generated tab. Synthetic workbooks (the tab skeletons from
test_landry_generate) except the last test, which writes to a scratch copy of
the repo's own tracked workbook and skips if it isn't there."""

import datetime
import io
import os
import shutil
import sqlite3
import sys
import types

import openpyxl
import pytest
from openpyxl.styles import Alignment
from openpyxl.worksheet.formula import ArrayFormula

from landry import cli, generate, ledger, models, xlsx_io
from landry.ledger import DriftError, ExcelOpenError, LedgerError
from test_landry_generate import DD_ENTRIES, ENTRIES, REPO_ROOT, _dd_skeleton, _skeleton

_REAL_EXCEL_CHECK = ledger._excel_has_open


@pytest.fixture(autouse=True)
def excel_is_closed(monkeypatch):
    """Don't let the tests depend on whether Excel is running on this machine."""
    monkeypatch.setattr(ledger, "_excel_has_open", lambda path: False)


def _journal_workbook(tmp_path, entries=ENTRIES, capacity=5):
    """A workbook whose Journal already holds ``entries``, plus the database that agrees with it."""
    wb, _ = _skeleton(capacity=capacity)
    xlsx, db = str(tmp_path / "wb.xlsx"), str(tmp_path / "landry.db")
    wb.save(xlsx)
    conn = models.init_db(db)
    for d, label, notes in entries:
        models.journal_add(conn, d, label, notes)
    conn.commit()
    conn.close()
    generate.generate_workbook(xlsx, db, recalc=False)
    return xlsx, db


def _drawdown_workbook(tmp_path, entries=DD_ENTRIES, capacity=6):
    wb, _ = _dd_skeleton(capacity=capacity)
    xlsx, db = str(tmp_path / "dd.xlsx"), str(tmp_path / "landry.db")
    wb.save(xlsx)
    conn = models.init_db(db)
    for d, v, n in entries:
        models.drawdown_add(conn, d, v, n)
    conn.commit()
    conn.close()
    generate.generate_workbook(xlsx, db, recalc=False)
    return xlsx, db


def _count(db, table):
    conn = sqlite3.connect(db)
    try:
        return conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
    finally:
        conn.close()


def _plain(value):
    """ArrayFormula (the STOCKHISTORY cells) compares by identity; unwrap it."""
    return (value.ref, value.text) if isinstance(value, ArrayFormula) else value


def _bytes(path):
    with open(path, "rb") as f:
        return f.read()


# ---------------------------------------------------------------- Journal --

def test_journal_add_appends_the_row_and_regenerates_the_tab(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    result = ledger.journal_add(xlsx, db, "a new decision", date="2026-10-03",
                                label="DCA-CATCHUP-2", recalc=False)
    assert (result["row"], result["id"]) == (6, 4)
    ws = openpyxl.load_workbook(xlsx)["Journal"]
    assert ws["A6"].value == datetime.datetime(2026, 10, 3)
    assert (ws["B6"].value, ws["C6"].value) == ("DCA-CATCHUP-2", "a new decision")
    assert ws["C6"].alignment.wrap_text is True                 # canonical formatting, not copied from a neighbour
    assert generate.verify_journal(xlsx, db) == []              # tab and database agree afterwards
    assert _count(db, "journal") == 4


def test_a_dateless_journal_add_is_dated_today(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    ledger.journal_add(xlsx, db, "no date given", recalc=False)
    assert openpyxl.load_workbook(xlsx)["Journal"]["A6"].value.date() == datetime.date.today()


def test_an_empty_note_is_refused(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    with pytest.raises(LedgerError, match="needs notes"):
        ledger.journal_add(xlsx, db, "   ", recalc=False)


def test_a_write_with_no_database_builds_it_from_the_workbook(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    os.remove(db)
    result = ledger.journal_add(xlsx, db, "first write on a fresh checkout", recalc=False)
    assert result["built_db_from_workbook"] == {"journal": 3}
    assert _count(db, "journal") == 4
    assert openpyxl.load_workbook(xlsx)["Journal"]["C6"].value == "first write on a fresh checkout"


def test_drift_blocks_the_write_and_changes_nothing(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    wb["Journal"]["C4"] = "hand-edited in Excel"
    wb.save(xlsx)
    before = _bytes(xlsx)
    with pytest.raises(DriftError, match="landry db pull") as err:
        ledger.journal_add(xlsx, db, "should not be written", recalc=False)
    assert "C4" in str(err.value)
    assert _bytes(xlsx) == before and _count(db, "journal") == 3


def test_pull_accepts_the_workbook_and_unblocks_the_write(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    wb["Journal"]["C4"] = "hand-edited in Excel"
    wb.save(xlsx)
    assert ledger.pull(xlsx, db) == {"journal": 3}
    ledger.journal_add(xlsx, db, "now it goes through", recalc=False)
    ws = openpyxl.load_workbook(xlsx)["Journal"]
    assert ws["C4"].value == "hand-edited in Excel" and ws["C6"].value == "now it goes through"


def test_format_differences_are_not_drift(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    wb["Journal"]["C3"].alignment = Alignment(wrap_text=False)   # "turned wrapping off to look"
    wb.save(xlsx)
    report = ledger.status(xlsx, db)["journal"]
    assert report["value_diffs"] == [] and report["format_diffs"] and report["db_rows"] == 3
    ledger.journal_add(xlsx, db, "formatting is reset, not refused", recalc=False)
    assert openpyxl.load_workbook(xlsx)["Journal"]["C3"].alignment.wrap_text is True


def test_a_failed_regeneration_rolls_back_and_leaves_no_trace(tmp_path, monkeypatch):
    xlsx, db = _journal_workbook(tmp_path)
    before = _bytes(xlsx)

    def boom(*args, **kwargs):
        raise RuntimeError("regeneration blew up")

    monkeypatch.setattr(generate, "generate_workbook", boom)
    with pytest.raises(RuntimeError, match="blew up"):
        ledger.journal_add(xlsx, db, "never lands", recalc=False)
    assert _bytes(xlsx) == before and _count(db, "journal") == 3
    assert not [n for n in os.listdir(tmp_path) if n.startswith(".landry-gen-")]


def test_recalc_errors_leave_the_workbook_unchanged(tmp_path, monkeypatch):
    xlsx, db = _journal_workbook(tmp_path)
    before = _bytes(xlsx)
    monkeypatch.setattr(generate, "generate_workbook", lambda *a, **k: {
        "tabs": {}, "capped_rows": [], "recalc": {"status": "success", "total_errors": 2}})
    with pytest.raises(LedgerError, match="recalc reported"):
        ledger.journal_add(xlsx, db, "never lands", recalc=True)
    assert _bytes(xlsx) == before and _count(db, "journal") == 3
    assert not [n for n in os.listdir(tmp_path) if n.startswith(".landry-gen-")]


def test_swapping_in_the_copy_keeps_the_workbooks_file_mode(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    os.chmod(xlsx, 0o644)
    ledger.journal_add(xlsx, db, "mode check", recalc=False)
    assert os.stat(xlsx).st_mode & 0o777 == 0o644               # mkstemp files start at 0600


def test_journal_edit_by_sheet_row(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    result = ledger.journal_edit(xlsx, db, 4, {"notes": "SUPERSEDED -- see row 6"}, recalc=False)
    assert result["row"] == 4
    ws = openpyxl.load_workbook(xlsx)["Journal"]
    assert ws["C4"].value == "SUPERSEDED -- see row 6"
    assert ws["C3"].value == ENTRIES[0][2] and ws["B4"].value == ENTRIES[1][1]   # neighbours untouched
    ledger.journal_edit(xlsx, db, 3, {"label": ""}, recalc=False)               # clearing a label is legitimate
    assert openpyxl.load_workbook(xlsx)["Journal"]["B3"].value is None


def test_journal_edit_rejects_a_row_that_does_not_exist(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    with pytest.raises(LedgerError, match="rows 3-5"):
        ledger.journal_edit(xlsx, db, 99, {"notes": "x"}, recalc=False)


# --------------------------------------------------------- Drawdown Log --

def test_drawdown_add_lands_in_date_order_with_its_formulas(tmp_path):
    xlsx, db = _drawdown_workbook(tmp_path)
    result = ledger.drawdown_add(xlsx, db, "2026-08-20", 720000, "backfilled", recalc=False)
    assert result["row"] == 4                                    # between 8/11 and 9/04, not appended
    ws = openpyxl.load_workbook(xlsx)[generate.DRAWDOWN_SHEET]
    assert ws["A4"].value == datetime.datetime(2026, 8, 20) and ws["B4"].value == 720000
    assert ws["H4"].value == "backfilled" and ws["A5"].value == datetime.datetime(2026, 9, 4)
    assert ws["C5"].value == '=IF(B5="","",MAX(C4,B5))'
    assert generate.verify_drawdown_log(xlsx, db) == []


def test_drawdown_add_refuses_a_duplicate_date(tmp_path):
    xlsx, db = _drawdown_workbook(tmp_path)
    before = _bytes(xlsx)
    with pytest.raises(LedgerError, match="drawdown edit"):
        ledger.drawdown_add(xlsx, db, "2026-09-04", 1.0, recalc=False)
    assert _bytes(xlsx) == before and _count(db, "drawdown_log") == 3


def test_drawdown_edit_changes_value_and_notes(tmp_path):
    xlsx, db = _drawdown_workbook(tmp_path)
    ledger.drawdown_edit(xlsx, db, "2026-09-30", {"portfolio_value": 777000.5, "notes": "restated"}, recalc=False)
    ws = openpyxl.load_workbook(xlsx)[generate.DRAWDOWN_SHEET]
    assert ws["B5"].value == 777000.5 and ws["H5"].value == "restated"
    with pytest.raises(LedgerError, match="no entry for 2026-01-01"):
        ledger.drawdown_edit(xlsx, db, "2026-01-01", {"notes": "x"}, recalc=False)


def test_a_hand_edited_drawdown_formula_is_drift(tmp_path):
    xlsx, db = _drawdown_workbook(tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    wb[generate.DRAWDOWN_SHEET]["E4"] = '=IF(D4>=-0.15,"Normal","Elevated")'     # someone moved a band by hand
    wb.save(xlsx)
    with pytest.raises(DriftError, match="regenerate restores"):
        ledger.drawdown_add(xlsx, db, "2026-10-31", 780000, recalc=False)


# ------------------------------------------------------------ the guards --

def test_excel_open_refuses_unless_forced(tmp_path, monkeypatch):
    xlsx, db = _journal_workbook(tmp_path)
    monkeypatch.setattr(ledger, "_excel_has_open", lambda path: True)
    before = _bytes(xlsx)
    with pytest.raises(ExcelOpenError, match="--force"):
        ledger.journal_add(xlsx, db, "blocked", recalc=False)
    assert _bytes(xlsx) == before
    ledger.journal_add(xlsx, db, "forced through", recalc=False, force=True)
    assert openpyxl.load_workbook(xlsx)["Journal"]["C6"].value == "forced through"


def test_a_stale_lock_file_only_counts_while_excel_is_running(tmp_path, monkeypatch):
    xlsx = str(tmp_path / "wb.xlsx")
    assert _REAL_EXCEL_CHECK(xlsx) is False                      # no lock file at all
    (tmp_path / "~$wb.xlsx").write_bytes(b"lock")
    monkeypatch.setattr(ledger.subprocess, "run", lambda *a, **k: types.SimpleNamespace(returncode=1))
    assert _REAL_EXCEL_CHECK(xlsx) is False                      # lock left behind, Excel not running
    monkeypatch.setattr(ledger.subprocess, "run", lambda *a, **k: types.SimpleNamespace(returncode=0))
    assert _REAL_EXCEL_CHECK(xlsx) is True


def test_an_old_schema_database_is_refused_clearly(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    os.remove(db)
    conn = sqlite3.connect(db)
    conn.execute("CREATE TABLE journal (id INTEGER PRIMARY KEY, date TEXT, ticker TEXT, notes TEXT)")
    conn.commit()
    conn.close()                                                  # a Phase A file: no user_version stamp
    with pytest.raises(models.SchemaMismatch, match="derived file"):
        ledger.journal_add(xlsx, db, "x", recalc=False)


def test_pull_replaces_stale_rows_and_restarts_ids(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    conn = sqlite3.connect(db)
    models.journal_add(conn, "2026-12-31", "STALE", "only in the database")
    conn.commit()
    conn.close()
    assert ledger.pull(xlsx, db) == {"journal": 3}
    conn = sqlite3.connect(db)
    assert [r[0] for r in conn.execute("SELECT id FROM journal ORDER BY id")] == [1, 2, 3]
    conn.close()


# ------------------------------------------------------------ regenerate --

def test_regenerate_resets_a_hand_edit_from_the_database(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    wb["Journal"]["C4"] = "hand-edited in Excel"
    wb.save(xlsx)
    assert ledger.status(xlsx, db)["journal"]["value_diffs"]
    ledger.regenerate(xlsx, db, recalc=False)
    assert openpyxl.load_workbook(xlsx)["Journal"]["C4"].value == ENTRIES[1][2]
    assert ledger.status(xlsx, db)["journal"]["value_diffs"] == []


def test_regenerate_needs_a_database_and_honours_the_excel_guard(tmp_path, monkeypatch):
    xlsx, db = _journal_workbook(tmp_path)
    with pytest.raises(LedgerError, match="db pull"):
        ledger.regenerate(xlsx, str(tmp_path / "nope.db"), recalc=False)
    monkeypatch.setattr(ledger, "_excel_has_open", lambda path: True)
    with pytest.raises(ExcelOpenError):
        ledger.regenerate(xlsx, db, recalc=False)


def test_an_empty_edit_note_is_refused(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    with pytest.raises(LedgerError, match="needs notes"):
        ledger.journal_edit(xlsx, db, 3, {"notes": " "}, recalc=False)


# -------------------------------------------------------------------- CLI --

def _run(*argv):
    return cli.main(list(argv))


def test_cli_journal_add_reads_multiline_notes_from_stdin(tmp_path, monkeypatch, capsys):
    xlsx, db = _journal_workbook(tmp_path)
    note = "first paragraph\n\nsecond has 'single' and \"double\" quotes"
    monkeypatch.setattr(sys, "stdin", io.StringIO(note + "\n"))
    code = _run("journal", "add", "--workbook", xlsx, "--db", db, "--label", "CLI-TEST",
                "--notes-file", "-", "--no-recalc")
    out = capsys.readouterr().out
    assert code == 0 and "Journal row 6 added" in out and "python -m landry audit" in out
    ws = openpyxl.load_workbook(xlsx)["Journal"]
    assert (ws["B6"].value, ws["C6"].value) == ("CLI-TEST", note)        # trailing newline trimmed


def test_cli_reports_a_refusal_without_a_traceback(tmp_path, capsys):
    xlsx, db = _journal_workbook(tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    wb["Journal"]["C4"] = "hand-edited"
    wb.save(xlsx)
    code = _run("journal", "add", "--workbook", xlsx, "--db", db, "--notes", "x", "--no-recalc")
    err = capsys.readouterr().err
    assert code == 1 and err.startswith("! the database and the workbook disagree")
    assert _run("journal", "add", "--workbook", xlsx, "--db", db, "--notes", "x",
                "--date", "not-a-date", "--no-recalc") == 1


def test_cli_db_status_pull_roundtrip(tmp_path, capsys):
    xlsx, db = _journal_workbook(tmp_path)
    assert _run("db", "status", "--workbook", xlsx, "--db", db) == 0
    assert "content in sync" in capsys.readouterr().out
    wb = openpyxl.load_workbook(xlsx)
    wb["Journal"]["C4"] = "hand-edited"
    wb.save(xlsx)
    assert _run("db", "status", "--workbook", xlsx, "--db", db) == 1
    out = capsys.readouterr().out
    assert "CONTENT DIFFERS" in out and "db pull" in out
    assert _run("db", "pull", "--workbook", xlsx, "--db", db) == 0
    assert _run("db", "status", "--workbook", xlsx, "--db", db) == 0
    assert _run("db", "status", "--workbook", xlsx, "--db", str(tmp_path / "none.db")) == 1


def test_cli_drawdown_add_and_edit(tmp_path, capsys):
    xlsx, db = _drawdown_workbook(tmp_path)
    assert _run("drawdown", "add", "--workbook", xlsx, "--db", db, "--date", "2026-10-31",
                "--value", "780000.5", "--notes", "month-end", "--no-recalc") == 0
    assert "Portfolio Drawdown Log row 6 added" in capsys.readouterr().out
    assert _run("drawdown", "edit", "--workbook", xlsx, "--db", db, "--date", "2026-10-31",
                "--value", "781000", "--no-recalc") == 0
    ws = openpyxl.load_workbook(xlsx)[generate.DRAWDOWN_SHEET]
    assert ws["B6"].value == 781000 and ws["H6"].value == "month-end"
    assert _run("drawdown", "add", "--workbook", xlsx, "--db", db, "--date", "2026-10-31",
                "--value", "1", "--no-recalc") == 1                      # duplicate date, refused cleanly


def test_cli_edit_needs_something_to_change(tmp_path):
    xlsx, db = _journal_workbook(tmp_path)
    with pytest.raises(SystemExit, match="nothing to change"):
        _run("journal", "edit", "--workbook", xlsx, "--db", db, "--row", "3", "--no-recalc")


# ---------------------------------------------- the repo's real workbook --

def test_a_write_to_a_scratch_copy_of_the_real_workbook_touches_only_the_journal(tmp_path):
    source = xlsx_io.latest_workbook(REPO_ROOT)
    if not source:
        pytest.skip("no tracked workbook in this checkout")
    xlsx, db = str(tmp_path / "wb.xlsx"), str(tmp_path / "landry.db")
    shutil.copy(source, xlsx)
    before = openpyxl.load_workbook(xlsx)
    counts = ledger.pull(xlsx, db)
    result = ledger.journal_add(xlsx, db, "integration-test entry", label="TEST", recalc=False)
    assert result["row"] == generate.FIRST_ROW + counts["journal"]

    after = openpyxl.load_workbook(xlsx)
    assert after["Journal"].cell(row=result["row"], column=3).value == "integration-test entry"
    for name in before.sheetnames:
        if name == "Journal":
            continue
        for row in before[name].iter_rows():
            for cell in row:
                assert _plain(after[name][cell.coordinate].value) == _plain(cell.value), (name, cell.coordinate)
    assert [d for d in generate.verify_journal(xlsx, db) if d[0] == "value"] == []
