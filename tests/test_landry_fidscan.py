"""landry.fidscan (scan parser, ranks, comparison, forward price change) and open_items.add, on synthetic inputs."""
import datetime as dt

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

from landry import fidscan, open_items

RAW = """AAA
PERF 1m -5.83%, 3m 9.65%, 6m 1.24%, YTD -30.78%, 1y -30.54%, 2y -51.18%, 5y -58.00%
EVT DEC 09 AAA to announce Q4 earnings (confirmed)
EPS Beat by $0.11; Missed by $-0.04
PE ttm 13.52 ~ 51.98; 5y 36.30 ~ 63.25; fwd(co~ind~peer) 9.90 ~ 51.71
ESS Very Bullish 9.3
OPS Trading Central:Sell(98), Zacks:Neutral(76), Jefferson:Buy(69), ISS-EVA:Outperform(7), I/B/E/S Estimate:--(N/A)
BBB
PERF 1m NA, 3m NA, 6m NA, YTD NA, 1y NA, 2y NA, 5y NA
EVT NA
PE ttm 5,500.80 ~ --; 5y -- ~ 63.25; fwd(co~ind~peer) NA
ESS Bearish 3
OPS Trading Central:Sell(94), McLean:Underperform(56)
CCC
PERF 1m 1.00%, 3m -20.00%, 6m 1.00%, YTD 1.00%, 1y 1.00%, 2y 1.00%, 5y 1.00%
ESS Neutral 5.1
OPS Zacks:Neutral(52)
"""
LANDRY = {"AAA": {"composite": 60.0, "decision": "BUY", "price": 100.0},
          "BBB": {"composite": 90.0, "decision": "STRONG BUY", "price": 50.0},
          "CCC": {"composite": 75.0, "decision": "BUY", "price": 10.0}}


def test_parse_raw_reads_every_field():
    d = fidscan.parse_raw(RAW)
    assert list(d) == ["AAA", "BBB", "CCC"]
    a = d["AAA"]
    assert a["ess"] == 9.3 and a["ess_label"] == "Very Bullish"
    assert a["perf"]["3m"] == 9.65 and a["perf"]["5y"] == -58.0
    assert (a["pe_ttm"], a["pe_ttm_industry"], a["pe_5y"]) == (13.52, 51.98, 36.30)
    assert a["earnings"].startswith("DEC 09") and len(a["surprises"]) == 2
    assert fidscan.split(a["opinions"]) == {"positive": 2, "neutral": 1, "negative": 1}    # the "--" row is dropped
    b = d["BBB"]
    assert b["ess"] == 3.0 and b["perf"]["3m"] is None and b["earnings"] is None
    assert b["pe_ttm"] == 5500.80 and b["pe_ttm_industry"] is None and b["pe_5y"] is None
    assert fidscan.split(b["opinions"])["negative"] == 2


def test_parse_raw_refuses_a_block_without_a_score():
    with pytest.raises(fidscan.FidScanError):
        fidscan.parse_raw("AAA\nPERF 1m 1.00%\n")
    with pytest.raises(fidscan.FidScanError):
        fidscan.parse_raw("nothing here")


def test_ranks_and_rank_correlation():
    assert fidscan.ranks([10, 30, 20]) == [3.0, 1.0, 2.0]
    assert fidscan.ranks([5, 5, 1]) == [1.5, 1.5, 3.0]
    assert fidscan.rank_correlation([1, 2, 3, 4], [10, 20, 30, 40]) == pytest.approx(1.0)
    assert fidscan.rank_correlation([1, 2, 3, 4], [40, 30, 20, 10]) == pytest.approx(-1.0)
    # 1 - 6*sum(d^2)/(n(n^2-1)) with ranks (1,2,3,4,5) vs (2,1,4,3,5): d^2 sum = 4 -> 0.8
    assert fidscan.rank_correlation([5, 4, 3, 2, 1], [4, 5, 2, 3, 1]) == pytest.approx(0.8)
    assert fidscan.rank_correlation([1, 2], [1, 2]) is None


def test_compare_flags_the_disagreements_and_forward_measures_them():
    store = {"scans": []}
    scan = fidscan.add(store, RAW, dt.date(2026, 10, 9), LANDRY, "test")
    rows = {r["ticker"]: r for r in fidscan.compare(scan)}
    assert rows["BBB"]["gap"] == -2 and rows["BBB"]["flag"] == "LANDRY HIGHER"       # Landry 1st, Fidelity 3rd
    assert rows["AAA"]["gap"] == 2 and rows["AAA"]["flag"] == "FIDELITY HIGHER"
    assert rows["CCC"]["gap"] == 0 and rows["CCC"]["flag"] == ""
    assert fidscan.forward(scan, {"AAA": 100.0, "BBB": 50.0, "CCC": 10.0}) is None   # no later prices yet
    fw = fidscan.forward(scan, {"AAA": 110.0, "BBB": 45.0, "CCC": 10.0})
    assert fw["groups"]["FIDELITY HIGHER"] == (pytest.approx(0.10), 1)
    assert fw["groups"]["LANDRY HIGHER"] == (pytest.approx(-0.10), 1)
    assert fw["groups"]["all"][0] == pytest.approx(0.0) and fw["groups"]["all"][1] == 3
    text = fidscan.format_report(store, {"AAA": 110.0, "BBB": 45.0, "CCC": 10.0})
    assert "since 2026-10-09" in text and "LANDRY HIGHER" in text
    fidscan.add(store, RAW, dt.date(2026, 10, 9), LANDRY)                              # same date replaces
    assert len(store["scans"]) == 1


def _oi_book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = open_items.SHEET
    ws.append(["Open items"]); ws.append([]); ws.append([])
    ws.append(["#", "Item", "Done (Y/N)", "Date Raised", "Date Resolved", "Notes"])
    ws.append([1, "first", None, dt.datetime(2026, 10, 1), None, "n1"])
    ws.append([2, "second", "Y", dt.datetime(2026, 10, 2), dt.datetime(2026, 10, 3), "n2"])
    ws.add_table(Table(displayName=open_items.TABLE, ref="A4:F6"))
    sr = wb.create_sheet("Schema Reference")
    sr.append(["Table", open_items.TABLE, "A4:F6 as of 2026-10-02 (2 items) -- grows"])
    wb.save(path)


def test_open_items_add_appends_and_grows_the_table(tmp_path):
    p = str(tmp_path / "wb.xlsx")
    _oi_book(p)
    assert open_items.add(p, "  third   item ", "a note", when=dt.date(2026, 10, 10), recalc=False) == "#3 added (row 7)"
    wb = openpyxl.load_workbook(p)
    ws = wb[open_items.SHEET]
    assert [ws.cell(7, c).value for c in (1, 2, 3, 6)] == [3, "third item", None, "a note"]
    assert ws.cell(7, 4).value == dt.datetime(2026, 10, 10)
    assert ws.tables[open_items.TABLE].ref == "A4:F7"
    assert wb["Schema Reference"]["C1"].value == "A4:F7 as of 2026-10-10 (3 items) -- grows"
    with pytest.raises(open_items.OpenItemsError):
        open_items.add(p, "   ", recalc=False)
