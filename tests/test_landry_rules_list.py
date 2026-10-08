"""Tests for landry.rules_list -- the printable Hard Rules list, and the check that keeps it in step with the rulebook.

The rule numbers are one running sequence across the whole rulebook, so inserting a rule renumbers everything after
it (v1.04's Rule 5 did) and every hand-kept cheat sheet goes stale without a sound. ``test_the_list_matches_the_live_
rulebook`` is the alarm: it reads the rulebook docx itself."""

import copy
import datetime as dt

import pytest

pytest.importorskip("docx")

from landry import cli, rules_list as rl                         # noqa: E402

needs_rulebook = pytest.mark.skipif(rl.rulebook_path() is None, reason="no rulebook docx in the repo root")


def _numbers(parts=None):
    return [n for part in (parts or rl.PARTS) for g in part["groups"] for n, _label, _text in g["rules"]]


def test_the_list_numbers_every_rule_once_from_one_in_order():
    assert _numbers() == list(range(1, 52))
    assert [p["title"].split(" — ")[0] for p in rl.PARTS] == [f"PART {n}" for n in (3, 4, 5, 6, 7, 8, 9, 10)]


@needs_rulebook
def test_the_list_matches_the_live_rulebook():
    """Fails when the rulebook gains, loses or renumbers a rule, changes a threshold, or moves to a new version: update
    PARTS in landry/rules_list.py (wording, numbering, VERSION, REVISED), then `python -m landry rules-list --pdf`."""
    assert rl.check() == []


@needs_rulebook
def test_the_rulebook_reader_finds_the_numbered_rules_of_the_index():
    rb = rl.rulebook_rules(rl.rulebook_path())
    assert sorted(rb) == list(range(1, 52))
    assert rb[5].startswith("Any Relative Strength vs. SPY or Technical Trend")        # v1.04's inserted rule
    assert rb[34].startswith("Early Partial Trim") and rb[38].startswith("Quarterly")   # shifted by it
    assert "Entry Checklist" in rb[36] and "documented before execution" in rb[36]      # the a)-e) sub-items belong to 36
    assert rb[51].startswith("Do not initiate a new financial-company position")
    assert rb[16].startswith("No single sector shall exceed 25% of portfolio") and "market value" in rb[16]   # v1.05's basis
    assert rl.rulebook_version(rl.rulebook_path()) == rl.VERSION


@needs_rulebook
def test_check_catches_a_threshold_the_rulebook_does_not_have(monkeypatch):
    parts = copy.deepcopy(rl.PARTS)
    for part in parts:
        for g in part["groups"]:
            g["rules"] = [(n, label, text.replace("8% of portfolio", "9% of portfolio") if n == 15 else text)
                          for n, label, text in g["rules"]]
    monkeypatch.setattr(rl, "PARTS", parts)
    problems = rl.check()
    assert len(problems) == 1 and problems[0].startswith("Rule 15:") and "9" in problems[0]


@needs_rulebook
def test_check_catches_a_rule_the_list_does_not_have(monkeypatch):
    parts = copy.deepcopy(rl.PARTS)
    parts[-1]["groups"][-1]["rules"] = parts[-1]["groups"][-1]["rules"][:-1]             # drop Rule 51
    monkeypatch.setattr(rl, "PARTS", parts)
    problems = rl.check()
    assert any("numbers 51 rules, this list 50" in p and "rulebook only: [51]" in p for p in problems)


@needs_rulebook
def test_check_catches_a_new_rulebook_version(monkeypatch):
    monkeypatch.setattr(rl, "VERSION", "1.04")
    assert rl.check()[0] == "the rulebook is v1.05, this list is v1.04"


def test_the_builder_writes_a_docx_with_every_rule_and_the_quick_reference(tmp_path):
    from docx import Document
    out = rl.build(str(tmp_path / "rules.docx"), today=dt.date(2026, 10, 5))
    doc = Document(out)
    cells = [c.text.strip() for t in doc.tables for r in t.rows for c in r.cells]
    assert all(str(n) in cells for n in range(1, 52))                                   # a number cell for every rule
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Version 1.05 (revised 2026-10-08)" in text and "51 numbered Hard Rules" in text
    assert any("Stop-Review Ladder" in p.text for p in doc.paragraphs) and any("Rung 3:" in p.text for p in doc.paragraphs)
    assert "Printed October 5, 2026" in text
    for heading in ("QUICK REFERENCE", "Decision bands (Part 3)", "Portfolio drawdown response (Part 7, Rule 39)",
                    "Macro overlay (Part 7)", "Hold Through: do NOT sell based on these alone (Part 6)",
                    "When rules conflict: higher rank wins (Part 10)", "What to avoid (Part 11)"):
        assert heading in text, heading
    assert any(c.startswith("PART 6") for c in cells) and any("Early Partial Trim" in c for c in cells)
    assert "Hard Rules List" in doc.sections[0].footer.paragraphs[0].text


@needs_rulebook
def test_the_command_checks_first_and_refuses_to_write_a_list_that_no_longer_matches(tmp_path, monkeypatch, capsys):
    assert cli.main(["rules-list", "--check"]) == 0
    out = tmp_path / "r.docx"
    assert cli.main(["rules-list", "--out", str(out)]) == 0 and out.exists()
    monkeypatch.setattr(rl, "VERSION", "1.03")                                           # now the list is out of date
    stale = tmp_path / "stale.docx"
    assert cli.main(["rules-list", "--out", str(stale)]) == 1 and not stale.exists()
    assert "not written" in capsys.readouterr().err
    assert cli.main(["rules-list", "--out", str(stale), "--force"]) == 0 and stale.exists()


@needs_rulebook
def test_the_stop_review_ladder_is_in_the_rulebook_and_unnumbered():
    """v1.05's ladder is an un-numbered Hard Rules block (like Hold Through), so it must not have moved any rule number."""
    from docx import Document
    text = "\n".join(p.text for p in Document(rl.rulebook_path()).paragraphs)
    assert "Hard Rules: Stop-Review Ladder" in text and "Rung 3 — 40% or more below cost" in text
    assert "unless a Stop-Review Ladder rung (below) has been reached" in text                 # Hold Through's first condition
    rb = rl.rulebook_rules(rl.rulebook_path())
    assert not any("Rung" in v for v in rb.values()) and len(rb) == 51
    assert rb[34].startswith("Early Partial Trim") and rb[38].startswith("Quarterly")           # nothing shifted
    lines = rl.STOP_REVIEW + [rl.STOP_REVIEW_NOTE]
    assert rl.check() == [] and all(l in rl._quick_reference_lines() for l in lines)
