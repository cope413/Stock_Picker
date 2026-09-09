"""Offline tests for landry.audit -- synthetic workbooks only, no
dependency on the real LANDRY_SYSTEM_WORKBOOK_*.xlsx."""

import subprocess

import openpyxl
from openpyxl.worksheet.table import Table, TableColumn

from landry.audit import (
    check_cross_tab_references,
    check_page_setup_vs_last_commit,
    check_reader_bounds,
    check_schema_reference,
    check_table_refs,
    report,
)
from landry.doctor import Check


def _add_table(ws, name, ref):
    tbl = Table(displayName=name, ref=ref)
    ws.add_table(tbl)


def test_table_ref_flags_a_genuinely_missed_row(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Ticker", "Company", "Qty"])
    ws.append(["MU", "Micron", 10])
    ws.append(["NEW", "New Position", 5])  # added past the table's ref
    _add_table(ws, "T", "A1:C2")
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    checks = check_table_refs(str(p))
    assert len(checks) == 1
    assert not checks[0].ok
    assert "row 3" in checks[0].detail


def test_table_ref_ignores_a_sparse_footnote_row(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Ticker", "Company", "Qty"])
    ws.append(["MU", "Micron", 10])
    ws.append(["This is a freeform footnote, not a data row", None, None])
    _add_table(ws, "T", "A1:C2")
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    checks = check_table_refs(str(p))
    assert len(checks) == 1
    assert checks[0].ok


def test_reader_bounds_flags_a_stale_hardcoded_bound(tmp_path, monkeypatch):
    import landry.xlsx_io as xio

    def fake_reader(path: str, sheet: str = "Fake") -> list:
        """A=Ticker,B=Company."""
        wb, ws = xio._open(path, sheet)
        out = []
        for r in ws.iter_rows(min_row=3, max_row=5, values_only=True):
            if not r or not r[0]:
                continue
            out.append(dict(ticker=r[0]))
        wb.close()
        return out

    monkeypatch.setattr(xio, "read_fake_tab", fake_reader, raising=False)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fake"
    ws.append(["header"])
    ws.append(["header2"])
    for i in range(3, 8):
        ws.append([f"T{i}", "Co"])
    p = tmp_path / "wb.xlsx"
    wb.save(p)

    checks = check_reader_bounds(str(p))
    hit = [c for c in checks if c.name == "reader_bound:read_fake_tab"]
    assert len(hit) == 1
    assert not hit[0].ok
    assert "row 7" in hit[0].detail


def test_cross_tab_reference_flags_blank_target_column(tmp_path):
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "A"
    a["A1"] = "=Target!C2"
    b = wb.create_sheet("Target")
    b["A2"] = "Ticker"
    b["B2"] = "Company"
    # C2 deliberately left blank -- the formula points at a dead column
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    checks = check_cross_tab_references(str(p))
    bad = [c for c in checks if not c.ok]
    assert any("Target" in c.detail and "C2" in c.detail for c in bad)


def test_cross_tab_reference_flags_truncated_range(tmp_path):
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "A"
    a["A1"] = "=COUNTIF(Target!B3:B5,\"x\")"
    b = wb.create_sheet("Target")
    for r in range(3, 8):
        b.cell(row=r, column=2, value="x")  # real data through row 7
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    checks = check_cross_tab_references(str(p))
    bad = [c for c in checks if not c.ok]
    assert any("row 7" in c.detail for c in bad)


def test_cross_tab_reference_clean_workbook_reports_ok(tmp_path):
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "A"
    a["A1"] = "=Target!B2"
    b = wb.create_sheet("Target")
    b["B2"] = "Header"
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    checks = check_cross_tab_references(str(p))
    assert len(checks) == 1
    assert checks[0].ok


def test_page_setup_skips_cleanly_when_not_a_git_repo(tmp_path, monkeypatch):
    def fake_run(*a, **k):
        raise subprocess.CalledProcessError(128, "git")
    monkeypatch.setattr(subprocess, "run", fake_run)

    wb = openpyxl.Workbook()
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    checks = check_page_setup_vs_last_commit(str(p), repo_dir=str(tmp_path))
    assert len(checks) == 1
    assert checks[0].ok
    assert "skipped" in checks[0].detail


def test_schema_reference_flags_wrong_table_name(tmp_path):
    wb = openpyxl.Workbook()
    sr = wb.active
    sr.title = "Schema Reference"
    sr["A4"] = "Target"
    sr["A6"] = "Table"
    sr["B6"] = "WrongName"
    sr["C6"] = "A2:B10"
    target = wb.create_sheet("Target")
    target["A2"] = "Ticker"
    target["B2"] = "Company"
    _add_table(target, "RealName", "A2:B10")
    p = tmp_path / "wb.xlsx"
    wb.save(p)
    checks = check_schema_reference(str(p))
    hit = [c for c in checks if c.name == "schema_ref:Target/table_name"]
    assert len(hit) == 1
    assert not hit[0].ok


def test_report_formats_pass_and_fail_counts():
    checks = [Check("a", True, "fine"), Check("b", False, "broken", fix="do X")]
    text = report(checks, "wb.xlsx")
    assert "wb.xlsx" in text
    assert "1 passed, 1 failed" in text
    assert "do X" in text
    assert "Review the failures" in text


def test_report_all_clean():
    text = report([Check("a", True, "fine")])
    assert "No structural drift detected." in text
