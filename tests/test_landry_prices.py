"""Offline tests for landry.prices -- synthetic workbooks only (no network, no LibreOffice).

The fixture mirrors the real tabs' structure: Price History (header row 2, Friday rows from 3, a
'NN Weeks' footer two rows under the data), Returns (Calc) (formulas ONE ROW AHEAD of the data, a
'trend' column, a chart below the data) and a 21x21 Correlation Matrix with its count column,
Rule 38 tally and colour rules."""

import datetime as dt

import openpyxl
import pandas as pd
import pytest
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Font, PatternFill

from landry import prices
from landry.prices import PricesError

FIRST_FRIDAY = dt.datetime(2026, 8, 21)


def _friday(i):
    return FIRST_FRIDAY + dt.timedelta(days=7 * i)


def _pf(col_letter, r):
    return (f"=IFERROR(IF(OR(NOT(ISNUMBER('Price History'!{col_letter}{r + 1})),"
            f"NOT(ISNUMBER('Price History'!{col_letter}{r})),'Price History'!{col_letter}{r}=0),\"\","
            f"'Price History'!{col_letter}{r + 1}/'Price History'!{col_letter}{r}-1),\"\")")


def build_workbook(tmp_path, header=None, n_weeks=6, name="wb.xlsx"):
    """header: 21 entries for columns B..V (None = a free slot)."""
    from openpyxl.utils import get_column_letter as L
    header = header or [f"T{i:02d}" for i in range(1, 22)]
    assert len(header) == 21
    wb = openpyxl.Workbook()
    ph = wb.active
    ph.title = "Price History"
    ph["A1"] = "Weekly closing price ($), week-ending Friday"
    ph["A2"] = "Week Ending"
    head_fill = PatternFill(fill_type="solid", fgColor="FFDDEBF7")
    for k, t in enumerate(header):
        c = ph.cell(row=2, column=2 + k)
        c.value = t
        c.fill = head_fill
        ph.column_dimensions[L(2 + k)].width = 6.7
    for i in range(n_weeks):
        r = 3 + i
        ph.cell(row=r, column=1, value=_friday(i)).number_format = "mm/dd/yyyy"
        for k, t in enumerate(header):
            if t:
                ph.cell(row=r, column=2 + k, value=round(100 + k + i * 1.5, 2)).number_format = "#,##0.00"
    last = 2 + n_weeks
    ph.cell(row=last + 3, column=1, value=f"=COUNT(A3:A{last + 1})")
    ph.cell(row=last + 3, column=2, value="Weeks")

    ret = wb.create_sheet("Returns (Calc)")
    ret["A2"] = "Week Ending"
    for k in range(21):
        ret.cell(row=2, column=2 + k, value=f"='Price History'!{L(2 + k)}2").fill = head_fill
        ret.column_dimensions[L(2 + k)].width = 5.7
    ret.cell(row=2, column=23, value="T01 trend")
    for r in range(3, last + 1):                                        # one row ahead: through the last data row
        ret.cell(row=r, column=1, value=f"=IF('Price History'!A{r + 1}=\"\",\"\",'Price History'!A{r + 1})").number_format = "mm/dd/yyyy"
        for k in range(21):
            ret.cell(row=r, column=2 + k, value=_pf(L(2 + k), r)).number_format = "0.0%"
        ret.cell(row=r, column=23, value="=IFERROR('Price History'!B3+(INDEX('Price History'!B:B,MATCH(9.99E+307,"
                                         "'Price History'!$A:$A))-'Price History'!B3)*(ROW()-3)/(MATCH(9.99E+307,"
                                         "'Returns (Calc)'!$A:$A)-3),\"\")")      # the real trend formula: identical in every row
    ch = LineChart()
    ch.add_data(Reference(ret, min_col=2, min_row=3, max_row=40))
    ret.add_chart(ch, f"A{last + 2}")

    cm = wb.create_sheet("Correlation Matrix")
    cm["A1"] = "Pairwise correlation"
    cm["A2"] = "Positions exceeding cap"
    cm["J2"] = '=COUNTIF(AG6:AG26,">1")'
    cm["J3"] = '=IF(J2>0,"REVIEW REQUIRED","OK")'
    cm["A5"] = "Ticker"
    cm["AG5"] = "Positions >0.70 with"
    label_fill = PatternFill(fill_type="solid", fgColor="FFDCE6F1")
    for k in range(21):
        col = 2 + k
        cm.cell(row=5, column=col, value=f"='Returns (Calc)'!{L(col)}2").fill = head_fill
        cm.cell(row=6 + k, column=1, value=f"='Returns (Calc)'!{L(col)}2").fill = label_fill
        cm.column_dimensions[L(col)].width = 7.5
        for j in range(21):
            cell = cm.cell(row=6 + k, column=2 + j)
            cell.number_format = "0.00"
            cell.value = 1 if j == k else (f"=IFERROR(ROUND(CORREL('Returns (Calc)'!${L(col)}$3:${L(col)}$500,"
                                           f"'Returns (Calc)'!${L(2 + j)}$3:${L(2 + j)}$500),2),\"\")")
        cm.cell(row=6 + k, column=33, value=f'=COUNTIF(B{6 + k}:V{6 + k},">0.7")-1')
    red = PatternFill(fill_type="solid", bgColor="FFFFC7CE")
    cm.conditional_formatting.add("B6:V26", FormulaRule(formula=["AND(B6<>1,ABS(B6)>0.7)"], fill=red))
    cm.conditional_formatting.add("AG6:AG26", CellIsRule(operator="greaterThan", formula=["1"], fill=red))
    path = tmp_path / name
    wb.save(path)
    return str(path)


def fake_fetch(extra_weeks=1, tickers=None, nan=(), jump=None):
    """A weekly frame covering the fixture's weeks plus ``extra_weeks`` more."""
    def fetch(requested):
        idx = pd.DatetimeIndex([_friday(i) for i in range(0, 6 + extra_weeks)])
        cols = tickers or list(requested)
        data = {t: [round(100 + (hash(t) % 7) + i * 1.5 + 0.004, 4) for i in range(len(idx))] for t in cols}
        df = pd.DataFrame(data, index=idx)
        for t in nan:
            if t in df.columns:
                df.loc[idx[-1], t] = float("nan")
        if jump:
            df.loc[idx[-1], jump[0]] = df.loc[idx[-2], jump[0]] * jump[1]
        return df
    return fetch


NOW_AFTER_FRIDAY_10_2 = dt.datetime(2026, 10, 3, 9, 0)                  # Saturday: the 10/2 week has closed


# ---------------------------------------------------------------- layout / dates --

def test_read_layout_rejects_a_duplicated_ticker(tmp_path):
    header = [f"T{i:02d}" for i in range(1, 22)]
    header[17] = header[13]                                              # the 10/2 failure: SPMO in two columns
    wb = openpyxl.load_workbook(build_workbook(tmp_path, header=header))
    with pytest.raises(PricesError, match="twice"):
        prices.read_layout(wb["Price History"])


def test_read_layout_finds_free_slots(tmp_path):
    header = [f"T{i:02d}" for i in range(1, 22)]
    header[14] = None
    header[19] = None
    layout = prices.read_layout(openpyxl.load_workbook(build_workbook(tmp_path, header=header))["Price History"])
    assert layout.gaps == [16, 21]                                       # P and U
    assert layout.last_row == 8 and len(layout.date_rows) == 6


def test_last_completed_friday_never_returns_an_unfinished_week():
    f = prices.last_completed_friday
    assert f(dt.datetime(2026, 10, 2, 17, 0)) == dt.date(2026, 10, 2)    # Friday after the close
    assert f(dt.datetime(2026, 10, 2, 15, 0)) == dt.date(2026, 9, 25)    # Friday during the session
    assert f(dt.datetime(2026, 10, 3, 9, 0)) == dt.date(2026, 10, 2)     # Saturday
    assert f(dt.datetime(2026, 10, 5, 9, 0)) == dt.date(2026, 10, 2)     # Monday
    assert f(dt.datetime(2026, 10, 8, 23, 0)) == dt.date(2026, 10, 2)    # Thursday night


# --------------------------------------------------------------------- append --

def test_append_adds_the_row_extends_returns_one_ahead_moves_footer_and_charts(tmp_path):
    path = build_workbook(tmp_path)
    assert openpyxl.load_workbook(path)["Returns (Calc)"]._charts[0].anchor._from.row == 9
    rep = prices.append_weeks(path, fetch=fake_fetch(), now=NOW_AFTER_FRIDAY_10_2)
    assert rep["appended"] == ["2026-10-02"] and not rep["warnings"]
    wb = openpyxl.load_workbook(path)
    ph, ret = wb["Price History"], wb["Returns (Calc)"]
    assert ph["A9"].value == dt.datetime(2026, 10, 2) and ph["A9"].number_format == "mm/dd/yyyy"
    assert isinstance(ph["B9"].value, (int, float)) and ph["B9"].value == round(ph["B9"].value, 2) and ph["B9"].number_format == "#,##0.00"
    assert ph["A11"].value is None and ph["B11"].value is None           # old footer cleared
    assert ph["A12"].value == "=COUNT(A3:A10)" and ph["B12"].value == "Weeks"
    # Returns (Calc): one row ahead -> formulas now reach row 9 and reference Price History rows 10 / 9
    assert "'Price History'!B10" in ret["B9"].value and "'Price History'!B9" in ret["B9"].value
    assert ret["A9"].value == "=IF('Price History'!A10=\"\",\"\",'Price History'!A10)"
    assert ret["W9"].value == ret["W8"].value and "'Price History'!B3+" in ret["W9"].value     # trend: verbatim, anchor stays on row 3
    assert ret["B9"].number_format == "0.0%"
    (chart,) = ret._charts
    assert chart.anchor._from.row == 10                                  # was 0-based 9 (row 10): pushed down with the data


def test_append_when_up_to_date_changes_nothing(tmp_path):
    path = build_workbook(tmp_path)
    before = openpyxl.load_workbook(path)["Price History"]["A8"].value
    rep = prices.append_weeks(path, fetch=fake_fetch(extra_weeks=0), now=dt.datetime(2026, 9, 26, 9, 0))
    assert rep["up_to_date"] and rep["appended"] == []
    assert openpyxl.load_workbook(path)["Price History"]["A8"].value == before


def test_append_several_weeks_lands_on_the_old_footer_without_losing_it(tmp_path):
    path = build_workbook(tmp_path)
    rep = prices.append_weeks(path, fetch=fake_fetch(extra_weeks=3), now=dt.datetime(2026, 10, 17, 9, 0))
    assert rep["appended"] == ["2026-10-02", "2026-10-09", "2026-10-16"]
    ph = openpyxl.load_workbook(path)["Price History"]
    assert [ph.cell(row=r, column=1).value.date().isoformat() for r in (9, 10, 11)] == ["2026-10-02", "2026-10-09", "2026-10-16"]
    assert ph["A14"].value == "=COUNT(A3:A12)" and ph["B14"].value == "Weeks"
    assert ph["B11"].value is not None                                    # a price, not the footer's 'Weeks'


def test_append_refuses_a_week_the_provider_does_not_have(tmp_path):
    path = build_workbook(tmp_path)
    with pytest.raises(PricesError, match="no close"):
        prices.append_weeks(path, fetch=fake_fetch(extra_weeks=1, nan=[f"T{i:02d}" for i in range(1, 12)]),
                            now=NOW_AFTER_FRIDAY_10_2)
    assert openpyxl.load_workbook(path)["Price History"]["A9"].value is None     # nothing written


def test_append_warns_and_leaves_one_missing_close_blank(tmp_path):
    path = build_workbook(tmp_path)
    rep = prices.append_weeks(path, fetch=fake_fetch(nan=["T03"]), now=NOW_AFTER_FRIDAY_10_2)
    assert any("T03" in w for w in rep["warnings"])
    ph = openpyxl.load_workbook(path)["Price History"]
    assert ph["D9"].value is None and ph["C9"].value is not None


def test_append_refuses_an_implausible_jump_unless_forced(tmp_path):
    path = build_workbook(tmp_path)
    with pytest.raises(PricesError, match="moved"):
        prices.append_weeks(path, fetch=fake_fetch(jump=("T02", 1.6)), now=NOW_AFTER_FRIDAY_10_2)
    rep = prices.append_weeks(path, fetch=fake_fetch(jump=("T02", 1.6)), now=NOW_AFTER_FRIDAY_10_2, force=True)
    assert rep["appended"] == ["2026-10-02"]


def test_append_dry_run_writes_nothing(tmp_path):
    path = build_workbook(tmp_path)
    rep = prices.append_weeks(path, fetch=fake_fetch(), now=NOW_AFTER_FRIDAY_10_2, write=False)
    assert rep["appended"] == ["2026-10-02"]
    assert openpyxl.load_workbook(path)["Price History"]["A9"].value is None


# -------------------------------------------------------------------- rebuild --

def test_rebuild_rewrites_prices_but_not_header_or_dates(tmp_path):
    path = build_workbook(tmp_path)
    rep = prices.rebuild(path, fetch=fake_fetch(extra_weeks=0))
    assert rep["weeks"] == 6 and rep["tickers"] == 21 and rep["cells_changed"] > 0
    ph = openpyxl.load_workbook(path)["Price History"]
    assert ph["B2"].value == "T01" and ph["A3"].value == _friday(0) and ph["A8"].value == _friday(5)
    assert ph["B3"].value == round(100 + (hash("T01") % 7) + 0.004, 2)


def test_rebuild_refuses_to_erase_history_the_provider_lacks(tmp_path):
    path = build_workbook(tmp_path)

    def gappy(requested):
        df = fake_fetch(extra_weeks=0)(requested)
        df.loc[df.index[:5], "T04"] = float("nan")
        return df
    with pytest.raises(PricesError, match="missing"):
        prices.rebuild(path, fetch=gappy)


# ----------------------------------------------------------------- add ticker --

def test_add_ticker_reuses_a_free_slot_without_touching_the_formula_tabs(tmp_path):
    header = [f"T{i:02d}" for i in range(1, 22)]
    header[14] = None                                                    # column P is a free slot
    path = build_workbook(tmp_path, header=header)
    info = prices.add_ticker(path, "ZZZ", fetch=fake_fetch(extra_weeks=0, tickers=["ZZZ"]))
    assert info["column"] == "P" and info["reused_slot"] and info["weeks_filled"] == 6
    wb = openpyxl.load_workbook(path)
    assert wb["Price History"]["P2"].value == "ZZZ" and wb["Price History"]["P3"].value is not None
    assert wb["Correlation Matrix"]["P20"].value == 1                    # diagonal for column 16, row 4 + 16


def test_add_ticker_beyond_v_extends_returns_and_the_matrix(tmp_path):
    path = build_workbook(tmp_path)                                      # B..V all taken
    info = prices.add_ticker(path, "NEWA", fetch=fake_fetch(extra_weeks=0, tickers=["NEWA"]))
    assert info["column"] == "W" and not info["reused_slot"]
    wb = openpyxl.load_workbook(path)
    ph, ret, cm = wb["Price History"], wb["Returns (Calc)"], wb["Correlation Matrix"]
    assert ph["W2"].value == "NEWA" and ph["W3"].value is not None and ph["W3"].number_format == "#,##0.00"
    # Returns (Calc) column AP (42) holds Price History column W, one row ahead like the rest
    assert ret["AP2"].value == "='Price History'!W2"
    assert "'Price History'!W4" in ret["AP3"].value and "'Price History'!W3" in ret["AP3"].value
    assert "'Price History'!W9" in ret["AP8"].value and "'Price History'!W8" in ret["AP8"].value   # one row ahead
    assert ret["AP9"].value is None                                      # and no further than the other columns
    # Correlation Matrix: row 27 / column W for the new ticker, via Returns column AP
    assert cm["A27"].value == "='Returns (Calc)'!AP2" and cm["W5"].value == "='Returns (Calc)'!AP2"
    assert cm["W27"].value == 1
    assert "$B$3:$B$500" in cm["W6"].value and "$AP$3:$AP$500" in cm["W6"].value
    assert "$AP$3:$AP$500" in cm["B27"].value and "$B$3:$B$500" in cm["B27"].value
    assert cm["AG6"].value == '=COUNTIF(B6:W6,">0.7")-1' and cm["AG27"].value == '=COUNTIF(B27:W27,">0.7")-1'
    assert cm["J2"].value == '=COUNTIF(AG6:AG27,">1")'
    ranges = sorted(str(cf.sqref) for cf in cm.conditional_formatting)
    assert "B6:W27" in ranges and "AG6:AG27" in ranges


def test_add_second_new_ticker_goes_to_the_next_column(tmp_path):
    path = build_workbook(tmp_path)
    prices.add_ticker(path, "NEWA", fetch=fake_fetch(extra_weeks=0, tickers=["NEWA"]))
    info = prices.add_ticker(path, "NEWB", fetch=fake_fetch(extra_weeks=0, tickers=["NEWB"]))
    assert info["column"] == "X"
    cm = openpyxl.load_workbook(path)["Correlation Matrix"]
    assert cm["A28"].value == "='Returns (Calc)'!AQ2" and cm["AG28"].value == '=COUNTIF(B28:X28,">0.7")-1'
    assert cm["J2"].value == '=COUNTIF(AG6:AG28,">1")'
    assert "$AQ$3:$AQ$500" in cm["W28"].value and "$AP$3:$AP$500" in cm["W28"].value


def test_add_ticker_rejects_a_duplicate_and_a_missing_provider_series(tmp_path):
    path = build_workbook(tmp_path)
    with pytest.raises(PricesError, match="already has a column"):
        prices.add_ticker(path, "T05", fetch=fake_fetch())
    with pytest.raises(PricesError, match="nothing"):
        prices.add_ticker(path, "NOPE", fetch=fake_fetch(tickers=["OTHER"]))


# --------------------------------------------------------------------- status --

def test_status_reports_a_duplicate_header_instead_of_raising(tmp_path):
    header = [f"T{i:02d}" for i in range(1, 22)]
    header[17] = header[13]
    st = prices.status(build_workbook(tmp_path, header=header))
    assert not st["ok"] and "twice" in st["problems"][0]


def test_status_counts_weeks_behind_and_missing_holdings(tmp_path, monkeypatch):
    class Pos:
        def __init__(self, ticker, quantity):
            self.ticker, self.quantity = ticker, quantity
    import landry.xlsx_io as xio
    monkeypatch.setattr(xio, "read_positions", lambda p: [Pos("T01", 5), Pos("GONE", 0), Pos("NEWHOLD", 3), Pos("FZDXX", 9)])
    st = prices.status(build_workbook(tmp_path), now=dt.datetime(2026, 10, 10, 9, 0))
    assert st["held_not_in_header"] == ["NEWHOLD"]
    assert st["weeks_behind"] == 2 and st["missing_fridays"][0] == dt.date(2026, 10, 2)
    assert not st["ok"]


# ---------------------------------------------------------- the old destroyer --

def test_export_no_longer_rewrites_price_history():
    from landry.export import export_workbook
    with pytest.raises(ValueError, match="landry prices"):
        export_workbook("does-not-matter.xlsx", weekly_closes=pd.DataFrame({"X": [1.0]}))
