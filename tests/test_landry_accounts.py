"""landry.accounts: the export parsers, the cash walk, pending transfers, the external-flow ledger and the write, on a
synthetic Current Positions sheet. No network, no LibreOffice, no live workbook."""

import datetime as dt
import json

import openpyxl
import pytest

from landry import accounts
from landry.accounts import Txn

D = dt.date
FID, CHA = "JT ULTRA (Fidelity)", "Self-Directed (Chase)"

CHASE_CSV = '''Trade Date,Post Date,Settlement Date,Account Name,Account Number,Account Type,Type,Description,Cusip,Ticker,Security Type,Local Currency,Price USD,Price Local,Quantity,G/L Short USD,G/L Short Local,G/L Long USDs,G/L Long Local,Amount USD,Amount Local
"10/9/2026","10/9/2026","10/9/2026","Self-Directed","...3693","Brokerage","BNK","BANKLINK ACH PUSH 82263702","","","Other","USD","0","0","0","","","","","-20000","-20000"
"10/9/2026","10/9/2026","10/9/2026","Self-Directed","...3693","Brokerage","WDL","CHASE DEPOSIT SWEEP INTRA-DAY WITHDRWAL","","QACDS","Money Market","USD","0","0","-19980.76","","","","","19980.76","19980.76"
"10/9/2026","10/9/2026","10/9/2026","Self-Directed","...3693","Brokerage","Dividend","VICTORYSHARES FREE CASH FLOW ETF CASH DIV","92647X830","VFLO","Stock","USD","0","0","0","","","","","19.24","19.24"
"10/1/2026","10/1/2026","10/1/2026","Self-Directed","...3693","Brokerage","Reinvest","VANGUARD FEDERAL REINVEST","922906300","VMFXX","Money Market","USD","0","0","111.93","","","","","-111.93","-111.93"
"10/1/2026","10/1/2026","10/1/2026","Self-Directed","...3693","Brokerage","Reinvest","JEPQ REINVEST","4","JEPQ","Stock","USD","61","61","2","","","","","-122","-122"
"10/1/2026","10/1/2026","10/1/2026","Self-Directed","...3693","Brokerage","XYZ","SOMETHING NEW","","","Other","USD","0","0","0","","","","","-1","-1"
'''

FIDELITY_CSV = '''

Run Date,Account,Account Number,Action,Symbol,Description,Type,Price ($),Quantity,Commission ($),Fees ($),Accrued Interest ($),Amount ($),Settlement Date
10/12/2026,JT ULTRA_3711,X64063711,JOURNALED JNL VS A/C TYPES (Cash),"",No Description,Cash,"",0,"","","",43.43,""
10/09/2026,JT ULTRA_3711,X64063711,DIVIDEND RECEIVED VICTORY PORTFOLIOS II (VFLO) (Margin),VFLO,VICTORY,Margin,"",0,"","","",43.43,""
10/08/2026,JT ULTRA_3711,X64063711,FOREIGN TAX PAID TAIWAN SEMICONDUCTOR (TSM) (Margin),TSM,TAIWAN,Margin,"",0,"","","","-5.76",""
10/08/2026,JT ULTRA_3711,X64063711,DIVIDEND RECEIVED TAIWAN SEMICONDUCTOR (TSM) (Margin),TSM,TAIWAN,Margin,"",0,"","","",27.41,""
10/08/2026,JT ULTRA_3711,X64063711,REDEMPTION FROM CORE ACCOUNT FIDELITY MMKT PREMIUM CLASS (FZDXX) (Cash),FZDXX,FIDELITY MMKT,Cash,1,"-23697.25","","","",23697.25,""
10/07/2026,JT ULTRA_3711,X64063711,YOU BOUGHT VEEVA SYSTEMS INC (VEEV) (Margin),VEEV,VEEVA SYSTEMS INC,Margin,283.25,100,"","","","-28325",10/08/2026
10/05/2026,LISA ROTH_6053,164416053,DIVIDEND RECEIVED J P MORGAN (JEPQ) (Cash),JEPQ,J P MORGAN,Cash,"",0,"","","",126.5,""
10/02/2026,JT ULTRA_3711,X64063711,Electronic Funds Transfer Paid (Cash),"",No Description,Cash,"",0,"","","","-15000",""

"The data and information in this spreadsheet is provided to you solely for your use"
Date downloaded 10/10/2026 10:07 am
'''


def sheet(fzdxx=648.04, fzfxx=0.0, qacds=20408.02, vmfxx=27966.40):
    """(formula-view sheet, value-view sheet) of a small Current Positions tab; one stock per account."""
    out = []
    for values in (False, True):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = accounts.SHEET
        rows = [(FID, "FZDXX", fzdxx, 1), (FID, "FZFXX", fzfxx, 1), (FID, "VEEV", 100, 280.0),
                (CHA, "VMFXX", vmfxx, 1), (CHA, "VFLO", 200, 50.0), (CHA, "QACDS", qacds, 1)]
        for i, (acct, tk, qty, px) in enumerate(rows, start=3):
            ws.cell(i, 1, acct)
            ws.cell(i, 2, tk)
            ws.cell(i, 5, qty)
            ws.cell(i, 6, px)
            ws.cell(i, 7, qty * px if values else f"=E{i}*F{i}")
            ws.cell(i, 8, qty)
            ws.cell(i, 13, "old note")
        out.append(ws)
    return out


def state(**pending):
    s = {"basis": "test", "accounts": {k: {"booked_through": "2026-10-08", "booked_keys": [], "booked_pending": 0.0}
                                       for k in accounts.ACCOUNTS},
         "flows": [{"date": "2026-10-02", "account": "fidelity", "amount": -15000.0, "status": "posted", "note": "EFT"}]}
    for acct, amt in pending.items():
        accounts.add_pending(s, acct, amt, D(2026, 10, 9), "announced")
    return s


SNAP = {"as_of": "2026-10-10", "source": "site read",
        "fidelity": {"total": 28713.12, "cash": {"FZDXX": 669.69, "FZFXX": 0}, "unswept": 43.43},
        "chase": {"total": 38393.66, "cash": {"QACDS": 427.26, "VMFXX": 27966.40}}}


def txns(tmp_path):
    c, f = tmp_path / "c.csv", tmp_path / "f.csv"
    c.write_text(CHASE_CSV)
    f.write_text(FIDELITY_CSV)
    return accounts.parse_chase_csv(str(c)), accounts.parse_fidelity_csv(str(f))


def test_chase_export_is_classified(tmp_path):
    chase, _ = txns(tmp_path)
    assert [t.kind for t in chase] == ["external", "internal", "income", "internal", "trade", "other"]
    assert chase[0].amount == -20000 and chase[0].date == D(2026, 10, 9)


def test_fidelity_export_reads_only_the_system_account(tmp_path):
    _, fid = txns(tmp_path)
    assert all(t.account == "fidelity" for t in fid) and len(fid) == 7          # the other household account is skipped
    assert [t.kind for t in fid] == ["internal", "income", "income", "income", "internal", "trade", "external"]


def test_cash_walk_closes_and_targets_include_unswept_cash(tmp_path):
    chase, fid = txns(tmp_path)
    live = [t for t in chase if t.date >= D(2026, 10, 8)] + fid
    wsf, wsv = sheet()
    fidelity, chase_res = accounts.reconcile(wsf, wsv, state(), SNAP, live)
    assert fidelity.unexplained == 0 and chase_res.unexplained == 0
    assert fidelity.target_cash == {"FZDXX": 713.12, "FZFXX": 0.0}             # unswept cash sits with the fund in use
    assert chase_res.target_cash == {"QACDS": 427.26, "VMFXX": 27966.40}
    assert [t.amount for t in chase_res.posted_flows] == [-20000.0]
    assert fidelity.projected_total == pytest.approx(28713.12) and not fidelity.attention and not chase_res.attention


def test_an_unexplained_difference_is_reported_and_blocks_the_write(tmp_path):
    chase, fid = txns(tmp_path)
    wsf, wsv = sheet(qacds=20000.00)
    res = accounts.reconcile(wsf, wsv, state(), SNAP, [t for t in chase if t.date >= D(2026, 10, 8)] + fid)[1]
    assert res.unexplained == pytest.approx(408.02) and not res.writable
    assert "cash walk is off by $408.02" in res.attention[0]


def test_trades_and_unrecognized_rows_need_attention(tmp_path):
    chase, _ = txns(tmp_path)
    s = state()
    s["accounts"]["chase"]["booked_through"] = "2026-10-01"
    wsf, wsv = sheet()
    res = accounts.reconcile(wsf, wsv, s, {"as_of": "2026-10-10", "chase": SNAP["chase"]}, chase)[0]
    text = " | ".join(res.attention)
    assert "JEPQ" in text and "NOT booked" in text and "unrecognized transaction 'SOMETHING NEW'" in text


def test_pending_transfer_is_carried_then_retired_when_it_posts():
    wsf, wsv = sheet(fzdxx=713.12)
    s = state(fidelity=15000.0)
    snap = {"as_of": "2026-10-10", "fidelity": {"total": 28713.12, "cash": {"FZDXX": 713.12}}}
    res = accounts.reconcile(wsf, wsv, s, snap, [])[0]
    assert res.pending == 15000.0 and res.target_cash["FZDXX"] == 15713.12 and res.unexplained == 0
    assert accounts.net_external_withdrawals(s) == 0.0                         # -(-15,000 + 15,000 pending)

    # booked with the pending inside; three days later it posts into the core fund
    wsf, wsv = sheet(fzdxx=15713.12)
    s["accounts"]["fidelity"]["booked_pending"] = 15000.0
    posted = Txn(D(2026, 10, 13), "fidelity", "external", "", 0, 15000.0, "Electronic Funds Transfer Received")
    snap = {"as_of": "2026-10-13", "fidelity": {"cash": {"FZDXX": 713.12, "FZFXX": 15000.0}}}
    res = accounts.reconcile(wsf, wsv, s, snap, [posted])[0]
    assert res.unexplained == 0 and res.pending == 0 and len(res.cleared_pending) == 1
    assert res.target_cash == {"FZDXX": 713.12, "FZFXX": 15000.0}
    assert any("FZFXX goes from zero" in c and "Performance Tracking lot" in c for c in res.crossings)


def test_identical_transactions_on_one_day_are_each_counted_once():
    a = Txn(D(2026, 10, 10), "chase", "income", "VFLO", 0, 5.0, "DIV")
    s = state()
    s["accounts"]["chase"]["booked_through"] = "2026-10-10"
    s["accounts"]["chase"]["booked_keys"] = [a.key()]
    assert accounts._unbooked(s, "chase", [a]) == []
    assert accounts._unbooked(s, "chase", [a, a]) == [a]


def test_apply_writes_cash_rows_notes_and_the_ledger(tmp_path):
    chase, fid = txns(tmp_path)
    live = [t for t in chase if t.date >= D(2026, 10, 8)] + fid
    wsf, _ = sheet()
    path, sp = str(tmp_path / "wb.xlsx"), str(tmp_path / "state.json")
    wsf.parent.save(path)
    s = state(fidelity=15000.0)
    accounts.save_state(s, sp)
    results, new = accounts.apply(path, SNAP, live, sp, recalc=False)
    ws = openpyxl.load_workbook(path)[accounts.SHEET]
    by = {ws.cell(r, 2).value: r for r in range(3, 9)}
    assert ws.cell(by["FZDXX"], 5).value == 15713.12 == ws.cell(by["FZDXX"], 8).value
    assert ws.cell(by["QACDS"], 5).value == 427.26
    note = ws.cell(by["FZDXX"], 13).value
    assert "estimate (site read, 10/10/26)" in note and "PENDING transfer of $15,000.00" in note and "$43.43" in note
    assert ws.cell(by["VMFXX"], 13).value == "old note"                        # an unchanged row keeps its note
    assert ws.cell(by["VEEV"], 5).value == 100
    assert accounts.net_external_withdrawals(new) == 20000.0                   # 15,000 out + 20,000 out - 15,000 pending in
    assert new["accounts"]["fidelity"]["booked_pending"] == 15000.0
    assert json.load(open(sp)) == new
    # a second run over the same inputs finds nothing new and changes nothing
    again, after = accounts.apply(path, SNAP, live, sp, recalc=False)
    assert all(r.unexplained == 0 and not [t for t in r.new_txns if t.kind != "internal"] for r in again)
    assert accounts.net_external_withdrawals(after) == 20000.0
    report = accounts.format_report(again, after, SNAP, applied=True, portfolio_total=100.0)
    assert "Drawdown Log value = portfolio total 100.00 + 20,000.00 = 20,100.00" in report


def test_open_orders_are_listed_kept_and_never_booked(tmp_path):
    wsf, _ = sheet(fzdxx=713.12, qacds=427.26)
    path, sp = str(tmp_path / "wb.xlsx"), str(tmp_path / "state.json")
    wsf.parent.save(path)
    accounts.save_state(state(), sp)
    snap = {"as_of": "2026-10-10", "transactions": [],
            "fidelity": {"cash": {"FZDXX": 713.12}, "open_orders": ["10/8 Sell 40 CRWD limit $285.60 (GTC)"]},
            "chase": {"cash": {"QACDS": 427.26, "VMFXX": 27966.40}, "open_orders": []}}
    results, new = accounts.apply(path, snap, [], sp, recalc=False)
    report = accounts.format_report(results, new, snap, applied=True)
    assert "open orders (1, not filled, nothing booked):" in report and "Sell 40 CRWD" in report
    assert "open orders: none" in report and not any(r.attention for r in results)
    assert openpyxl.load_workbook(path)[accounts.SHEET].cell(5, 5).value == 100       # VEEV untouched
    listing = accounts.format_orders(new)
    assert "Sell 40 CRWD" in listing and "Chase Self-Directed ...3693 (2026-10-10): none" in listing
    del snap["fidelity"]["open_orders"]
    assert "open orders: not read" in accounts.format_report(
        accounts.apply(path, snap, [], sp, recalc=False)[0], new, snap, applied=True)


def test_a_fund_the_account_does_not_have_is_refused():
    wsf, wsv = sheet()
    with pytest.raises(accounts.AccountsError, match="SPAXX"):
        accounts.reconcile(wsf, wsv, state(), {"as_of": "2026-10-10", "fidelity": {"cash": {"SPAXX": 1}}}, [])
