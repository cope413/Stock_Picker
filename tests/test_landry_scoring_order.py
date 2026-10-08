"""audit.check_scoring_row_order: sorting the Scoring tab silently re-labels every tab that keys to its rows (found 10/8/26 in a review copy)."""

import openpyxl

from landry.audit import check_scoring_row_order


def book(tmp_path, scoring, ir):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sc = wb.create_sheet("Scoring")
    sc["B2"] = "Company"
    for i, t in enumerate(scoring):
        sc.cell(row=3 + i, column=1, value=t)
    ic = wb.create_sheet("Implied-Return Calculator")
    ic["A3"] = "Ticker"
    for row, t in ir.items():
        ic.cell(row=row, column=1, value=t)
    ic["A40"] = "=Scoring!A39"                                                        # a formula is not an anchor
    p = str(tmp_path / "b.xlsx")
    wb.save(p)
    return p


def test_matching_order_passes(tmp_path):
    p = book(tmp_path, ["MU", "AVGO", "PLTR", "NVDA"], {5: "AVGO", 6: "PLTR"})        # Implied-Return row = Scoring row + 1
    c = check_scoring_row_order(p)[0]
    assert c.ok and "2 anchors" in c.detail


def test_a_sorted_scoring_tab_fails(tmp_path):
    p = book(tmp_path, ["NVDA", "PLTR", "AVGO", "MU"], {5: "AVGO", 6: "PLTR"})        # same tickers, new order
    c = check_scoring_row_order(p)[0]
    assert not c.ok and "Implied-Return row 5 says AVGO but Scoring row 4 is PLTR" in c.detail and "never sort Scoring" in c.fix


def test_the_live_workbook_is_in_order():
    from landry.xlsx_io import latest_workbook
    import os
    wb = latest_workbook(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    if wb:
        assert check_scoring_row_order(wb)[0].ok
