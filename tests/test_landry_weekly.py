"""Offline tests for landry.weekly and the `weekly` / `market` CLI commands -- synthetic workbooks, fake
providers, fake recalc and audit (no network, no LibreOffice)."""

import datetime as dt
from types import SimpleNamespace

import openpyxl
import pytest
from test_landry_market import GOOD, add_market_tabs, digest, snap, table
from test_landry_prices import NOW_AFTER_FRIDAY_10_2, build_workbook, fake_fetch

from landry import cli, ledger, prices, weekly
from landry.doctor import Check

TICKERS = ("T01", "T02", "T03")
SAT = NOW_AFTER_FRIDAY_10_2                                      # Saturday 10/3: the 10/2 week has closed
WED = dt.datetime(2026, 10, 7, 9, 0)
NOSLEEP = lambda s: None                                         # noqa: E731


def combined(tmp_path, name="w.xlsx"):
    """Price History + Returns (Calc) + Correlation Matrix from the prices fixture, plus the two tabs
    the market refresh writes."""
    path = build_workbook(tmp_path, name=name)
    wb = openpyxl.load_workbook(path)
    add_market_tabs(wb, tickers=TICKERS)
    wb.save(path)
    return path


def run(path, **kw):
    kw.setdefault("now", SAT)
    kw.setdefault("fetch_weeks", fake_fetch())
    kw.setdefault("snapshot", table(dict(zip(TICKERS, GOOD.values()))))
    kw.setdefault("earnings", lambda t: None)
    kw.setdefault("sleep", NOSLEEP)
    return weekly.run(path, **kw)


def checks(*failing):
    return [Check("fine", True, "ok")] + [Check(n, False, f"{n} broke") for n in failing]


OK_RECALC = lambda p: {"status": "success", "total_errors": 0, "clamped_rows": []}      # noqa: E731


# ------------------------------------------------------------------------ run --

def test_run_appends_the_friday_then_refreshes_market_data(tmp_path):
    path = combined(tmp_path)
    rep = run(path)
    assert rep["problems"] == [] and rep["wrote"] and not rep["dry_run"]
    assert rep["prices"]["appended"] == ["2026-10-02"]
    wb = openpyxl.load_workbook(path)
    assert wb["Price History"]["A9"].value == dt.datetime(2026, 10, 2)
    assert wb["Market Data"]["C3"].value == 101.5 and wb["Market Data"]["E3"].value == 6000.0


def test_quotes_are_cross_checked_against_the_friday_close_just_appended(tmp_path):
    path = combined(tmp_path)
    rep = run(path)                                              # fixture closes are ~109-115; the fake quote is 101.5
    assert any("Price History close" in w for w in rep["warnings"])


def test_the_cross_check_only_applies_on_a_weekend_with_a_current_price_history(tmp_path):
    path = combined(tmp_path)
    assert weekly._weekend_reference(path, SAT) is None          # Price History still ends 9/25: behind the 10/2 close
    prices.append_weeks(path, fetch=fake_fetch(), now=SAT)
    ref = weekly._weekend_reference(path, SAT)
    assert ref and set(ref) >= {"T01", "T21"} and ref["T01"] == prices.latest_closes(path)[1]["T01"]
    assert weekly._weekend_reference(path, WED) is None          # midweek the quote has moved on


def test_a_price_history_failure_does_not_stop_the_market_refresh(tmp_path):
    path = combined(tmp_path)
    rep = run(path, fetch_weeks=fake_fetch(jump=("T01", 1.6)))   # +60% in a week: refused unless forced
    assert rep["prices"] is None and rep["problems"][0].startswith("Price History:")
    assert rep["market"] is not None and rep["wrote"]            # Market Data was still written
    assert openpyxl.load_workbook(path)["Price History"]["A9"].value is None   # nothing appended


def test_no_quote_for_any_ticker_is_a_problem(tmp_path):
    path = combined(tmp_path)
    rep = run(path, snapshot=lambda t: None, do_prices=False)
    assert any("no quote came back for any of 3 tickers" in p for p in rep["problems"])
    assert not rep["wrote"]


def test_a_few_missing_quotes_are_a_warning_not_a_problem(tmp_path):
    path = combined(tmp_path)
    rows = dict(zip(TICKERS, GOOD.values()))
    rows["T02"] = None
    rep = run(path, snapshot=table(rows), do_prices=False)
    assert rep["problems"] == [] and any("no quote for T02" in w for w in rep["warnings"])
    assert rep["wrote"]                                          # the other two rows were refreshed


def test_dry_run_writes_nothing_but_reports_what_it_would_do(tmp_path):
    path = combined(tmp_path)
    before = digest(path)
    rep = run(path, write=False)
    assert digest(path) == before and rep["dry_run"] and not rep["wrote"]
    assert rep["prices"]["appended"] == ["2026-10-02"]
    assert rep["market"]["market"]["cells_changed"] > 0


def test_a_layout_problem_in_market_data_is_reported_not_raised(tmp_path):
    path = combined(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb["Market Data"]["C2"] = "Company (long)"
    wb.save(path)
    rep = run(path)
    assert any(p.startswith("Market Data:") for p in rep["problems"])
    assert rep["prices"]["appended"] == ["2026-10-02"]           # the prices step still ran


def test_market_only_run_leaves_price_history_alone(tmp_path):
    path = combined(tmp_path)
    rep = run(path, do_prices=False)
    assert rep["prices"] is None and rep["wrote"]
    assert openpyxl.load_workbook(path)["Price History"]["A9"].value is None


# --------------------------------------------------------------------- verify --

def test_verify_recalcs_once_when_something_was_written():
    calls = []

    def recalc(p):
        calls.append(p)
        return OK_RECALC(p)
    rep = {"wrote": True, "problems": [], "warnings": []}
    weekly.verify("wb.xlsx", rep, recalc_fn=recalc, audit_fn=lambda p, r: checks())
    assert calls == ["wb.xlsx"] and rep["recalc"]["status"] == "success"
    assert rep["audit"] == {"checks": 1, "failed": [], "standing": []} and rep["problems"] == []


def test_verify_does_not_recalc_when_nothing_was_written():
    rep = {"wrote": False, "problems": [], "warnings": []}
    weekly.verify("wb.xlsx", rep, recalc_fn=lambda p: pytest.fail("recalc ran"), audit_fn=lambda p, r: checks())
    assert "recalc" not in rep and rep["audit"]["checks"] == 1   # the audit still runs


def test_verify_turns_recalc_and_audit_trouble_into_problems():
    rep = {"wrote": True, "problems": [], "warnings": []}
    weekly.verify("wb.xlsx", rep, recalc_fn=lambda p: {"error": "soffice not found"},
                  audit_fn=lambda p, r: checks("table_ref"))
    assert any("recalc failed: soffice not found" in p for p in rep["problems"])
    assert any("audit: 1 of 2 checks failed (table_ref)" in p for p in rep["problems"])

    rep = {"wrote": True, "problems": [], "warnings": []}
    weekly.verify("wb.xlsx", rep, recalc_fn=lambda p: {"status": "errors_found", "total_errors": 3,
                                                       "error_summary": {"#REF!": ["X1"]}, "clamped_rows": [5]},
                  audit_fn=lambda p, r: checks())
    assert any("3 formula error(s)" in p for p in rep["problems"])
    assert any("capped 1 row height" in p for p in rep["problems"])


def test_standing_scoring_findings_are_reported_but_do_not_fail_the_run():
    rep = {"wrote": False, "problems": [], "warnings": []}
    weekly.verify("wb.xlsx", rep, audit_fn=lambda p, r: checks("scoring_verification:no_audit_trail",
                                                               "scoring_verification:V/bulk_import"))
    assert rep["problems"] == [] and rep["audit"]["failed"] == [] and len(rep["audit"]["standing"]) == 2
    text = weekly.format_report(dict(rep, prices=None, market=None, dry_run=False),
                                when=dt.datetime(2026, 10, 3, 9, 0))
    assert "2 standing Scoring-provenance finding(s)" in text and "NEEDS ATTENTION" not in text
    rep = {"wrote": False, "problems": [], "warnings": []}                      # a structural failure still gates
    weekly.verify("wb.xlsx", rep, audit_fn=lambda p, r: checks("scoring_verification:x", "table_ref"))
    assert len(rep["problems"]) == 1 and "table_ref" in rep["problems"][0]


def test_verify_survives_an_audit_that_cannot_run():
    def boom(p, r):
        raise RuntimeError("no git")
    rep = {"wrote": False, "problems": [], "warnings": []}
    weekly.verify("wb.xlsx", rep, audit_fn=boom)
    assert rep["audit"] is None and any("audit could not run: no git" in p for p in rep["problems"])


def test_verify_warns_when_the_recalc_was_skipped():
    rep = {"wrote": True, "problems": [], "warnings": []}
    weekly.verify("wb.xlsx", rep, audit_fn=lambda p, r: checks(), recalc_now=False)
    assert any("NOT safe to commit" in w for w in rep["warnings"]) and "recalc" not in rep


def test_rule38_status_reads_the_matrix_tally(tmp_path):
    wb = openpyxl.Workbook()
    cm = wb.active
    cm.title = "Correlation Matrix"
    cm["J2"], cm["J3"] = 11, "REVIEW REQUIRED"
    path = str(tmp_path / "cm.xlsx")
    wb.save(path)
    assert weekly.rule38_status(path) == {"over_cap": 11, "status": "REVIEW REQUIRED"}
    assert weekly.rule38_status(str(tmp_path / "missing.xlsx")) is None


# --------------------------------------------------------------------- report --

def _report(tmp_path, **kw):
    path = combined(tmp_path)
    rep = run(path, earnings=table({"T01": dt.date(2026, 10, 9), "T02": dt.date(2026, 12, 1)}), **kw)
    weekly.verify(path, rep, recalc_fn=OK_RECALC, audit_fn=lambda p, r: checks())
    return rep


def test_report_says_what_happened(tmp_path):
    text = weekly.format_report(_report(tmp_path), workbook="w.xlsx", when=dt.datetime(2026, 10, 3, 9, 0))
    for needle in ("Weekly routine", "Price History   appended 2026-10-02", "Market Data     3 of 3 tickers refreshed",
                   "Earnings dates  3 tickers: 2 updated", "T01 2026-09-01 -> 2026-10-09",
                   "held names reporting within 14 days: T01 10/09",
                   "Recalc          success, 0 formula errors", "Audit           1 checks, 0 failed",
                   "Nothing is committed. Review the diff"):
        assert needle in text, needle
    assert "NEEDS ATTENTION" not in text and "DRY RUN" not in text


def test_report_for_an_up_to_date_week_and_a_dry_run(tmp_path):
    path = combined(tmp_path)
    prices.append_weeks(path, fetch=fake_fetch(), now=SAT)
    rep = run(path, write=False)
    text = weekly.format_report(rep, when=dt.datetime(2026, 10, 3, 9, 0))
    assert "up to date through 2026-10-02" in text and "DRY RUN" in text and "Nothing was changed." in text
    assert "to update" in text


def test_report_leads_with_what_needs_attention(tmp_path):
    path = combined(tmp_path)
    rep = run(path, fetch_weeks=fake_fetch(jump=("T01", 1.6)))
    weekly.verify(path, rep, recalc_fn=OK_RECALC, audit_fn=lambda p, r: checks("table_ref"))
    text = weekly.format_report(rep, when=dt.datetime(2026, 10, 3, 9, 0))
    assert "Price History   FAILED (see below)" in text and "NEEDS ATTENTION" in text
    assert "[FAIL] table_ref: table_ref broke" in text


# ------------------------------------------------------------------------ CLI --

def _args(cmd="weekly", **kw):
    d = dict(cmd=cmd, workbook="wb.xlsx", dry_run=False, no_recalc=False, force=False, allow_big_moves=False)
    d.update(kw)
    return SimpleNamespace(**d)


def test_cli_refuses_while_excel_has_the_workbook_open(monkeypatch, capsys):
    monkeypatch.setattr(ledger, "_excel_has_open", lambda p: True)
    monkeypatch.setattr(weekly, "run", lambda *a, **k: pytest.fail("run was called"))
    assert cli._cmd_weekly(_args()) == 1
    assert "Excel appears to have the workbook open" in capsys.readouterr().err
    # a dry run reads only, so it is allowed through (and --force overrides the guard)
    monkeypatch.setattr(weekly, "run", lambda *a, **k: {"problems": [], "warnings": [], "wrote": False,
                                                       "prices": None, "market": None, "dry_run": True})
    monkeypatch.setattr(weekly, "verify", lambda *a, **k: None)
    assert cli._cmd_weekly(_args(dry_run=True)) == 0
    assert cli._cmd_weekly(_args(force=True)) == 0


def test_cli_exit_status_and_flags_reach_the_routine(monkeypatch, capsys):
    monkeypatch.setattr(ledger, "_excel_has_open", lambda p: False)
    seen = {}

    def fake_run(wb, **kw):
        seen["run"] = kw
        return {"problems": seen.get("problems", []), "warnings": [], "wrote": False, "prices": None,
                "market": None, "dry_run": kw["write"] is False}

    def fake_verify(wb, rep, **kw):
        seen["verify"] = kw

    monkeypatch.setattr(weekly, "run", fake_run)
    monkeypatch.setattr(weekly, "verify", fake_verify)
    assert cli._cmd_weekly(_args(allow_big_moves=True, no_recalc=True)) == 0
    assert seen["run"]["do_prices"] is True and seen["run"]["allow_big_moves"] is True
    assert seen["verify"]["recalc_now"] is False
    assert cli._cmd_weekly(_args(cmd="market")) == 0
    assert seen["run"]["do_prices"] is False                     # `market` is the routine minus Price History
    seen["problems"] = ["Market Data: no quote came back for any of 36 tickers"]
    assert cli._cmd_weekly(_args()) == 2
    assert "NEEDS ATTENTION" in capsys.readouterr().out


def test_cli_parser_accepts_the_new_commands(monkeypatch):
    monkeypatch.setattr(cli, "_cmd_weekly", lambda a: (a.cmd, a.dry_run, getattr(a, "allow_big_moves", None)))
    assert cli.main(["weekly", "--dry-run", "--workbook", "x.xlsx"]) == ("weekly", True, False)
    assert cli.main(["market", "--no-recalc"]) == ("market", False, None)
