"""landry.cp_tab: adding a position row to Current Positions without openpyxl's insert_rows (the 9/30/26 insertion left a merge
inside the Table, stale heights and formulas pointing at the old rows; Excel asked to repair the file)."""

import os
import shutil

import pytest

openpyxl = pytest.importorskip("openpyxl")

from openpyxl.worksheet.table import Table, TableStyleInfo      # noqa: E402

from landry import cp_tab                                        # noqa: E402
from landry.xlsx_recalc import recalc, soffice_path              # noqa: E402

JT, CH = "JT ULTRA (Fidelity)", "Self-Directed (Chase)"
LIVE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "LANDRY_SYSTEM_WORKBOOK_25.xlsx")


# --------------------------------------------------------------------------- pure formula rewriting

@pytest.mark.parametrize("before,after", [
    ("=SUM(G3:G23)", "=SUM(G3:G24)"),                                           # a block sum ending above the new row grows
    ("=SUM(G3:G23,G27:G46)", "=SUM(G3:G24,G28:G47)"),                           # the next block moves down whole
    ("=IFERROR(G27/SUM(G27:G46),\"\")", "=IFERROR(G28/SUM(G28:G47),\"\")"),
    ("=SUM(G14:G21)+G27+G28+G46", "=SUM(G14:G21)+G28+G29+G47"),                # a list of rows names each one
    ("=SUM(G14:G21)", "=SUM(G14:G21)"),                                         # nothing to do above the insertion point
    ("=SUM($G$3:$G$23)", "=SUM($G$3:$G$24)"),
    ("=$B27", "=$B28"),
    ("=SUM(G1:G23)", "=SUM(G1:G23)"),                                           # a range that starts above the block is not a block sum
    ("=I24/(H24-H14-H13)", "=I25/(H25-H14-H13)"),                               # the row-24 subtotal itself moves
    ("=SUBTOTAL(103,B3:B23)", "=SUBTOTAL(103,B3:B24)"),
])
def test_shift_formula(before, after):
    assert cp_tab.shift_formula(before, 24) == after


def test_other_sheets_and_structured_references_are_left_alone():
    f = ("=IF(IFERROR(INDEX('Monitor & Recheck Triggers'!$H$3:$H$51,MATCH($B30,'Monitor & Recheck Triggers'!$A$3:$A$51,0)),\"\")=\"\","
         "61.04,INDEX('Monitor & Recheck Triggers'!$H$3:$H$51,MATCH($B30,'Monitor & Recheck Triggers'!$A$3:$A$51,0)))")
    assert cp_tab.shift_formula(f, 24) == f.replace("$B30", "$B31")
    g = "=CurrentPositionsTable[[#This Row],[Unrealized G/L ($)]]/CurrentPositionsTable[[#This Row],[Cost Basis\n($)]]"
    assert cp_tab.shift_formula(g, 24) == g
    assert cp_tab.shift_formula("text", 24) == "text" and cp_tab.shift_formula(None, 24) is None and cp_tab.shift_formula(5, 24) == 5


def test_retarget_moves_a_rows_own_cells_and_nothing_else():
    assert cp_tab.retarget("=E23*F23", 23, 24) == "=E24*F24"
    assert cp_tab.retarget("=IFERROR(G23/SUM(G3:G24),\"\")", 23, 24) == "=IFERROR(G24/SUM(G3:G24),\"\")"
    f = "=IF(IFERROR(INDEX('Monitor & Recheck Triggers'!$H$3:$H$51,MATCH($B23,'Monitor & Recheck Triggers'!$A$3:$A$51,0)),\"\")=\"\",67.06,1)"
    assert cp_tab.retarget(f, 23, 24) == f.replace("$B23", "$B24")


# --------------------------------------------------------------------------- a small sheet with the real structure

PRICE = ("=IF(IFERROR(INDEX('Monitor & Recheck Triggers'!$H$3:$H$51,MATCH($B{r},'Monitor & Recheck Triggers'!$A$3:$A$51,0)),\"\")=\"\","
         "{fb},INDEX('Monitor & Recheck Triggers'!$H$3:$H$51,MATCH($B{r},'Monitor & Recheck Triggers'!$A$3:$A$51,0)))")


def _small_workbook():
    """JT block rows 3-5 (two stocks + a cash fund) with its subtotal on 6 and Excluding-Cash line on 7; Chase block rows 9-10,
    subtotal 11, Excluding-Cash 12; combined total 14; Cash / Cash Equivalents 15; a merged note on 17."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = cp_tab.SHEET
    rows = {3: (JT, "AAA", 10, 100.0, 900.0), 4: (JT, "BBB", 5, 50.0, 200.0), 5: (JT, "CASH", 1000, 1.0, 1000.0),
            9: (CH, "CCC", 20, 10.0, 150.0), 10: (CH, "CSH2", 500, 1.0, 500.0)}
    for c, h in enumerate(["Account", "Ticker", "Description", "Asset Class", "Quantity", "Price", "Market Value", "Cost Basis", "Unrealized", "Unrealized %",
                           "% of Account", "% of Combined", "Notes"], 1):
        ws.cell(2, c).value = h
    for r, (acct, tk, q, px, cost) in rows.items():
        sub = (3, 5) if acct == JT else (9, 10)
        other = (9, 10) if acct == JT else (3, 5)
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value = acct, tk, tk + " name"
        ws.cell(r, 4).value = "Cash" if tk.startswith("C") and tk in ("CASH", "CSH2") else "Equity"
        ws.cell(r, 5).value = q
        ws.cell(r, 6).value = PRICE.format(r=r, fb=px)
        ws.cell(r, 7).value = f"=E{r}*F{r}"
        ws.cell(r, 8).value = cost
        ws.cell(r, 9).value = f"=G{r}-H{r}"
        ws.cell(r, 10).value = f'=IF(H{r}=0,"",G{r}/H{r}-1)'
        ws.cell(r, 11).value = f'=IFERROR(G{r}/SUM(G{sub[0]}:G{sub[1]}),"")'
        ws.cell(r, 12).value = f'=IFERROR(G{r}/(SUM(G3:G5,G9:G10)),"")'
        ws.row_dimensions[r].height = 15.0
    for sr, (a, b), ex_cash in ((6, (3, 5), (5,)), (11, (9, 10), (10,))):
        ws.cell(sr, 1).value = "Positions"
        ws.cell(sr, 2).value = f"=SUBTOTAL(103,B{a}:B{b})"
        ws.cell(sr, 3).value = f"{JT if sr == 6 else CH} SUBTOTAL"
        for col, L in ((7, "G"), (8, "H"), (9, "I"), (11, "K"), (12, "L")):
            ws.cell(sr, col).value = f"=SUM({L}{a}:{L}{b})"
        ws.cell(sr, 10).value = f"=I{sr}/H{sr}"
        ws.cell(sr + 1, 9).value = "Excluding Cash"
        ws.cell(sr + 1, 10).value = f"=I{sr}/(H{sr}" + "".join(f"-H{x}" for x in ex_cash) + ")"
    ws.cell(14, 3).value = "COMBINED PORTFOLIO TOTAL"
    ws.cell(14, 7).value, ws.cell(14, 8).value = "=SUM(G3:G5)+SUM(G9:G10)", "=SUM(H3:H5)+SUM(H9:H10)"
    ws.cell(15, 3).value = "Cash / Cash Equivalents"
    ws.cell(15, 7).value, ws.cell(15, 8).value = "=G5+G10", "=H5+H10"
    ws.cell(17, 1).value = "NOTE: a long note"
    ws.merge_cells("A17:M17")
    ws.row_dimensions[17].height = 142.5
    t = Table(displayName=cp_tab.TABLE, ref="A2:M15")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    return wb, ws


def test_block_bounds_finds_each_accounts_rows():
    wb, ws = _small_workbook()
    assert cp_tab.block_bounds(ws, JT) == (3, 5, 6)
    assert cp_tab.block_bounds(ws, CH) == (9, 10, 11)
    with pytest.raises(cp_tab.CurrentPositionsError):
        cp_tab.block_bounds(ws, "Nobody")


def test_insert_position_moves_everything_below_and_keeps_every_reference_true():
    wb, ws = _small_workbook()
    new = cp_tab.insert_position(wb, JT, "NEW", "New Corp", "Equity", 7, 700.0, 123.45, note="added")
    assert new == 6
    # the new row: typed values, its own formulas, the Monitor lookup with the new fallback price
    assert [ws.cell(6, c).value for c in (1, 2, 3, 4, 5, 8, 13)] == [JT, "NEW", "New Corp", "Equity", 7, 700.0, "added"]
    assert ws.cell(6, 7).value == "=E6*F6" and ws.cell(6, 9).value == "=G6-H6"
    assert ws.cell(6, 11).value == '=IFERROR(G6/SUM(G3:G6),"")' and ws.cell(6, 12).value == '=IFERROR(G6/(SUM(G3:G6,G10:G11)),"")'
    assert '="",123.45,INDEX' in ws.cell(6, 6).value and "$B6" in ws.cell(6, 6).value and "$B5" not in ws.cell(6, 6).value
    # the JT block sums grew; the old rows are untouched
    assert ws.cell(7, 2).value == "=SUBTOTAL(103,B3:B6)" and ws.cell(7, 7).value == "=SUM(G3:G6)" and ws.cell(7, 12).value == "=SUM(L3:L6)"
    assert ws.cell(5, 11).value == '=IFERROR(G5/SUM(G3:G6),"")'                      # an old JT row's % of account now includes the new row
    assert ws.cell(8, 10).value == "=I7/(H7-H5)"                                      # the Excluding-Cash line followed its subtotal
    # the Chase block moved down whole, its own formulas followed
    assert [ws.cell(10, c).value for c in (1, 2)] == [CH, "CCC"] and ws.cell(10, 7).value == "=E10*F10"
    assert ws.cell(10, 11).value == '=IFERROR(G10/SUM(G10:G11),"")' and ws.cell(10, 12).value == '=IFERROR(G10/(SUM(G3:G6,G10:G11)),"")'
    assert ws.cell(12, 7).value == "=SUM(G10:G11)" and ws.cell(13, 10).value == "=I12/(H12-H11)"
    assert ws.cell(15, 7).value == "=SUM(G3:G6)+SUM(G10:G11)" and ws.cell(16, 7).value == "=G5+G11" and ws.cell(16, 8).value == "=H5+H11"
    # structure: Table, merge, heights
    assert ws.tables[cp_tab.TABLE].ref == "A2:M16"
    assert [str(m) for m in ws.merged_cells.ranges] == ["A18:M18"]
    assert ws.row_dimensions[18].height == 142.5 and ws.row_dimensions[17].height != 142.5 and ws.row_dimensions[6].height == 15.0


def test_the_excluding_cash_line_is_rederived_by_ticker():
    wb, ws = _small_workbook()
    cp_tab.insert_position(wb, JT, "NEW", "New Corp", "Equity", 7, 700.0, 123.45)
    ws.cell(8, 10).value = "=I7/(H7-H5-H4)"                                           # a line gone wrong, as JT ULTRA's was after a row was added above it
    out = cp_tab.fix_excluding_cash(ws, JT, ("CASH",))
    assert out == {"was": "=I7/(H7-H5-H4)", "now": "=I7/(H7-H5)"}
    with pytest.raises(cp_tab.CurrentPositionsError):
        cp_tab.fix_excluding_cash(ws, JT, ("MISSING",))


def test_it_refuses_instead_of_guessing():
    wb, ws = _small_workbook()
    cp_tab.insert_position(wb, JT, "NEW", "New Corp", "Equity", 7, 700.0, 123.45)
    with pytest.raises(cp_tab.CurrentPositionsError, match="already has a row"):
        cp_tab.insert_position(wb, JT, "NEW", "New Corp", "Equity", 7, 700.0, 123.45)
    wb, ws = _small_workbook()
    ws.cell(10, 6).value = "=E10*2"                                                     # the Chase block's last row is the template; its Price cell is not one this module knows
    with pytest.raises(cp_tab.CurrentPositionsError, match="unexpected shape"):
        cp_tab.insert_position(wb, CH, "NEW", "x", "Equity", 1, 1.0, 1.0)


# --------------------------------------------------------------------------- the live workbook

@pytest.mark.skipif(soffice_path() is None or not os.path.exists(LIVE), reason="needs LibreOffice and the live workbook")
def test_inserting_a_zero_share_row_into_the_live_workbook_changes_no_number(tmp_path):
    """The invariant that matters: a position of zero shares added to the JT ULTRA block must change nothing on the sheet
    except the block's position count -- every subtotal, percentage and combined total reads what it read before, one row
    lower where it moved. Both copies are recalculated by LibreOffice."""
    before, after = str(tmp_path / "before.xlsx"), str(tmp_path / "after.xlsx")
    shutil.copy(LIVE, before)
    wb = openpyxl.load_workbook(LIVE)
    cp_tab.insert_position(wb, JT, "ZZZZ", "Test position", "Equity", 0, 0.0, 1.0)
    wb.save(after)
    assert recalc(before, 180)["total_errors"] == 0 and recalc(after, 180)["total_errors"] == 0
    a = openpyxl.load_workbook(before, data_only=True)[cp_tab.SHEET]
    b = openpyxl.load_workbook(after, data_only=True)[cp_tab.SHEET]
    at = cp_tab.block_bounds(openpyxl.load_workbook(before)[cp_tab.SHEET], JT)[2]
    diffs = []
    for r in range(1, a.max_row + 1):
        r2 = r + 1 if r >= at else r
        for c in range(1, 17):
            va, vb = a.cell(r, c).value, b.cell(r2, c).value
            if isinstance(va, (int, float)) and isinstance(vb, (int, float)):
                if abs(va - vb) > 1e-9 * max(1.0, abs(va)):
                    diffs.append((r, c, va, vb))
            elif va != vb and not (va in (None, "") and vb in (None, "")):
                diffs.append((r, c, va, vb))
    assert diffs == [(at, 2, a.cell(at, 2).value, a.cell(at, 2).value + 1)], diffs     # only the position count moved
