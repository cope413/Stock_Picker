"""landry.stops: the stop-review ladder (Journal row 109) -- pure classification, the registry, a synthetic workbook, the audit check
and the `landry stops` command. No network, no LibreOffice."""

import datetime as dt
import json
import os

import openpyxl
import pytest

from landry import audit, cli, stops
from landry.xlsx_io import latest_workbook

D = dt.date
_WB = latest_workbook(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
needs_workbook = pytest.mark.skipif(not _WB, reason="no workbook file")
FRIDAY = D(2026, 10, 2)


def weekly(prices, end=FRIDAY):
    """Weekly closes, oldest first, the last on ``end``."""
    n = len(prices)
    return [(end - dt.timedelta(weeks=n - 1 - i), float(p)) for i, p in enumerate(prices)]


def flat_then(level, n_flat=60, tail=()):
    return [level] * n_flat + list(tail)


def holding(price_path, cost_per_share, qty=100, price=None, rs_score=4, spy=None, weight=0.03):
    closes = weekly(price_path)
    return stops.Holding("TST", qty, cost_per_share * qty, price if price is not None else closes[-1][1], weight, closes, rs_score, spy)


def review(date, conclusion="reaffirmed", deterioration=False, nxt=None, adds=True):
    return {"date": date.isoformat(), "conclusion": conclusion, "adds_frozen": adds, "deterioration": deterioration,
            "next_check": nxt.isoformat() if nxt else None, "journal_row": 110}


# --------------------------------------------------------------------------- the numbers

def test_the_average_and_weeks_under():
    c = weekly([10] * 40)
    assert stops.ma(c) == 10
    assert stops.ma(weekly([10] * 39)) is None                                    # not enough history
    falling = weekly([10] * 40 + [9, 8, 7])                                       # each of the last three closes is under its own average
    assert stops.weeks_under(falling) == 3
    assert stops.weeks_under(weekly([10] * 45)) == 0


def test_relative_strength_points():
    c = weekly([100] * 30 + [80] * 3)                                             # down 20% over the window
    spy = {d: 500.0 for d, _ in c}
    assert stops.rs_points(c, spy) == pytest.approx(-20.0)
    spy2 = {d: 500.0 + i for i, (d, _) in enumerate(c)}                          # SPY up 26 / 500 = 5.2% over the same 26 weeks
    assert stops.rs_points(c, spy2) == pytest.approx(-20.0 - 5.2, abs=0.1)
    assert stops.rs_points(c, None) is None
    assert stops.rs_points(weekly([100] * 10), spy) is None


# --------------------------------------------------------------------------- rungs

def test_a_price_decline_alone_is_never_a_rung():
    h = holding(flat_then(100, tail=[100]), cost_per_share=140)                   # 28.6% below cost, above its average, RS fine
    s = stops.classify(h, None, FRIDAY)
    assert s.rung == 0 and not s.needs_attention and not s.adds_frozen


@pytest.mark.parametrize("price,rung", [(80.1, 0), (80.0, 1), (75.0, 1), (70.0, 2), (61.0, 2), (60.0, 2)])
def test_the_loss_thresholds_with_the_price_under_its_average(price, rung):
    path = flat_then(120, tail=[110, 100, 90, price])
    h = holding(path, cost_per_share=100, price=price)
    s = stops.classify(h, review(FRIDAY), FRIDAY)
    assert s.under_ma is True
    assert s.rung == rung                                                        # 20% is rung 1, 30% and beyond rung 2 (rung 3 needs 8 weeks under)


def test_either_leg_confirms():
    above_but_weak = holding(flat_then(60, tail=[75, 75]), cost_per_share=100, price=75, rs_score=1)    # 25% down, over its average, RS score 1
    assert stops.classify(above_but_weak, None, FRIDAY).under_ma is False
    assert stops.classify(above_but_weak, None, FRIDAY).rung == 1
    spy = {d: 1000.0 for d, _ in weekly(flat_then(60, tail=[75, 75]))}
    live_ok = holding(flat_then(60, tail=[75, 75]), cost_per_share=100, price=75, rs_score=1, spy=spy)
    # six months ago the stock was 60, now 75: ahead of a flat SPY, so the live number overrides the recorded score
    assert stops.classify(live_ok, None, FRIDAY).rs1 is False
    assert stops.classify(live_ok, None, FRIDAY).rung == 0


def test_approaching_rung_one():
    h = holding(flat_then(120, tail=[100, 90, 82.5]), cost_per_share=100, price=82.5)
    s = stops.classify(h, None, FRIDAY)
    assert s.rung == 0 and s.approaching and "approaching" in s.action


def test_rung_two_trims_only_when_the_re_underwrite_found_deterioration():
    path = flat_then(120, tail=[100, 85, 70])
    h = holding(path, cost_per_share=100, price=70)
    ok = stops.classify(h, review(FRIDAY - dt.timedelta(days=3)), FRIDAY)
    assert ok.rung == 2 and not ok.needs_attention and "no deterioration" in ok.action
    bad = stops.classify(h, review(FRIDAY - dt.timedelta(days=3), deterioration=True), FRIDAY)
    assert bad.needs_attention and "TRIM ONE THIRD" in bad.action
    done = stops.classify(h, review(FRIDAY, conclusion="resized", deterioration=True), FRIDAY)
    assert not done.needs_attention                                               # the trim has been recorded
    none = stops.classify(h, None, FRIDAY)
    assert none.needs_attention and none.review_state in ("due", "overdue")


def test_rung_three_is_an_exit_review():
    path = flat_then(120, n_flat=40, tail=[110, 100, 95, 90, 85, 80, 75, 70, 65, 60])        # ten weeks under, 40% below cost
    h = holding(path, cost_per_share=100, price=60, rs_score=1)
    s = stops.classify(h, review(FRIDAY, "reaffirmed"), FRIDAY)
    assert s.rung == 3 and s.needs_attention and "EXIT REVIEW" in s.action
    assert not stops.classify(h, review(FRIDAY, "referred"), FRIDAY).needs_attention
    short = holding(flat_then(120, n_flat=40, tail=[100, 80, 60]), cost_per_share=100, price=60, rs_score=1)
    assert stops.classify(short, review(FRIDAY), FRIDAY).rung == 2                # only three weeks under: not rung 3 yet


def test_the_review_on_file_must_be_current():
    path = flat_then(120, tail=[100, 90, 78])
    h = holding(path, cost_per_share=100, price=78)
    assert stops.classify(h, review(FRIDAY - dt.timedelta(days=20), nxt=FRIDAY + dt.timedelta(days=30)), FRIDAY).review_state == "current"
    assert stops.classify(h, review(FRIDAY - dt.timedelta(days=95)), FRIDAY).review_state == "expired"
    assert stops.classify(h, review(FRIDAY - dt.timedelta(days=20), nxt=FRIDAY - dt.timedelta(days=1)), FRIDAY).review_state == "expired"
    assert stops.classify(h, review(FRIDAY - dt.timedelta(days=95)), FRIDAY).needs_attention
    fresh = stops.classify(h, None, FRIDAY)                                       # the run began within the last two weeks: due, not yet overdue
    assert fresh.review_state == "due" and fresh.needs_attention
    old = holding(flat_then(120, tail=[78] * 8), cost_per_share=100, price=78)
    assert stops.classify(old, None, FRIDAY).review_state == "overdue"


def test_additions_are_frozen_on_a_rung_and_by_an_explicit_freeze():
    on = stops.classify(holding(flat_then(120, tail=[100, 90, 78]), 100, price=78), review(FRIDAY), FRIDAY)
    assert on.adds_frozen
    off = stops.classify(holding(flat_then(100, tail=[100]), 100, price=100), review(FRIDAY, adds=True), FRIDAY)
    assert off.rung == 0 and off.adds_frozen                                      # the review itself froze additions, and it is current
    lifted = stops.classify(holding(flat_then(100, tail=[100]), 100, price=100), review(FRIDAY, adds=False), FRIDAY)
    assert not lifted.adds_frozen


# --------------------------------------------------------------------------- the registry

def test_registry_roundtrip(tmp_path):
    p = str(tmp_path / "s.json")
    assert stops.load_registry(p) == {}
    e = stops.record_review("adbe", "reaffirmed", on=D(2026, 10, 8), next_check=D(2026, 12, 9), journal_row=110, path=p)
    assert e["conclusion"] == "reaffirmed" and e["adds_frozen"] and not e["deterioration"]
    reg = stops.load_registry(p)
    assert reg["ADBE"]["next_check"] == "2026-12-09" and reg["ADBE"]["journal_row"] == 110
    stops.record_review("ADBE", "resized", on=D(2027, 1, 6), deterioration=True, path=p)            # replaces
    assert stops.load_registry(p)["ADBE"]["conclusion"] == "resized"
    with pytest.raises(stops.StopsError):
        stops.record_review("ADBE", "sold", path=p)
    assert json.load(open(p))["policy"]["source"].startswith("Journal row 109")


# --------------------------------------------------------------------------- a synthetic workbook, the audit check, the command

def make_book(tmp_path, bad_price=75.0):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cp = wb.create_sheet("Current Positions")
    for c, h in enumerate(["Account", "Ticker", "Description", "Asset Class", "Quantity", "Price", "Market Value", "Cost Basis", "G/L", "G/L %",
                           "% Acct", "% Combined", "Notes"], 1):
        cp.cell(row=2, column=c, value=h)
    rows = [("A", "FALL", "Falling Inc", "Equity", 100, bad_price, bad_price * 100, 10000.0),            # cost $100 a share
            ("A", "OKAY", "Okay Inc", "Equity", 100, 100.0, 10000.0, 9000.0),
            ("A", "ETFX", "Some ETF", "Equity", 50, 40.0, 2000.0, 1000.0),                              # unscored: not a stock for the ladder
            ("A", "FZ", "Cash fund", "Cash", 5000, 1.0, 5000.0, 5000.0)]
    for i, r in enumerate(rows):
        for c, v in enumerate(r, 1):
            cp.cell(row=3 + i, column=c, value=v)
        cp.cell(row=3 + i, column=12, value=r[6] / 27000.0)
    sc = wb.create_sheet("Scoring")
    sc.cell(row=2, column=1, value="Ticker")
    sc.cell(row=2, column=40, value="Industry")                                     # the reader indexes columns out to AM
    for i, (t, rs) in enumerate((("FALL", 1), ("OKAY", 4))):
        sc.cell(row=3 + i, column=1, value=t)
        sc.cell(row=3 + i, column=2, value=t + " Inc")
        sc.cell(row=3 + i, column=3, value=dt.datetime(2026, 9, 13))
        sc.cell(row=3 + i, column=20, value=rs)
        sc.cell(row=3 + i, column=21, value="M")
    md = wb.create_sheet("Market Data")
    for c, h in enumerate(["Ticker", "Company", "Price", "Volume", "Mkt Cap", "P/E", "52w Low", "52w High", "Div Yield"], 1):
        md.cell(row=2, column=c, value=h)
    for i, (t, p) in enumerate((("FALL", bad_price), ("OKAY", 100.0))):
        md.cell(row=3 + i, column=1, value=t)
        md.cell(row=3 + i, column=3, value=p)
    ph = wb.create_sheet("Price History")
    ph.cell(row=2, column=1, value="Week Ending")
    ph.cell(row=2, column=2, value="FALL")
    ph.cell(row=2, column=3, value="OKAY")
    fall = flat_then(120, n_flat=50, tail=[110, 100, 90, bad_price])
    ok = [100.0] * len(fall)
    for i, (d, c) in enumerate(weekly(fall)):
        ph.cell(row=3 + i, column=1, value=dt.datetime.combine(d, dt.time()))
        ph.cell(row=3 + i, column=2, value=c)
        ph.cell(row=3 + i, column=3, value=ok[i])
    path = str(tmp_path / "book.xlsx")
    wb.save(path)
    return path


def test_holdings_are_the_scored_stocks_only(tmp_path):
    hs = stops.read_holdings(make_book(tmp_path), today=FRIDAY)
    assert [h.ticker for h in hs] == ["FALL", "OKAY"]
    f = [h for h in hs if h.ticker == "FALL"][0]
    assert f.quantity == 100 and f.cost == 10000.0 and f.rs_score == 1 and len(f.closes) == 54


def test_the_ladder_the_audit_check_and_the_command(tmp_path, monkeypatch, capsys):
    reg = str(tmp_path / "stops.json")
    monkeypatch.setattr(stops, "REGISTRY_FILE", reg)
    book = make_book(tmp_path)                                                     # FALL is 25% below cost, under its average, RS score 1
    st = stops.ladder(book, today=FRIDAY)
    assert [(s.ticker, s.rung) for s in st] == [("FALL", 1), ("OKAY", 0)]
    assert st[0].needs_attention and st[0].adds_frozen
    chk = audit.check_stop_ladder(book, today=FRIDAY)[0]
    assert not chk.ok and "FALL" in chk.detail and "stops record" in chk.fix
    assert cli.main(["stops", "--workbook", book]) == 2
    out = capsys.readouterr().out
    assert "FALL" in out and "RE-UNDERWRITE" in out
    assert cli.main(["stops", "record", "FALL", "--conclusion", "reaffirmed", "--next", (D.today() + dt.timedelta(days=40)).isoformat(),
                     "--journal-row", "200"]) == 0
    chk = audit.check_stop_ladder(book, today=D.today())[0]
    assert chk.ok and "FALL rung 1" in chk.detail
    assert cli.main(["stops", "--workbook", book]) == 0
    assert cli.main(["stops", "record", "FALL"]) == 1                              # a conclusion is required


def test_the_summary_line():
    sm = {"n": 17, "on_rung": [], "approaching": [], "attention": []}
    assert stops.one_line(sm) == "stop-review ladder: none of 17 held scored stocks is on a rung"
    sm = {"n": 17, "on_rung": [("ADBE", 1, -0.21, "current")], "approaching": [("XYZ", -0.19)], "attention": []}
    line = stops.one_line(sm)
    assert "ADBE rung 1 (21.0% below cost, review current)" in line and "XYZ approaching (19.0% below cost)" in line


# --------------------------------------------------------------------------- the live workbook and registry

@needs_workbook
def test_live_workbook_covers_every_held_scored_stock_and_adbe_has_its_review():
    reg = stops.load_registry()
    assert reg["ADBE"]["conclusion"] == "reaffirmed" and reg["ADBE"]["adds_frozen"] and reg["ADBE"]["journal_row"] >= 110
    assert reg["ADBE"]["next_check"] == "2026-12-09"
    st = stops.ladder(_WB)
    held = {s.ticker for s in st}
    assert {"ADBE", "NVDA", "VEEV", "VRTX"} <= held
    assert all(s.rung in (0, 1, 2, 3) for s in st)
    adbe = [s for s in st if s.ticker == "ADBE"][0]
    assert adbe.loss < -0.15                                                       # about 19-21% below cost at the weekly prices
