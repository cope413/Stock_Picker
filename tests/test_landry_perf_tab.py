"""Tests for landry.perf_tab (the Performance Tracking lot ledger and its summary block) and for the weekly run's
SPY benchmark hook. Synthetic workbooks shaped like the live one; the formula tests go through a real LibreOffice
recalc (skipped when soffice is missing) and compare every summary line with arithmetic done here, not read back
from the sheet."""

import datetime as dt
import hashlib
import os

import openpyxl
import pytest
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table

from landry import market, perf_tab as pt, weekly
from landry.audit import check_merges_inside_tables, check_performance_tracking_ties
from landry.xlsx_io import latest_workbook
from landry.xlsx_recalc import recalc, soffice_path

D = dt.date
HEADERS = ["Ticker", "Company", "Entry\nDate", "Entry\nPrice", "Entry\nScore", "Entry\nConfidence",
           "Entry\nDecision Band", "SPY Price\nat Entry", "Status\n(Held / Exited)", "Exit\nDate", "Exit\nPrice",
           "Exit\nReason", "Current /\nExit Price", "SPY Price\n(current / exit)", "Total Return\nSince Entry",
           "SPY Return\nSame Period", "Excess\nReturn"]
PRICES = {"AAA": 120.0, "BBB": 55.0, "CCC": 9.0}                 # Market Data
BENCH = {"as_of": D(2026, 10, 2), "spy_now": 750.0, "spy_0805": 700.0, "spy_1231": 650.0, "inception": pt.INCEPTION}
needs_soffice = pytest.mark.skipif(soffice_path() is None, reason="soffice not installed")
_WB = latest_workbook(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
needs_workbook = pytest.mark.skipif(not _WB, reason="no workbook file")


def lots():
    return [
        pt.Lot("AAA", "Alpha", "Stock", "System", "Held", D(2026, 8, 7), 100.0, lot_shares=12, spy_entry=700.0,
               score=85.0, confidence="L", band="Strong Buy"),
        pt.Lot("AAA", "Alpha (pre-System)", "Stock", "Baseline", "Held", pt.INCEPTION, 90.0, lot_shares=38, ytd_price=80.0),
        pt.Lot("BBB", "Beta ETF", "ETF", "Baseline", "Held", pt.INCEPTION, 50.0, lot_shares=10, ytd_price=45.0),
        pt.Lot("CCC", "Gamma", "Stock", "Baseline", "Held", pt.INCEPTION, 10.0, lot_shares=5, ytd_price=8.0),
        pt.Lot("FZ", "Cash fund", "Cash", "Baseline", "Held", pt.INCEPTION, 1.0, ytd_price=1.0),
        pt.Lot("OLD", "Sold long ago", "Stock", "Baseline", "Exited", pt.INCEPTION, 20.0, lot_shares=10, ytd_price=15.0,
               exit_date=D(2026, 8, 7), exit_price=25.0, exit_reason="test", spy_exit=710.0),
    ]


def make_book(spare=2, with_lots=True):
    """Market Data, Current Positions (a Table), and Performance Tracking converted to the lot ledger the way the
    live tab was: the original A:Q Table, R:AC added, lots written, the summary block built."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    md = wb.create_sheet("Market Data")
    md["A1"] = "Market Data"
    for c, h in enumerate(["Ticker", "Company", "Price ($)", "Volume"], 1):
        md.cell(row=2, column=c, value=h)
    for i, (t, p) in enumerate(PRICES.items()):
        md.cell(row=3 + i, column=1, value=t)
        md.cell(row=3 + i, column=3, value=p)
    cp = wb.create_sheet("Current Positions")
    cp["A1"] = "Combined"
    heads = ["Account", "Ticker", "Description", "Asset Class", "Quantity", "Price ($)", "Market Value ($)",
             "Cost Basis ($)", "Unrealized G/L ($)", "Unrealized G/L (%)", "% of Account", "% of Combined",
             "Notes"]
    for c, h in enumerate(heads, 1):
        cp.cell(row=2, column=c, value=h)
    rows = [("Acct A", "AAA", "Alpha", "Equity", 30), ("Acct A", "BBB", "Beta ETF", "Equity", 10),
            ("Acct A", "FZ", "Cash fund", "Cash", 1000), ("Acct B", "AAA", "Alpha", "Equity", 20),
            ("Acct B", "CCC", "Gamma", "Equity", 5)]
    for i, (a, t, d, k, q) in enumerate(rows):
        r = 3 + i
        price = PRICES.get(t, 1.0)
        for c, v in enumerate((a, t, d, k, q, price, q * price, q * price * 0.9, 1.0, 0.1, 0.1, 0.1, ""), 1):
            cp.cell(row=r, column=c, value=v)
    cp.add_table(Table(displayName="CurrentPositionsTable", ref=f"A2:M{2 + len(rows)}"))

    ws = wb.create_sheet(pt.SHEET)
    ws["A1"] = "note"
    for c, h in enumerate(HEADERS, 1):
        ws.cell(row=2, column=c, value=h)
    n = len(lots()) if with_lots else 0
    last = 2 + n + spare
    tbl = Table(displayName=pt.TABLE, ref=f"A2:Q{last}")
    tbl.autoFilter = AutoFilter(ref=f"A2:Q{last}")
    ws.add_table(tbl)
    tbl._initialise_columns()                                     # openpyxl does this at save; ensure_columns needs it now
    for cell, col in zip(ws[tbl.ref][0], tbl.tableColumns):
        col.name = str(cell.value)
    pt.ensure_columns(ws)
    styles = pt.capture_styles(ws)
    for i, lot in enumerate(lots() if with_lots else []):
        pt.write_lot(ws, 3 + i, lot, styles)
    for r in range(3 + n, last + 1):
        pt.blank_row(ws, r, styles)
    pt.build_block(wb, ws, last, BENCH)
    return wb


def save(wb, tmp_path, name="p.xlsx"):
    path = str(tmp_path / name)
    wb.save(path)
    return path


def digest(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


# ----------------------------------------------------------------- the rows --

def test_a_lot_must_be_consistent():
    base = dict(ticker="X", company="X", type="Stock", basis="System", status="Held", entry_date=D(2026, 9, 1),
                entry_price=10.0, lot_shares=5)
    pt.Lot(**base).check()
    for change in ({"type": "Bond"}, {"basis": "Mixed"}, {"status": "Sold"}, {"lot_shares": None},
                   {"basis": "Baseline", "ytd_price": None},                       # no 12/31/25 close
                   {"basis": "Baseline", "ytd_price": 5.0, "type": "Cash"},       # a cash fund has no share snapshot
                   {"status": "Exited"}):                                          # no exit data
        with pytest.raises(pt.PerfTabError):
            pt.Lot(**{**base, **change}).check()


def test_write_lot_types_the_inputs_and_leaves_the_rest_to_formulas():
    wb = make_book()
    ws = wb[pt.SHEET]
    cell = lambda r, c: ws.cell(row=r, column=c)                              # noqa: E731
    # row 3: the System lot
    assert [cell(3, c).value for c in (pt.TICKER, pt.ENTRY_PRICE, pt.SPY_ENTRY, pt.LOT_SHARES, pt.TYPE, pt.BASIS)] == \
        ["AAA", 100.0, 700.0, 12, "Stock", "System"]
    assert cell(3, pt.YTD_PRICE).value == "=$D3"                              # System lots start the year at entry
    assert "INDEX('Market Data'!$C$3:$C$300,MATCH($A3" in cell(3, pt.CURRENT).value
    assert cell(3, pt.SPY_NOW).value == "=PT_SPY_Now"
    # row 4: a held baseline lot -- its 8/5 share count typed (the audit's snapshot), SPY at entry from the benchmark
    # cell, a typed 12/31/25 close; its Shares are a formula that follows Current Positions
    assert cell(4, pt.LOT_SHARES).value == 38 and cell(4, pt.SPY_ENTRY).value == "=PT_SPY_0805"
    assert cell(4, pt.YTD_PRICE).value == 80.0 and cell(4, pt.ENTRY_DATE).value == dt.datetime(2026, 8, 5)
    assert "CurrentPositionsTable[Quantity]" in cell(4, pt.SHARES).value
    # row 7: a cash fund is priced at $1.00, not looked up; row 8: a sold lot keeps its sale price and SPY at exit
    assert cell(7, pt.CURRENT).value == 1.0
    assert (cell(8, pt.STATUS).value, cell(8, pt.CURRENT).value, cell(8, pt.SPY_NOW).value, cell(8, pt.LOT_SHARES).value) \
        == ("Exited", 25.0, 710.0, 10)
    # a spare row: formulas in place, nothing typed
    assert cell(9, pt.TICKER).value is None and cell(9, pt.SHARES).value.startswith('=IF($A9=""')


def test_etf_and_cash_tickers_are_green_unless_sold():
    wb = make_book()
    ws = wb[pt.SHEET]
    def green(r):
        c = ws.cell(row=r, column=pt.TICKER)
        return c.fill.fill_type == "solid" and c.fill.fgColor.rgb == pt.LIGHT_GREEN and c.font.color.rgb == pt.DARK_GREEN
    assert [green(r) for r in range(3, 8)] == [False, False, True, False, True]    # BBB (ETF) and FZ (Cash); not OLD
    # selling a green lot turns it back
    pt.close_lot(ws, 5, D(2026, 10, 9), 60.0, "test", 760.0, shares=10)
    assert not green(5) and ws.cell(row=5, column=pt.STATUS).value == "Exited" and ws.cell(row=5, column=pt.LOT_SHARES).value == 10


def test_close_lot_needs_the_shares_for_a_baseline_lot_and_keeps_a_system_lots_own():
    wb = make_book()
    ws = wb[pt.SHEET]
    with pytest.raises(pt.PerfTabError):
        pt.close_lot(ws, 4, D(2026, 10, 9), 130.0, "test", 760.0)             # baseline: how many shares were sold?
    pt.close_lot(ws, 3, D(2026, 10, 9), 130.0, "test", 760.0)                  # the System lot knows its 12
    assert ws.cell(row=3, column=pt.LOT_SHARES).value == 12 and ws.cell(row=3, column=pt.EXIT_PRICE).value == 130.0
    with pytest.raises(pt.PerfTabError):
        pt.close_lot(ws, 3, D(2026, 10, 9), 130.0, "again", 760.0)             # no longer held


# ------------------------------------------------------------ the summary block --

def test_the_block_has_alans_lines_in_order_and_formulas_bounded_to_the_table():
    wb = make_book()
    ws = wb[pt.SHEET]
    _first, last = pt.table_bounds(ws)
    top = pt.block_top(last)
    labels = [ws.cell(row=top + i, column=1).value for i in range(1, 8)]
    assert labels == ["Subtotal", "TOTAL", "YTD Subtotal", "YTD TOTAL", "Cumulative Subtotal",
                      "Cumulative TOTAL (since inception 8/5)", "Avg Annual Gain/(Loss)"]
    sub = ws.cell(row=top + 1, column=4).value
    assert f"$3:$V${last}" in sub or f"$V$3:$V${last}" in sub
    assert '"Stock"' in sub and '"Held"' in sub                                # stocks only, held only
    assert '"Stock"' not in ws.cell(row=top + 2, column=4).value               # TOTAL: every held position
    assert '"Held"' not in ws.cell(row=top + 4, column=4).value                # YTD TOTAL: held and sold
    assert ws.cell(row=top + 1, column=1).value == "Subtotal" and ws.cell(row=top + 6, column=1).alignment.horizontal == "left"
    for key, name in pt.NAMES.items():                                          # every benchmark name points into the block
        sheet, ref = pt._name_ref(wb, key)
        row = int(ref[1:])
        assert sheet == pt.SHEET and top < row < top + pt.BLOCK_ROWS
    assert [ws[pt._name_ref(wb, k)[1]].value for k in pt.NAMES] == \
        [dt.datetime(2026, 10, 2), 750.0, 700.0, 650.0, dt.datetime(2026, 8, 5)]


def test_the_tabs_formulas_have_no_circular_reference_and_no_whole_column_range():
    """Excel reports a range-level cycle; LibreOffice does not (it evaluates SUMIFS lazily) -- so this is the only
    thing that would have caught Shares reading $I:$I, a column the summary's Excess formulas also live in."""
    wb = make_book()
    assert pt.find_cycles(wb) == []
    ws = wb[pt.SHEET]
    import re
    for r in (3, 4, 9):
        for c in range(1, pt.LAST_COL + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.startswith("="):
                assert not re.search(r"\$[A-Z]+:\$[A-Z]+", v), (r, c, v)
    # the old formula, with whole columns, is a cycle through the summary block
    ws.cell(row=4, column=pt.SHARES).value = (
        '=IF(OR($S4="System",$I4="Exited"),$T4,SUMIFS(CurrentPositionsTable[Quantity],CurrentPositionsTable[Ticker],$A4)'
        '-SUMIFS($T:$T,$A:$A,$A4,$S:$S,"System",$I:$I,"Held"))')
    cycles = pt.find_cycles(wb)
    assert cycles and "U4" in cycles[0]


def test_no_merged_cell_touches_the_table(tmp_path):
    path = save(make_book(), tmp_path)
    (check,) = check_merges_inside_tables(path)
    assert check.ok, check.detail


def test_add_lot_uses_a_spare_row_then_grows_the_table_and_moves_the_block():
    wb = make_book(spare=1)
    ws = wb[pt.SHEET]
    _first, last0 = pt.table_bounds(ws)
    new = lambda t: pt.Lot(t, t, "Stock", "System", "Held", D(2026, 10, 7), 10.0, lot_shares=3, spy_entry=745.0,   # noqa: E731
                           score=80.0, confidence="L", band="Buy")
    assert pt.add_lot(wb, new("DDD")) == last0                                 # the one spare row
    assert pt.table_bounds(ws)[1] == last0
    r = pt.add_lot(wb, new("EEE"), grow_by=3)                                   # none left: grow by 3
    _first, last1 = pt.table_bounds(ws)
    assert r == last0 + 1 and last1 == last0 + 3
    assert ws.tables[pt.TABLE].ref == f"A2:{pt.L(pt.LAST_COL)}{last1}" == ws.tables[pt.TABLE].autoFilter.ref
    # the block moved intact: header, lines, benchmark cells (values carried), formulas on the new bounds
    top = pt.block_top(last1)
    assert ws.cell(row=top, column=1).value == "Performance summary"
    assert f"${last1}" in ws.cell(row=top + 1, column=4).value and f"${last0}" not in ws.cell(row=top + 1, column=4).value
    # every row's Shares formula was rewritten to the new bounds, not just the new rows'
    for rr in (3, 4, last1):
        assert f"$T$3:$T${last1}" in ws.cell(row=rr, column=pt.SHARES).value and f"${last0}," not in ws.cell(row=rr, column=pt.SHARES).value
    assert pt.find_cycles(wb) == []
    assert [ws[pt._name_ref(wb, k)[1]].value for k in pt.NAMES] == \
        [dt.datetime(2026, 10, 2), 750.0, 700.0, 650.0, dt.datetime(2026, 8, 5)]
    # the old block's rows are table rows now: no label, no merge left behind
    old_top = pt.block_top(last0)
    for rr in range(last0 + 1, last1 + 1):
        assert ws.cell(row=rr, column=1).value in (None, "EEE")
    assert not [m for m in ws.merged_cells.ranges if m.min_row <= last1]
    # validations and colour rules reach the new last row
    assert {str(dv.sqref) for dv in ws.data_validations.dataValidation} == {
        f"I3:I{last1}", f"G3:G{last1}", f"R3:R{last1}", f"S3:S{last1}"}
    assert f"O3:O{last1}" in {str(cf.sqref) for cf in ws.conditional_formatting}
    assert old_top != top


@needs_soffice
def test_growing_the_table_keeps_the_workbook_calculating_and_openable(tmp_path):
    wb = make_book(spare=0)
    pt.add_lot(wb, pt.Lot("CCC", "Gamma add", "Stock", "System", "Held", D(2026, 10, 7), 9.5, lot_shares=2,
                          spy_entry=745.0), grow_by=2)
    path = save(wb, tmp_path)
    res = recalc(path)
    assert res["status"] == "success" and res["total_errors"] == 0, res
    (check,) = check_merges_inside_tables(path)
    assert check.ok
    # CCC is now 3 baseline shares (5 on Current Positions less the 2 the System lot took) plus the 2-share lot
    v = openpyxl.load_workbook(path, data_only=True)[pt.SHEET]
    rows = {(v.cell(row=r, column=1).value, v.cell(row=r, column=pt.BASIS).value): v.cell(row=r, column=pt.SHARES).value
            for r in pt.lot_rows(v)}
    assert rows[("CCC", "Baseline")] == 3 and rows[("CCC", "System")] == 2


# ------------------------------------- the formulas, through a real recalc, against arithmetic done here --

@needs_soffice
def test_every_summary_line_matches_independent_arithmetic(tmp_path):
    path = save(make_book(), tmp_path)
    res = recalc(path)
    assert res["status"] == "success" and res["total_errors"] == 0, res
    v = openpyxl.load_workbook(path, data_only=True)[pt.SHEET]
    approx = lambda x: pytest.approx(x, rel=1e-9, abs=1e-9)                    # noqa: E731

    spy_now, spy0, spy1231 = 750.0, 700.0, 650.0
    # (type, status, shares, entry price, spy at entry, current/exit price, ytd base price, spy at exit-or-now)
    L = [("Stock", "Held", 12, 100.0, 700.0, 120.0, 100.0, spy_now),
         ("Stock", "Held", 38, 90.0, spy0, 120.0, 80.0, spy_now),            # 50 on Current Positions less the System lot's 12
         ("ETF", "Held", 10, 50.0, spy0, 55.0, 45.0, spy_now),
         ("Stock", "Held", 5, 10.0, spy0, 9.0, 8.0, spy_now),
         ("Cash", "Held", 1000, 1.0, spy0, 1.0, 1.0, spy_now),
         ("Stock", "Exited", 10, 20.0, spy0, 25.0, 15.0, 710.0)]
    got = {r: [v.cell(row=r, column=c).value for c in (pt.SHARES, pt.ENTRY_VALUE, pt.CUR_VALUE, pt.GAIN, pt.RET,
                                                       pt.SPY_RET, pt.YTD_VALUE, pt.YTD_GAIN, pt.YTD_RET, pt.SPY_GAIN)]
           for r in range(3, 9)}
    for r, (typ, st, sh, entry, spy_e, cur, ytd, spy_x) in zip(range(3, 9), L):
        ev, cv = sh * entry, sh * cur
        spy_ret = spy_x / spy_e - 1
        want = [sh, ev, cv, cv - ev, cur / entry - 1, spy_ret, sh * ytd, cv - sh * ytd, cur / ytd - 1, ev * spy_ret]
        assert [approx(x) for x in want] == got[r], (r, typ)

    def agg(select, base):                                                     # base: "entry" or "ytd"
        b = sum((sh * (entry if base == "entry" else ytd)) for (typ, st, sh, entry, _s, _c, ytd, _x) in L if select(typ, st))
        c = sum(sh * cur for (typ, st, sh, _e, _s, cur, _y, _x) in L if select(typ, st))
        return b, c, c - b, (c - b) / b

    top = pt.block_top(pt.table_bounds(v)[1])
    line = lambda i: [v.cell(row=top + i, column=c).value for c in range(4, 10)]       # noqa: E731
    sp_lot = lambda select: sum(sh * entry * (spx / se - 1) for (typ, st, sh, entry, se, _c, _y, spx) in L if select(typ, st))  # noqa: E731

    held_stock = lambda t, s: t == "Stock" and s == "Held"                    # noqa: E731
    held = lambda t, s: s == "Held"                                          # noqa: E731
    stock = lambda t, s: t == "Stock"                                        # noqa: E731
    everything = lambda t, s: True                                           # noqa: E731
    for i, (sel, base, spy) in enumerate([(held_stock, "entry", None), (held, "entry", None), (stock, "ytd", spy_now / spy1231 - 1),
                                          (everything, "ytd", spy_now / spy1231 - 1), (stock, "entry", spy_now / spy0 - 1),
                                          (everything, "entry", spy_now / spy0 - 1)], start=1):
        b, c, g, ret = agg(sel, base)
        if spy is None:                                                        # the first pair: each lot's own SPY period
            spy = sp_lot(sel) / b
        assert line(i) == [approx(b), approx(c), approx(g), approx(ret), approx(spy), approx(ret - spy)], i
    # Avg Annual: Cumulative TOTAL / years since 8/5/26 (to the as-of date), simple
    b, c, g, ret = agg(everything, "entry")
    yrs = (D(2026, 10, 2) - D(2026, 8, 5)).days / 365
    spy_c = spy_now / spy0 - 1
    f = line(7)
    assert f[2:] == [approx(g / yrs), approx(ret / yrs), approx(spy_c / yrs), approx(ret / yrs - spy_c / yrs)]
    assert f[0] in (None, "") and f[1] in (None, "")


@needs_soffice
def test_the_audit_ties_hold_on_a_consistent_book_and_catch_a_lot_that_drifts(tmp_path):
    path = save(make_book(), tmp_path)
    assert recalc(path)["total_errors"] == 0
    checks = check_performance_tracking_ties(path)
    assert [c.name.split(":")[1] for c in checks] == ["shares", "baseline", "value", "summary", "cycles", "benchmark"]
    assert all(c.ok for c in checks), [c.detail for c in checks if not c.ok]
    # a System lot bigger than the position: Current Positions holds 5 CCC, the System says it bought 9
    wb = openpyxl.load_workbook(path)
    pt.add_lot(wb, pt.Lot("CCC", "Gamma add", "Stock", "System", "Held", D(2026, 10, 7), 9.5, lot_shares=9, spy_entry=745.0))
    path2 = save(wb, tmp_path, "p2.xlsx")
    assert recalc(path2)["total_errors"] == 0
    bad = {c.name.split(":")[1]: c for c in check_performance_tracking_ties(path2) if not c.ok}
    assert "shares" in bad and "CCC" in bad["shares"].detail


@needs_soffice
def test_a_purchase_recorded_only_on_current_positions_is_caught_until_it_is_added_as_a_lot(tmp_path):
    """The silent failure the typed 8/5 share count exists for: the baseline lot follows Current Positions, so +100
    shares there would simply be measured from the 8/5 close."""
    wb = make_book()
    cp = wb["Current Positions"]
    cp["E3"].value = 130                                         # AAA, Acct A: 30 -> 130 shares
    cp["G3"].value = 130 * 120.0
    path = save(wb, tmp_path)
    assert recalc(path)["total_errors"] == 0
    bad = {c.name.split(":")[1]: c for c in check_performance_tracking_ties(path) if not c.ok}
    assert list(bad) == ["baseline"] and "AAA" in bad["baseline"].detail and "+100" in bad["baseline"].detail
    wb = openpyxl.load_workbook(path)
    pt.add_lot(wb, pt.Lot("AAA", "Alpha add", "Stock", "System", "Held", D(2026, 10, 7), 118.0, lot_shares=100,
                          spy_entry=745.0))
    path2 = save(wb, tmp_path, "p2.xlsx")
    assert recalc(path2)["total_errors"] == 0
    assert all(c.ok for c in check_performance_tracking_ties(path2))        # the baseline gave the 100 up


def test_rebase_baseline_resets_only_the_lots_whose_quantity_moved():
    wb = make_book()
    wb["Current Positions"]["E4"].value = 10.5                   # BBB: dividend reinvestment, half a share
    assert pt.rebase_baseline(wb) == [("BBB", 10, 10.5)]
    ws = wb[pt.SHEET]
    assert ws.cell(row=5, column=pt.LOT_SHARES).value == 10.5
    assert ws.cell(row=4, column=pt.LOT_SHARES).value == 38 and ws.cell(row=7, column=pt.LOT_SHARES).value is None   # AAA, cash
    assert pt.rebase_baseline(wb) == []


# ------------------------------------------------------------- the weekly hook --

def refresh_book(tmp_path, wb=None, **kw):
    wb = wb or make_book()
    mon = wb.create_sheet("Monitor & Recheck Triggers")
    for c, h in enumerate(["Ticker", "Category", "x", "x", "x", "x", "x", "Current\nPrice", "x", "Next/Last\nEarnings Date"], 1):
        mon.cell(row=2, column=c, value=h)
    for c, h in enumerate(["Ticker", "Company", "Price ($)", "Volume", "Market Cap ($M)", "P/E", "52-Wk Low",
                           "52-Wk High", "Div Yield (%)"], 1):
        wb["Market Data"].cell(row=2, column=c, value=h)
    path = save(wb, tmp_path)
    snaps = {t: {"price": p * 1.01, "volume": 10} for t, p in PRICES.items()}
    snaps["SPY"] = {"price": 760.0}
    snaps.update(kw.pop("snaps", {}))
    args = dict(snapshot=lambda t: snaps.get(t), earnings=lambda t: None, today=D(2026, 10, 10), sleep=lambda s: None)
    args.update(kw)
    return path, args


def tab_cells(path):
    ws = openpyxl.load_workbook(path)[pt.SHEET]
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}


def test_the_weekly_run_writes_spy_and_the_as_of_date_and_nothing_else_on_the_tab(tmp_path):
    path, kw = refresh_book(tmp_path)
    before = tab_cells(path)
    rep = market.refresh(path, **kw)
    perf = rep["performance"]
    assert perf["present"] and perf["changed"] and rep["wrote"]
    assert (perf["spy_old"], perf["spy_new"]) == (750.0, 760.0)
    assert perf["as_of_new"] == D(2026, 10, 9)                                   # Saturday 10/10 -> Friday's close
    assert perf["held"] == 4 and perf["no_price"] == []                          # the cash fund and the sold lot are exempt
    after = tab_cells(path)
    changed = {k for k in set(before) | set(after) if before.get(k) != after.get(k)}
    wb = openpyxl.load_workbook(path)
    spy_cell, date_cell = pt._name_ref(wb, "spy_now")[1], pt._name_ref(wb, "as_of")[1]
    assert changed == {spy_cell, date_cell}                                       # nothing else on the tab moved
    assert wb[pt.SHEET][spy_cell].value == 760.0 and wb[pt.SHEET][date_cell].value == dt.datetime(2026, 10, 9)


def test_a_second_weekly_run_over_unchanged_data_does_not_save(tmp_path):
    path, kw = refresh_book(tmp_path)
    market.refresh(path, **kw)
    h = digest(path)
    rep = market.refresh(path, **kw)
    assert not rep["performance"]["changed"] and not rep["wrote"] and digest(path) == h


def test_a_missing_spy_quote_keeps_the_benchmark_and_is_reported(tmp_path):
    path, kw = refresh_book(tmp_path, snaps={"SPY": None})
    rep = market.refresh(path, **kw)
    assert rep["performance"]["spy_failed"] and not rep["performance"]["changed"]
    wb = openpyxl.load_workbook(path)
    assert wb[pt.SHEET][pt._name_ref(wb, "spy_now")[1]].value == 750.0
    text = weekly.format_report({"prices": None, "market": rep, "problems": [], "warnings": [], "wrote": rep["wrote"],
                                 "dry_run": False, "positions_only": False})
    assert "SPY NOT refreshed (no quote)" in text


def test_positions_only_never_fetches_spy_but_still_finds_unpriced_lots(tmp_path):
    wb = make_book()
    pt.add_lot(wb, pt.Lot("ZZZ", "Not in Market Data", "Stock", "System", "Held", D(2026, 10, 7), 10.0, lot_shares=1,
                          spy_entry=745.0))
    path, kw = refresh_book(tmp_path, wb=wb, do_market=False, do_earnings=False,
                            snapshot=lambda t: (_ for _ in ()).throw(AssertionError("positions-only must not fetch")))
    rep = market.refresh(path, **kw)
    perf = rep["performance"]
    assert perf["present"] and not perf["changed"] and not perf["spy_failed"]
    assert perf["no_price"] == ["ZZZ"]                                           # the cash fund is exempt: it is $1.00


def test_a_big_spy_move_is_a_warning_not_a_refusal(tmp_path):
    path, kw = refresh_book(tmp_path, snaps={"SPY": {"price": 900.0}})
    rep = market.refresh(path, **kw)
    assert rep["performance"]["changed"] and any("from the sheet's previous 750" in w for w in rep["performance"]["warnings"])
    wb = openpyxl.load_workbook(path)
    assert wb[pt.SHEET][pt._name_ref(wb, "spy_now")[1]].value == 900.0


def test_the_report_says_what_happened_on_the_performance_tab():
    base = {"present": True, "held": 30, "no_price": [], "warnings": [], "spy_failed": False, "changed": True,
            "spy_old": 750.0, "spy_new": 760.0, "as_of_old": D(2026, 9, 25), "as_of_new": D(2026, 10, 2)}
    rep = {"dry_run": False, "positions_only": False}
    (line,) = weekly._performance_lines(base, rep)
    assert line == ("Performance     SPY now $760.00 as of 2026-10-02 (was $750.00 as of 2026-09-25); "
                    "30 held lots price from Market Data")
    assert "would be" in weekly._performance_lines(base, {**rep, "dry_run": True})[0]
    assert "unchanged" in weekly._performance_lines({**base, "changed": False}, rep)[0]
    assert "positions only" in weekly._performance_lines(base, {**rep, "positions_only": True})[0]
    assert weekly._performance_lines({"present": False}, rep) == []


def test_a_workbook_without_the_ledger_is_untouched_by_the_hook(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    from test_landry_market import add_market_tabs
    add_market_tabs(wb)
    path = save(wb, tmp_path)
    rep = market.refresh(path, snapshot=lambda t: {"price": 1.0}, earnings=lambda t: None, today=D(2026, 10, 3),
                         sleep=lambda s: None, do_positions=True)
    assert rep["performance"]["present"] is False and not rep["performance"]["changed"]


# --------------------------------------------------------------- the live workbook --

@needs_workbook
def test_live_ledger_is_consistent_and_its_summary_matches_arithmetic_on_its_own_rows():
    wb = openpyxl.load_workbook(_WB, data_only=True)
    ws = wb[pt.SHEET]
    ls = pt.read_lots(wb)
    assert len(ls) >= 39 and {l["basis"] for l in ls} == {"System", "Baseline"} and {l["type"] for l in ls} == set(pt.TYPES)
    # System lots carry the whole Part 9 record; baseline lots never carry a score or a band
    for l in ls:
        if l["basis"] == "System":
            assert isinstance(l["lot_shares"], (int, float)) and l["lot_shares"] > 0
        else:
            assert l["score"] is None and l["band"] is None and l["entry_date"] == dt.datetime(2026, 8, 5)
            if l["status"] == "Held" and l["type"] != "Cash":                       # the typed 8/5 snapshot
                assert l["lot_shares"] == pytest.approx(l["shares"], abs=0.5)
    num = lambda v: v if isinstance(v, (int, float)) else 0.0                    # noqa: E731
    top = pt.block_top(pt.table_bounds(ws)[1])
    got = {ws.cell(row=top + i, column=1).value: [ws.cell(row=top + i, column=c).value for c in range(4, 10)] for i in range(1, 8)}
    sel = {"Subtotal": (lambda l: l["type"] == "Stock" and l["status"] == "Held", "entry_value"),
           "TOTAL": (lambda l: l["status"] == "Held", "entry_value"),
           "YTD Subtotal": (lambda l: l["type"] == "Stock", "ytd_value"),
           "YTD TOTAL": (lambda l: True, "ytd_value"),
           "Cumulative Subtotal": (lambda l: l["type"] == "Stock", "entry_value"),
           "Cumulative TOTAL (since inception 8/5)": (lambda l: True, "entry_value")}
    for label, (pick, base_key) in sel.items():
        rows = [l for l in ls if pick(l)]
        base, cur = sum(num(l[base_key]) for l in rows), sum(num(l["cur_value"]) for l in rows)
        b, c, g, r = got[label][:4]
        assert (b, c) == (pytest.approx(base), pytest.approx(cur)), label
        assert g == pytest.approx(cur - base) and r == pytest.approx((cur - base) / base), label
    assert got["TOTAL"][1] == pytest.approx(sum(num(l["cur_value"]) for l in ls if l["status"] == "Held"))


@needs_workbook
def test_live_ledger_has_no_circular_reference():
    assert pt.find_cycles(openpyxl.load_workbook(_WB)) == []


@needs_workbook
def test_live_audit_ties_hold():
    checks = check_performance_tracking_ties(_WB)
    assert [c.ok for c in checks] == [True] * 6, [c.detail for c in checks if not c.ok]


@needs_workbook
def test_live_type_decides_the_green_tickers_and_matches_current_positions():
    wb = openpyxl.load_workbook(_WB)
    ws = wb[pt.SHEET]
    for r in pt.lot_rows(ws):
        a = ws.cell(row=r, column=pt.TICKER)
        green = a.fill.fill_type == "solid" and a.fill.fgColor.rgb == pt.LIGHT_GREEN and a.font.color.rgb == pt.DARK_GREEN
        held = ws.cell(row=r, column=pt.STATUS).value == "Held"
        assert green == (held and ws.cell(row=r, column=pt.TYPE).value in pt.CASH_LIKE), a.value
