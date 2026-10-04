"""Offline tests for landry.market -- synthetic workbooks and fake providers (no network, no LibreOffice)."""

import datetime as dt
import hashlib

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from landry import market
from landry.market import MarketError

TODAY = dt.date(2026, 10, 4)
D = dt.date

MD_HEAD = ["Ticker", "Company", "Price ($)", "Volume", "Market Cap ($M)", "P/E", "52-Wk Low", "52-Wk High",
           "Div Yield (%)"]
MON_HEAD = ["Ticker", "Category", "Last Score\nDate", "Last\nComposite", "Last Tier 1\nWtd Avg", "Last\nDecision",
            "Price at\nLast Score", "Current\nPrice", "% Price\nChange", "Next/Last\nEarnings Date"]


def add_market_tabs(wb, tickers=("AAA", "BBB", "CCC")):
    """Market Data (title row 1, header row 2, rows from 3, a footer formula and a label under the data)
    and the Monitor tab, shaped like the live ones."""
    md = wb.create_sheet("Market Data")
    md["A1"] = "Market Data -- Claude-populated via yfinance"
    md.merge_cells("A1:I1")
    for c, h in enumerate(MD_HEAD, 1):
        md.cell(row=2, column=c, value=h)
    fill = PatternFill(fill_type="solid", fgColor="FFDDEBF7")
    rows = [(tickers[0], "Alpha Corp", 100.0, 1000, 5000.5, 20.0, 80.0, 120.0, 1.5),
            (tickers[1], "Beta ETF", 50.0, 500, 9999.0, 15.0, 40.0, 60.0, 3.0),
            (tickers[2], "Gamma Inc", 10.0, 100, None, None, 5.0, 12.0, None)]
    for i, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            md.cell(row=3 + i, column=c, value=v).fill = fill
        md.cell(row=3 + i, column=3).number_format = "\\$#,##0.00"
    md["K3"] = "=C3*2"                                               # a formula outside the columns written
    md.cell(row=7, column=1, value="=SUBTOTAL(103,A3:A6)")           # a footer formula is not a ticker
    md.cell(row=8, column=1, value="Note: refreshed weekly")         # nor is a label

    mon = wb.create_sheet("Monitor & Recheck Triggers")
    for c, h in enumerate(MON_HEAD, 1):
        mon.cell(row=2, column=c, value=h)
    for i, (t, cat, d) in enumerate([(tickers[0], "Owned", dt.datetime(2026, 9, 1)),
                                     (tickers[1], "Owned", dt.datetime(2026, 10, 10)),
                                     (tickers[2], "Watch List", dt.datetime(2026, 11, 20))]):
        mon.cell(row=3 + i, column=1, value=t)
        mon.cell(row=3 + i, column=2, value=cat)
        mon.cell(row=3 + i, column=10, value=d).number_format = "mm/dd/yyyy"
    other = wb.create_sheet("Other")
    other["A1"] = "untouched"
    return wb


def build(tmp_path, name="m.xlsx", **kw):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    add_market_tabs(wb, **kw)
    path = str(tmp_path / name)
    wb.save(path)
    return path


def digest(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def snap(price=101.5, volume=2000, cap=6.0e9, pe=21.0, lo=80.0, hi=125.0, dy=1.5, rate=1.5, trail=None):
    return {"price": price, "volume": volume, "market_cap": cap, "pe": pe, "wk52_low": lo, "wk52_high": hi,
            "dividend_yield": dy, "dividend_rate": rate, "dividend_trailing_rate": trail}


def table(rows):
    return lambda t: rows.get(t)


GOOD = {"AAA": snap(), "BBB": snap(price=51.0, cap=None, dy=3.1, rate=None),
        "CCC": snap(price=10.5, cap=2.0e9, pe=None, dy=None, rate=None)}


def values(path, sheet, row, cols):
    ws = openpyxl.load_workbook(path)[sheet]
    return [ws.cell(row=row, column=c).value for c in cols]


# ------------------------------------------------------------------ normalize --

def test_normalize_converts_units_and_drops_unusable_values():
    out, warns = market.normalize(snap(price=101.5, cap=6.0e9, volume=1999.6, pe="Infinity"))
    assert out["price"] == 101.5 and out["market_cap_m"] == 6000.0 and out["volume"] == 2000
    assert "pe" not in out and warns == []                           # 'Infinity' is not a P/E
    out, _ = market.normalize(snap(pe=-4.0))
    assert "pe" not in out                                           # negative earnings: no P/E, not a negative one
    out, warns = market.normalize(snap(price=0))
    assert "price" not in out and any("unusable price" in w for w in warns)
    out, warns = market.normalize(snap(lo=130.0, hi=120.0))
    assert "wk52_low" not in out and "wk52_high" not in out and any("52-week" in w for w in warns)
    assert market.normalize({})[0] == {}                             # a failed fetch yields nothing, not zeros


def test_normalize_warns_about_a_big_price_move_but_keeps_the_value():
    out, warns = market.normalize(snap(price=7.0), previous={"price": 70.0})      # a 10:1 split
    assert out["price"] == 7.0 and any("-90%" in w for w in warns)


def test_dividend_yield_unit_change_is_caught_for_stocks_and_funds():
    out, warns = market.normalize(snap(price=100.0, dy=0.015, rate=1.5))          # a fraction, rate/price = 1.5%
    assert out["dividend_yield"] == pytest.approx(1.5) and any("fraction" in w for w in warns)
    out, warns = market.normalize(snap(price=100.0, dy=1.48, rate=1.5))           # already a percent: untouched
    assert out["dividend_yield"] == 1.48 and warns == []
    out, warns = market.normalize(snap(dy=0.03, rate=None), previous={"dividend_yield": 3.0})   # a fund: no rate
    assert "dividend_yield" not in out and any("unit change" in w for w in warns)
    out, warns = market.normalize(snap(dy=3.2, rate=None), previous={"dividend_yield": 3.0})
    assert out["dividend_yield"] == 3.2 and warns == []


def test_a_yield_of_exactly_zero_with_no_rate_reads_as_nothing_paid():
    out, _ = market.normalize(snap(dy=0.0, rate=None, trail=None))
    assert "dividend_yield" in out and out["dividend_yield"] is None              # Yahoo's 0.0 for a non-payer: blank
    out, _ = market.normalize(snap(dy=0.0, rate=0.04, trail=None))
    assert out["dividend_yield"] == 0.0                                           # a (tiny) payer keeps its zero


def test_yahoo_flipping_between_none_and_zero_for_a_non_payer_leaves_the_file_alone(tmp_path):
    path = build(tmp_path)
    rows = {"AAA": snap(dy=None, rate=None, trail=0.0), "BBB": GOOD["BBB"], "CCC": GOOD["CCC"]}
    first = market.refresh(path, snapshot=table(rows), earnings=lambda t: None, today=TODAY, sleep=lambda s: None)
    assert first["wrote"] and values(path, "Market Data", 3, [9]) == [None]
    rows["AAA"] = snap(dy=0.0, rate=None, trail=0.0)                              # the same non-payer, as 0.0 this time
    second = market.refresh(path, snapshot=table(rows), earnings=lambda t: None, today=TODAY, sleep=lambda s: None)
    assert not second["wrote"] and second["market"]["cells_changed"] == 0
    assert values(path, "Market Data", 3, [9]) == [None]


def test_a_yield_is_cleared_only_when_the_provider_says_nothing_was_paid():
    out, _ = market.normalize(snap(dy=None, rate=None, trail=0.0))
    assert "dividend_yield" in out and out["dividend_yield"] is None             # affirmative "nothing": clear it
    out, _ = market.normalize(snap(dy=None, rate=None, trail=None))
    assert "dividend_yield" not in out                                           # silence: keep what is there
    out, _ = market.normalize(snap(dy=None, rate=None, trail=2.4))
    assert "dividend_yield" not in out                                           # it paid, the yield field is missing: a glitch


@pytest.mark.parametrize("old,new,expected", [
    (None, D(2026, 11, 2), D(2026, 11, 2)),                           # nothing there yet
    (dt.datetime(2026, 11, 2), D(2026, 11, 2), None),                 # same
    (dt.datetime(2026, 11, 2), D(2026, 11, 5), D(2026, 11, 5)),       # a later-confirmed date replaces a future one
    (dt.datetime(2026, 11, 10), D(2026, 11, 5), D(2026, 11, 5)),      # ... including an earlier one
    (dt.datetime(2026, 11, 2), D(2026, 8, 6), None),                  # a future date never gives way to a past one
    (dt.datetime(2026, 9, 10), D(2026, 12, 9), D(2026, 12, 9)),       # a past date gives way to the next one
    (dt.datetime(2026, 9, 10), D(2026, 8, 6), None),                  # ... but never moves further back
    (dt.datetime(2026, 8, 6), D(2026, 9, 10), D(2026, 9, 10)),        # a newer past date is fine
    (dt.datetime(2026, 9, 10), None, None),
    (dt.datetime(2026, 9, 10), "soon", None),
])
def test_pick_earnings_date(old, new, expected):
    assert market.pick_earnings_date(old, new, TODAY) == expected


# -------------------------------------------------------------------- refresh --

def test_refresh_writes_each_field_to_its_column_and_nothing_else(tmp_path):
    path = build(tmp_path)
    rep = market.refresh(path, snapshot=table(GOOD), earnings=lambda t: None, today=TODAY)
    assert rep["wrote"] and rep["market"]["refreshed"] == 3 and rep["market"]["rows"] == 3
    assert values(path, "Market Data", 3, range(3, 10)) == [101.5, 2000, 6000.0, 21.0, 80.0, 125.0, 1.5]
    assert values(path, "Market Data", 5, range(3, 10)) == [10.5, 2000, 2000.0, None, 80.0, 125.0, None]   # no P/E: stays blank
    ws = openpyxl.load_workbook(path)["Market Data"]
    assert ws["B3"].value == "Alpha Corp" and ws["A3"].value == "AAA"           # identity columns untouched
    assert ws["C3"].number_format == "\\$#,##0.00" and ws["C3"].fill.fgColor.rgb == "FFDDEBF7"
    assert ws["K3"].value == "=C3*2"                                            # a formula elsewhere survives
    assert "A1:I1" in {str(r) for r in ws.merged_cells.ranges}
    assert openpyxl.load_workbook(path)["Other"]["A1"].value == "untouched"


def test_a_value_the_provider_lacks_is_kept_not_blanked(tmp_path):
    path = build(tmp_path)
    rep = market.refresh(path, snapshot=table(GOOD), earnings=lambda t: None, today=TODAY)
    assert values(path, "Market Data", 4, [5]) == [9999.0]                      # BBB's market cap: provider had none
    assert ("BBB", "market_cap_m") in rep["market"]["kept"]


def test_a_stale_yield_is_cleared_when_the_provider_reports_no_dividend_and_kept_when_it_is_silent(tmp_path):
    rows = {"AAA": snap(dy=None, rate=None, trail=0.0), "BBB": GOOD["BBB"], "CCC": GOOD["CCC"]}
    path = build(tmp_path)                                                       # AAA carries a 1.5 yield
    rep = market.refresh(path, snapshot=table(rows), earnings=lambda t: None, today=TODAY, sleep=lambda s: None)
    assert values(path, "Market Data", 3, [9]) == [None]
    assert rep["market"]["cleared"] == [("AAA", "dividend_yield")] and rep["market"]["cells_changed"] > 0
    rows["AAA"] = snap(dy=None, rate=None, trail=None)
    path2 = build(tmp_path, name="m2.xlsx")
    rep = market.refresh(path2, snapshot=table(rows), earnings=lambda t: None, today=TODAY, sleep=lambda s: None)
    assert values(path2, "Market Data", 3, [9]) == [1.5]
    assert ("AAA", "dividend_yield") in rep["market"]["kept"] and rep["market"]["cleared"] == []


def test_a_failed_ticker_is_retried_then_reported_and_its_row_left_alone(tmp_path):
    path = build(tmp_path)
    calls = {"AAA": 0}

    def flaky(t):
        if t == "AAA":
            calls["AAA"] += 1
            return None if calls["AAA"] == 1 else snap()
        if t == "BBB":
            return {k: None for k in snap()}                                   # what market_snapshot returns on error
        return GOOD["CCC"]

    before = values(path, "Market Data", 4, range(3, 10))
    rep = market.refresh(path, snapshot=flaky, earnings=lambda t: None, today=TODAY, sleep=lambda s: None)
    assert calls["AAA"] == 2 and rep["market"]["failed"] == ["BBB"] and rep["market"]["refreshed"] == 2
    assert values(path, "Market Data", 3, [3]) == [101.5]
    assert values(path, "Market Data", 4, range(3, 10)) == before


def test_only_ticker_rows_are_fetched(tmp_path):
    path = build(tmp_path)
    seen = []
    market.refresh(path, snapshot=lambda t: seen.append(t), earnings=lambda t: seen.append("E" + t),
                   today=TODAY, attempts=1, write=False)
    assert sorted(set(seen)) == ["AAA", "BBB", "CCC", "EAAA", "EBBB", "ECCC"]            # no footer, no label


def test_tickers_filter_limits_both_tabs(tmp_path):
    path = build(tmp_path)
    seen = []
    market.refresh(path, snapshot=lambda t: seen.append(t) or snap(), earnings=lambda t: None, today=TODAY,
                   tickers=["bbb"], write=False)
    assert set(seen) == {"BBB"}


def test_earnings_dates_keep_their_cell_format_and_the_upcoming_list_is_held_names_only(tmp_path):
    path = build(tmp_path)
    dates = {"AAA": D(2026, 10, 9), "BBB": D(2026, 12, 1), "CCC": None}
    rep = market.refresh(path, snapshot=table(GOOD), earnings=table(dates), today=TODAY, sleep=lambda s: None)
    ws = openpyxl.load_workbook(path)["Monitor & Recheck Triggers"]
    assert ws["J3"].value == dt.datetime(2026, 10, 9) and ws["J3"].number_format == "mm/dd/yyyy"
    assert ws["J4"].value == dt.datetime(2026, 12, 1)
    assert ws["J5"].value == dt.datetime(2026, 11, 20)                          # no date from the provider: kept
    em = rep["earnings"]
    assert [(t, n) for t, _o, n in em["updated"]] == [("AAA", D(2026, 10, 9)), ("BBB", D(2026, 12, 1))]
    assert em["no_date"] == ["CCC"]
    assert em["upcoming"] == [("AAA", D(2026, 10, 9))]                          # BBB is later, CCC is not held


def test_a_past_provider_date_does_not_replace_a_future_one(tmp_path):
    path = build(tmp_path)
    rep = market.refresh(path, snapshot=table(GOOD), earnings=table({"BBB": D(2026, 8, 6)}), today=TODAY,
                         sleep=lambda s: None)
    assert values(path, "Monitor & Recheck Triggers", 4, [10]) == [dt.datetime(2026, 10, 10)]
    assert rep["earnings"]["kept"] and rep["earnings"]["kept"][0][0] == "BBB"


def test_dry_run_reports_the_change_and_writes_nothing(tmp_path):
    path = build(tmp_path)
    before = digest(path)
    rep = market.refresh(path, snapshot=table(GOOD), earnings=table({"AAA": D(2026, 10, 9)}), today=TODAY,
                         write=False, sleep=lambda s: None)
    assert digest(path) == before and not rep["wrote"]
    assert rep["market"]["cells_changed"] > 0 and rep["earnings"]["updated"]


def test_a_second_run_over_unchanged_data_does_not_save(tmp_path):
    path = build(tmp_path)
    dates = {"AAA": D(2026, 10, 9)}
    first = market.refresh(path, snapshot=table(GOOD), earnings=table(dates), today=TODAY, sleep=lambda s: None)
    assert first["wrote"]
    mid = digest(path)
    second = market.refresh(path, snapshot=table(GOOD), earnings=table(dates), today=TODAY, sleep=lambda s: None)
    assert not second["wrote"] and second["market"]["cells_changed"] == 0 and not second["earnings"]["updated"]
    assert digest(path) == mid


def test_a_moved_column_is_refused_before_anything_is_written(tmp_path):
    path = build(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb["Market Data"]["C2"] = "Company (long)"                                  # a column was inserted ahead of Price
    wb.save(path)
    before = digest(path)
    with pytest.raises(MarketError, match="column C"):
        market.refresh(path, snapshot=table(GOOD), earnings=lambda t: None, today=TODAY)
    assert digest(path) == before


def test_a_missing_tab_is_refused(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Market Data"
    path = str(tmp_path / "x.xlsx")
    wb.save(path)
    with pytest.raises(MarketError, match="Monitor"):
        market.refresh(path, snapshot=table(GOOD), earnings=lambda t: None, today=TODAY)


def test_a_quote_far_from_the_newest_friday_close_is_reported(tmp_path):
    path = build(tmp_path)
    rep = market.refresh(path, snapshot=table(GOOD), earnings=lambda t: None, today=TODAY, write=False,
                         reference={"AAA": 90.0, "BBB": 51.1})
    warns = rep["market"]["warnings"]
    assert any(w.startswith("AAA:") and "Price History close" in w for w in warns)
    assert not any(w.startswith("BBB:") for w in warns)                         # within 3%


# ------------------------------------------------------------ the reader --

def test_read_market_data_has_no_row_bound_and_ignores_footers(tmp_path):
    """It was hardcoded to rows 3-38 and went stale when four tickers were added (2026-10-04); the audit
    caught it. A row counts if column A holds a ticker-shaped string -- not a footer formula, not a note."""
    from landry import xlsx_io
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    add_market_tabs(wb)
    md = wb["Market Data"]
    md.cell(row=7, column=1).value = None                                        # the fixture's footer formula
    md.cell(row=8, column=1).value = None                                        # ... and its label
    for i in range(60):                                                          # far past the old row-38 bound
        md.cell(row=6 + i, column=1, value=f"T{i:02d}")
        md.cell(row=6 + i, column=3, value=10.0 + i)
    md.cell(row=70, column=1, value="=SUBTOTAL(103,A3:A69)")
    md.cell(row=71, column=1, value="Note: refreshed weekly by landry weekly")
    path = str(tmp_path / "rd.xlsx")
    wb.save(path)
    rows = xlsx_io.read_market_data(path)
    assert [r["ticker"] for r in rows][:4] == ["AAA", "BBB", "CCC", "T00"] and len(rows) == 63
    assert rows[0]["price"] == 100.0 and rows[0]["company"] == "Alpha Corp" and rows[-1]["ticker"] == "T59"


def test_audit_finds_no_hardcoded_bound_left_in_read_market_data():
    import inspect
    import re

    from landry import xlsx_io
    assert not re.search(r"max_row\s*=\s*\d+", inspect.getsource(xlsx_io.read_market_data))
