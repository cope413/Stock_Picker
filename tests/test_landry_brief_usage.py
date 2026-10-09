"""landry brief and landry usage: smoke tests on the live workbook and a synthetic transcript."""
import json
import glob

from landry import brief, usage


def _wb():
    return sorted(glob.glob("LANDRY_SYSTEM_WORKBOOK_25.xlsx"))[0]


def test_brief_is_compact_and_has_sections():
    t = brief.build(_wb())
    assert "JOURNAL:" in t and "OPEN ITEMS:" in t and "PROCESS CHECKLIST:" in t and "GIT:" in t
    assert len(t) < 12000


def test_usage_reports_thinking_and_compactions(tmp_path):
    p = tmp_path / "s.jsonl"
    rows = []
    for k in range(4):
        rows.append({"type": "assistant", "message": {"id": f"m{k}", "usage": {"input_tokens": 1, "cache_read_input_tokens": 1000 * (k + 1),
                     "cache_creation_input_tokens": 10, "output_tokens": 500},
                     "content": [{"type": "text", "text": "x" * 36}]}})
        rows.append({"type": "user", "message": {"content": [{"type": "tool_result", "content": "y" * 360}]}})
        if k == 1:
            rows.append({"type": "system", "subtype": "compact_boundary"})
    p.write_text("\n".join(json.dumps(r) for r in rows))
    calls = usage.parse(str(p))
    assert len(calls) == 4 and calls[-1]["seg"] == 1
    r = usage.report(calls)
    assert "compactions 1" in r and "thinking" in r


def test_rule12_arithmetic_matches_the_rule():
    from landry import rule12
    hi = rule12.screen(100.0, 0.0, 0.15)          # 15% growth, no dividend: Base ~15%
    assert hi["passes"] and abs(hi["base"] - 0.15) < 0.005 and hi["bear"] > 0
    lo = rule12.screen(100.0, 2.0, 0.05)          # 5% growth + 2% yield cannot reach 10%
    assert not lo["passes"] and 0.06 < lo["base"] < 0.08
    assert 0.07 < lo["need"] < 0.10               # growth needed with a 2% yield


def test_open_items_update_on_a_scratch_copy(tmp_path):
    import shutil
    from landry import open_items
    import openpyxl
    p = tmp_path / "wb.xlsx"
    shutil.copy(_wb(), p)
    before = openpyxl.load_workbook(p)["Open Items"]
    r = next(i for i in range(5, before.max_row + 1) if str(before.cell(i, 1).value) == "4")
    old = before.cell(r, 6).value
    msg = open_items.update(str(p), 4, "test note", recalc=False)
    ws = openpyxl.load_workbook(p)["Open Items"]
    assert "noted" in msg and str(ws.cell(r, 6).value).endswith("test note")
    assert str(ws.cell(r, 3).value or "").upper() != "Y"
    assert (old or "") in str(ws.cell(r, 6).value)
    try:
        open_items.update(str(p), 99999, "x", recalc=False)
        assert False
    except open_items.OpenItemsError:
        pass
