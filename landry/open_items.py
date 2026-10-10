"""``landry open-items``: add a note to, or close, an Open Items row (the tab is a Table edited in place).

Replaces the throwaway openpyxl scripts that were rewritten for every Open Items update (ten of them on 10/8).
A note is APPENDED to column F (never replaces it) and the row height is re-fit; ``close`` also sets Done = Y and
the resolved date. Refuses if Excel has the workbook open or the row is missing. ``add`` appends a NEW numbered row
directly below the last one (no row insertion), grows ``OpenItemsTable`` to cover it and restates the Schema Reference
line. Recalc runs after the save; run ``landry audit`` before a commit."""
import datetime as dt
import copy
import math
import re
from typing import Optional

import openpyxl

from landry import ledger, xlsx_recalc

SHEET = "Open Items"
TABLE = "OpenItemsTable"


class OpenItemsError(Exception):
    pass


def _row_of(ws, n: int) -> int:
    for r in range(5, ws.max_row + 1):
        if str(ws.cell(r, 1).value) == str(n):
            return r
    raise OpenItemsError(f"no Open Items row numbered #{n}")


def _fit(ws, r: int) -> None:
    item = len(str(ws.cell(r, 2).value or ""))
    notes = len(str(ws.cell(r, 6).value or ""))
    ws.row_dimensions[r].height = min(400.0, 13.5 * max(math.ceil(item / 92), math.ceil(notes / 82), 1) + 3)


def update(path: str, n: int, note: str, close: bool = False, when: Optional[dt.date] = None,
           recalc: bool = True) -> str:
    if ledger._excel_has_open(path):
        raise OpenItemsError("Excel has the workbook open; close it and retry")
    note = " ".join(note.split())
    if not note and not close:
        raise OpenItemsError("nothing to do: give a note and/or --close")
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[SHEET]
        r = _row_of(ws, n)
        if close and str(ws.cell(r, 3).value or "").upper() == "Y":
            raise OpenItemsError(f"#{n} is already Done")
        if note:
            old = str(ws.cell(r, 6).value or "").rstrip()
            ws.cell(r, 6).value = (old + " " + note).strip()
        if close:
            ws.cell(r, 3).value = "Y"
            ws.cell(r, 5).value = dt.datetime.combine(when or dt.date.today(), dt.time())
        _fit(ws, r)
        wb.save(path)
    finally:
        wb.close()
    if recalc:
        xlsx_recalc.recalc(path)
    return f"#{n} {'closed' if close else 'noted'} (row {r})"


def add(path: str, item: str, note: str = "", when: Optional[dt.date] = None, recalc: bool = True) -> str:
    """Append a new Open Items row: next number, Date Raised = ``when`` (today), styled like the row above it."""
    if ledger._excel_has_open(path):
        raise OpenItemsError("Excel has the workbook open; close it and retry")
    item, note = " ".join(item.split()), " ".join(note.split())
    if not item:
        raise OpenItemsError("an Open Item needs its text")
    when = when or dt.date.today()
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[SHEET]
        used = [r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value is not None]
        if not used:
            raise OpenItemsError("no existing Open Items rows to continue from")
        last = used[-1]
        r, n = last + 1, max(int(ws.cell(x, 1).value) for x in used) + 1
        if any(ws.cell(r, c).value is not None for c in range(1, 7)):
            raise OpenItemsError(f"row {r} below the last item is not empty")
        tbl = ws.tables[TABLE]
        if tbl.ref != f"A4:F{last}":
            raise OpenItemsError(f"{TABLE} covers {tbl.ref}, expected A4:F{last}; fix that first")
        for c in range(1, 7):
            ws.cell(r, c)._style = copy.copy(ws.cell(last, c)._style)
        first = ws.cell(last, 1).value
        ws.cell(r, 1).value = n if isinstance(first, int) else str(n)
        ws.cell(r, 2).value = item
        ws.cell(r, 4).value = dt.datetime.combine(when, dt.time())
        ws.cell(r, 6).value = note or None
        _fit(ws, r)
        tbl.ref = f"A4:F{r}"
        if "Schema Reference" in wb.sheetnames:
            for row in wb["Schema Reference"].iter_rows():
                for cell in row:
                    if cell.value == TABLE:
                        nxt = cell.offset(0, 1)
                        nxt.value = re.sub(r"^A4:F\d+ as of \S+ \(\d+ items\)",
                                           f"A4:F{r} as of {when.isoformat()} ({len(used) + 1} items)", str(nxt.value or ""))
        wb.save(path)
    finally:
        wb.close()
    if recalc:
        xlsx_recalc.recalc(path)
    return f"#{n} added (row {r})"
