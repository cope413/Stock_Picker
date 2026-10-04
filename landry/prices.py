"""Weekly closes for the Price History tab, and everything that has to move with them.

Why this exists (2026-10-04): the tab used to be filled by ``landry export``'s Price History
block, which wipes and rewrites A3:Q162 from the FIRST 16 TICKERS ALPHABETICALLY. Around 10/1
that replaced the documented layout (19 tickers in B-V, free slots at P and U, 19 charts keyed
to those columns) with an alphabetical 16, silently dropped VRT/VRTX/HELO/JEPQ-era columns' place
in line, left the old R-V columns behind five weeks out of alignment, and left eight held
positions with no column at all -- so Rule 38's correlation check was partly blind.

This module only ever does four things, none of which reorder or drop a column:

* ``status``      read-only health check (duplicate/blank headers, holdings with no column,
                  weeks behind, last-row gaps);
* ``append_weeks``  add the completed Friday(s) after the last row, extend Returns (Calc)
                  one row ahead (it does not auto-extend), move the footer, push the charts
                  below the data;
* ``rebuild``     rewrite every price under the header that is already there from one
                  consistent dividend-adjusted pull (the quarterly restatement);
* ``add_ticker``  give a new holding a column, extending Returns (Calc) and the Correlation
                  Matrix to match.

Prices are dividend-adjusted weekly closes (``data_auto.weekly_closes``), so week-over-week
changes are total returns, as Rule 38 asks. A weekly append is raw = adjusted at the moment of
the pull; yfinance later restates earlier closes for each dividend, so ``rebuild`` each quarter.
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from typing import Callable, Dict, List, Optional, Sequence, Tuple

PH_SHEET = "Price History"
RET_SHEET = "Returns (Calc)"
CM_SHEET = "Correlation Matrix"

HEADER_ROW = 2
FIRST_ROW = 3
LAST_PH_COL_ORIGINAL = 22        # V: Price History columns B..V map to Returns (Calc) columns B..V
PH_FIRST_EXTRA_COL = 23          # W
RET_EXTRA_FIRST_COL = 42         # AP: Returns (Calc)'s W..AO hold the 19 chart "trend" columns
MATRIX_HEADER_ROW = 5            # Correlation Matrix: header row; ticker k's row is 4 + (its Price History column)
MATRIX_FIRST_ROW = 6
COUNT_COL = 33                   # AG: "Positions >0.70 with"
MAX_WEEKLY_MOVE = 0.35           # a new close further than this from last week's is refused unless forced
MAX_PH_COL = 32                  # AF: the matrix's buffer columns W..AF are the room for new tickers

CLOSE_AFTER_ET = dt.time(16, 15)


class PricesError(RuntimeError):
    pass


def returns_col(ph_col: int) -> int:
    """Returns (Calc) column that holds the week-over-week return for Price History column ``ph_col``."""
    if ph_col <= LAST_PH_COL_ORIGINAL:
        return ph_col
    return RET_EXTRA_FIRST_COL + (ph_col - PH_FIRST_EXTRA_COL)


def _L(col: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(col)


# ------------------------------------------------------------------ layout --

@dataclass
class Layout:
    tickers: Dict[int, str]                    # column -> ticker, header row, left to right
    date_rows: List[Tuple[int, dt.datetime]]   # (row, week-ending date)
    gaps: List[int] = field(default_factory=list)       # blank header slots between used columns

    @property
    def last_row(self) -> int:
        return self.date_rows[-1][0]

    @property
    def last_date(self) -> dt.datetime:
        return self.date_rows[-1][1]


def read_layout(ws) -> Layout:
    tickers: Dict[int, str] = {}
    for c in range(2, ws.max_column + 1):
        v = ws.cell(row=HEADER_ROW, column=c).value
        if isinstance(v, str) and v.strip():
            tickers[c] = v.strip().upper()
    seen: Dict[str, int] = {}
    for c, t in tickers.items():
        if t in seen:
            raise PricesError(f"Price History header has {t} twice (columns {_L(seen[t])} and {_L(c)}) -- "
                              f"a duplicate column is a second copy of one instrument, not a second position")
        seen[t] = c
    date_rows = [(r, ws.cell(row=r, column=1).value) for r in range(FIRST_ROW, ws.max_row + 1)
                 if isinstance(ws.cell(row=r, column=1).value, dt.datetime)]
    if not date_rows:
        raise PricesError("Price History has no dated rows")
    for (r0, d0), (r1, d1) in zip(date_rows, date_rows[1:]):
        if d1 <= d0:
            raise PricesError(f"Price History rows {r0} and {r1} are not in chronological order "
                              f"({d0:%Y-%m-%d} then {d1:%Y-%m-%d}); Returns (Calc) compares row n with row n-1")
    cols = sorted(tickers)
    gaps = [c for c in range(cols[0], cols[-1] + 1) if c not in tickers] if cols else []
    return Layout(tickers, date_rows, gaps)


# ------------------------------------------------------------------- dates --

def last_completed_friday(now: Optional[dt.datetime] = None) -> dt.date:
    """The most recent Friday whose close is in: today if it is Friday and after ~4:15pm ET, else the
    Friday before. A week that has not closed must never be written as a weekly close."""
    from zoneinfo import ZoneInfo
    et_zone = ZoneInfo("America/New_York")
    now = now or dt.datetime.now(et_zone)
    et = now.astimezone(et_zone) if now.tzinfo else now.replace(tzinfo=et_zone)
    day = et.date()
    fri = day - dt.timedelta(days=(day.weekday() - 4) % 7)
    if fri == day and et.time() < CLOSE_AFTER_ET:
        fri -= dt.timedelta(days=7)
    return fri


def missing_fridays(last_date: dt.datetime, now: Optional[dt.datetime] = None) -> List[dt.date]:
    out, d, stop = [], last_date.date() + dt.timedelta(days=7), last_completed_friday(now)
    while d <= stop:
        out.append(d)
        d += dt.timedelta(days=7)
    return out


# ------------------------------------------------------------------- fetch --

def fetch_weekly(tickers: Sequence[str], years: int = 4, refresh: bool = True):
    """Dividend-adjusted Friday closes, one column per ticker (``data_auto.weekly_closes``)."""
    from landry.data_auto import fetch_daily, weekly_closes
    return weekly_closes(fetch_daily(list(tickers), years=years, refresh=refresh, min_bars=20))


def _val(wk, date, ticker) -> Optional[float]:
    import pandas as pd
    ts = pd.Timestamp(date)
    if ticker not in wk.columns or ts not in wk.index:
        return None
    v = wk.at[ts, ticker]
    return None if v != v else round(float(v), 2)


# ------------------------------------------------------------------ status --

def status(path: str, now: Optional[dt.datetime] = None) -> dict:
    """Read-only health check. Never raises on a layout problem -- it reports it."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[PH_SHEET]
    problems: List[str] = []
    try:
        layout = read_layout(ws)
    except PricesError as e:
        return {"ok": False, "problems": [str(e)]}
    held: List[str] = []
    try:
        from landry.xlsx_io import read_positions
        cash = {"FZDXX", "FZFXX", "VMFXX", "QACDS"}
        held = sorted({p.ticker for p in read_positions(path)
                       if getattr(p, "quantity", 0) and p.ticker not in cash})
    except Exception as e:                      # status must still say something without positions
        problems.append(f"could not read Current Positions: {e}")
    header = set(layout.tickers.values())
    missing = sorted(set(held) - header)
    if missing:
        problems.append(f"held but with no Price History column: {', '.join(missing)}")
    unheld = sorted(header - set(held)) if held else []
    behind = missing_fridays(layout.last_date, now)
    last = layout.last_row
    blanks = [t for c, t in layout.tickers.items() if ws.cell(row=last, column=c).value is None]
    if blanks:
        problems.append(f"no value in the last row ({layout.last_date:%Y-%m-%d}) for: {', '.join(blanks)}")
    fridays_off = [d for _, d in layout.date_rows if d.weekday() != 4]
    if fridays_off:
        problems.append(f"{len(fridays_off)} row date(s) are not Fridays, first {fridays_off[0]:%Y-%m-%d}")
    return {"ok": not problems, "problems": problems, "last_date": layout.last_date.date(),
            "last_row": last, "weeks": len(layout.date_rows), "tickers": list(layout.tickers.values()),
            "gaps": [_L(c) for c in layout.gaps], "weeks_behind": len(behind), "missing_fridays": behind,
            "held_not_in_header": missing, "in_header_not_held": unheld}


# ------------------------------------------------------------------ helpers --

def _copy_style(src, dst) -> None:
    from copy import copy
    dst._style = copy(src._style)


def _ret_last_formula_row(ret) -> int:
    last = FIRST_ROW - 1
    for r in range(FIRST_ROW, ret.max_row + 1):
        v = ret.cell(row=r, column=2).value
        if isinstance(v, str) and v.startswith("="):
            last = r
    return last


def _extend_returns(ret, new_last_row: int) -> int:
    """Returns (Calc) runs ONE ROW AHEAD of Price History: row n holds the return of Price History
    row n+1 over row n. Translate the last formula row down so formulas reach ``new_last_row``."""
    from openpyxl.formula.translate import Translator
    src_row = _ret_last_formula_row(ret)
    added = 0
    for r in range(src_row + 1, new_last_row + 1):
        for c in range(1, ret.max_column + 1):
            cell = ret.cell(row=src_row, column=c)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                dst = ret.cell(row=r, column=c)
                above = ret.cell(row=src_row - 1, column=c).value if src_row > FIRST_ROW else None
                # Row-relative formulas (the returns, the date) differ from row to row and must be translated. The
                # chart "trend" columns are the SAME formula in every row, anchored to Price History's first row
                # ('Price History'!B3 ... ROW()-3): translating those would slide the anchor down a row.
                if above == cell.value:
                    dst.value = cell.value
                else:
                    dst.value = Translator(cell.value, origin=cell.coordinate).translate_formula(dst.coordinate)
                _copy_style(cell, dst)
        added += 1
    return added


def _shift_charts(ws, rows: int) -> int:
    """Keep the charts below the data: move every chart on the sheet down ``rows`` rows."""
    moved = 0
    for ch in getattr(ws, "_charts", []):
        a = ch.anchor
        if isinstance(a, str):
            m = re.match(r"([A-Z]+)(\d+)$", a)
            if m:
                ch.anchor = f"{m.group(1)}{int(m.group(2)) + rows}"
                moved += 1
        else:
            a._from.row += rows
            if getattr(a, "to", None) is not None:
                a.to.row += rows
            moved += 1
    return moved


def _move_footer(ws, old_last: int, new_last: int) -> Optional[int]:
    """The 'NN Weeks' footer sits two blank rows under the data; keep that gap."""
    footer = None
    for r in range(old_last + 1, old_last + 12):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.upper().startswith("=COUNT(") and str(ws.cell(row=r, column=2).value).strip() == "Weeks":
            footer = r
            break
    if footer is None:
        return None
    new_row = new_last + (footer - old_last)
    if new_row == footer:
        return footer
    blank = ws.cell(row=old_last + 1, column=1)
    for c in (1, 2):
        src, dst = ws.cell(row=footer, column=c), ws.cell(row=new_row, column=c)
        dst.value = src.value
        _copy_style(src, dst)
        src.value = None
        _copy_style(ws.cell(row=footer - 1, column=c) if footer - 1 > old_last else blank, src)
    ws.cell(row=new_row, column=1).value = f"=COUNT(A{FIRST_ROW}:A{new_last + 1})"
    return new_row


# ------------------------------------------------------------------- append --

def append_weeks(path: str, *, fetch: Callable = fetch_weekly, now: Optional[dt.datetime] = None,
                 write: bool = True, force: bool = False) -> dict:
    """Append every completed Friday after the last row. Never rewrites an existing row."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ph, ret = wb[PH_SHEET], wb[RET_SHEET]
    layout = read_layout(ph)
    targets = missing_fridays(layout.last_date, now)
    report = {"appended": [], "warnings": [], "up_to_date": not targets, "last_date_before": layout.last_date.date()}
    if not targets:
        return report
    wk = fetch(list(layout.tickers.values()))
    new_vals: Dict[dt.date, Dict[int, Optional[float]]] = {}
    for d in targets:
        row = {c: _val(wk, d, t) for c, t in layout.tickers.items()}
        blanks = [layout.tickers[c] for c, v in row.items() if v is None]
        if len(blanks) > max(1, len(row) // 5):
            raise PricesError(f"no close for {', '.join(blanks)} in the week ending {d:%Y-%m-%d} -- "
                              f"the data provider has not caught up (or the week is not complete); nothing written")
        if blanks:
            report["warnings"].append(f"{d:%Y-%m-%d}: no close for {', '.join(blanks)} -- left blank")
        new_vals[d] = row
    prev = {c: ph.cell(row=layout.last_row, column=c).value for c in layout.tickers}
    for d in targets:
        for c, v in new_vals[d].items():
            p = prev.get(c)
            if v is not None and isinstance(p, (int, float)) and p and abs(v / p - 1) > MAX_WEEKLY_MOVE and not force:
                raise PricesError(f"{layout.tickers[c]} moved {v / p - 1:+.0%} ({p} -> {v}) in the week ending "
                                  f"{d:%Y-%m-%d} -- more than {MAX_WEEKLY_MOVE:.0%}; check the column, or pass force")
            if v is not None:
                prev[c] = v
    if not write:
        report["appended"] = [d.isoformat() for d in targets]
        return report
    old_last = layout.last_row
    new_last = old_last + len(targets)
    max_col = max(layout.tickers)
    report["footer_row"] = _move_footer(ph, old_last, new_last)     # first: the new rows may land on the old footer cell
    for i, d in enumerate(targets):
        r = old_last + 1 + i
        for c in range(1, max_col + 1):
            _copy_style(ph.cell(row=old_last, column=c), ph.cell(row=r, column=c))
        ph.cell(row=r, column=1).value = dt.datetime(d.year, d.month, d.day)
        for c, v in new_vals[d].items():
            ph.cell(row=r, column=c).value = v
        report["appended"].append(d.isoformat())
    added = _extend_returns(ret, new_last)
    report["returns_rows_added"] = added
    report["charts_moved"] = _shift_charts(ret, added) if added else 0
    wb.save(path)
    return report


# ------------------------------------------------------------------ rebuild --

def rebuild(path: str, *, fetch: Callable = fetch_weekly, write: bool = True) -> dict:
    """Rewrite every price under the existing header and dates from one adjusted pull."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ph = wb[PH_SHEET]
    layout = read_layout(ph)
    wk = fetch(list(layout.tickers.values()))
    missing = [t for t in layout.tickers.values() if t not in wk.columns]
    if missing:
        raise PricesError(f"the data provider returned nothing for {', '.join(missing)}; nothing written")
    for c, t in layout.tickers.items():
        had_none = sum(1 for r, _ in layout.date_rows if ph.cell(row=r, column=c).value is None)
        now_none = sum(1 for _, d in layout.date_rows if _val(wk, d, t) is None)
        if now_none > had_none + 3:
            raise PricesError(f"the data provider is missing {now_none - had_none} weeks of {t} that the tab has; "
                              f"refusing to erase them -- nothing written")
    changed, biggest, filled = 0, 0.0, 0
    for r, d in layout.date_rows:
        for c, t in layout.tickers.items():
            new = _val(wk, d, t)
            cell = ph.cell(row=r, column=c)
            old = cell.value
            if new is None and old is None:
                continue
            if isinstance(old, (int, float)) and new is not None and old:
                biggest = max(biggest, abs(new / old - 1))
            if new != old:
                changed += 1
                filled += 1 if old is None else 0
                if write:
                    cell.value = new
    if write:
        wb.save(path)
    return {"cells_changed": changed, "cells_newly_filled": filled, "largest_relative_change": round(biggest, 4),
            "weeks": len(layout.date_rows), "tickers": len(layout.tickers)}


# ------------------------------------------------------------------- matrix --

def _ensure_returns_column(ret, ph_col: int, last_formula_row: int) -> int:
    rc = returns_col(ph_col)
    if ph_col <= LAST_PH_COL_ORIGINAL:
        return rc
    L_new, L_src = _L(ph_col), _L(LAST_PH_COL_ORIGINAL)
    ret.cell(row=2, column=rc).value = f"='Price History'!{L_new}2"
    _copy_style(ret.cell(row=2, column=LAST_PH_COL_ORIGINAL), ret.cell(row=2, column=rc))
    for r in range(FIRST_ROW, last_formula_row + 1):
        src = ret.cell(row=r, column=LAST_PH_COL_ORIGINAL)
        if isinstance(src.value, str):
            dst = ret.cell(row=r, column=rc)
            dst.value = re.sub(rf"'Price History'!{L_src}(\d+)", rf"'Price History'!{L_new}\1", src.value)
            _copy_style(src, dst)
    ret.column_dimensions[_L(rc)].width = ret.column_dimensions[L_src].width or 5.7
    return rc


def _ensure_matrix(cm, ret, up_to_ph_col: int) -> dict:
    """Make the Correlation Matrix cover Price History columns B..``up_to_ph_col``: header, ticker row,
    every pairwise CORREL, the >0.70 count column, the Rule 38 tally and the colour rules."""
    from openpyxl.formatting.formatting import ConditionalFormattingList
    last_row = 4 + up_to_ph_col
    last_col_letter = _L(up_to_ph_col)
    interior, label, head, count = cm["C7"], cm["A26"], cm.cell(row=MATRIX_HEADER_ROW, column=22), cm.cell(row=26, column=COUNT_COL)
    made = 0
    for k in range(2, up_to_ph_col + 1):
        rk = _L(returns_col(k))
        row = 4 + k
        if cm.cell(row=row, column=1).value is None:
            cm.cell(row=row, column=1).value = f"='Returns (Calc)'!{rk}2"
            _copy_style(label, cm.cell(row=row, column=1))
            made += 1
        if cm.cell(row=MATRIX_HEADER_ROW, column=k).value is None:
            cm.cell(row=MATRIX_HEADER_ROW, column=k).value = f"='Returns (Calc)'!{rk}2"
            _copy_style(head, cm.cell(row=MATRIX_HEADER_ROW, column=k))
    for i in range(2, up_to_ph_col + 1):
        ri = _L(returns_col(i))
        for j in range(2, up_to_ph_col + 1):
            cell = cm.cell(row=4 + i, column=j)
            if cell.value is not None:
                continue
            rj = _L(returns_col(j))
            cell.value = 1 if i == j else (f"=IFERROR(ROUND(CORREL('Returns (Calc)'!${ri}$3:${ri}$500,"
                                           f"'Returns (Calc)'!${rj}$3:${rj}$500),2),\"\")")
            _copy_style(interior, cell)
    for row in range(MATRIX_FIRST_ROW, last_row + 1):
        c = cm.cell(row=row, column=COUNT_COL)
        c.value = f'=COUNTIF(B{row}:{last_col_letter}{row},">0.7")-1'
        if row > 26:
            _copy_style(count, c)
    cm["J2"].value = f'=COUNTIF(AG{MATRIX_FIRST_ROW}:AG{last_row},">1")'
    for col in range(23, up_to_ph_col + 1):
        cm.column_dimensions[_L(col)].width = cm.column_dimensions["C"].width or 8.7
    new_cf = ConditionalFormattingList()
    for cf in cm.conditional_formatting:
        rng = str(cf.sqref)
        if rng.startswith("B6:"):
            rng = f"B6:{last_col_letter}{last_row}"
        elif rng.startswith("AG6:"):
            rng = f"AG6:AG{last_row}"
        for rule in cf.rules:
            new_cf.add(rng, rule)
    cm.conditional_formatting = new_cf
    return {"matrix_rows": last_row - MATRIX_FIRST_ROW + 1, "new_labels": made}


# --------------------------------------------------------------- add ticker --

def add_ticker(path: str, ticker: str, *, col: Optional[int] = None, fetch: Callable = fetch_weekly,
               write: bool = True) -> dict:
    """Give ``ticker`` a Price History column (a free slot in B..V if there is one, else the next column),
    fill its history, and extend Returns (Calc) and the Correlation Matrix to cover it."""
    import openpyxl
    ticker = ticker.strip().upper()
    wb = openpyxl.load_workbook(path)
    ph, ret, cm = wb[PH_SHEET], wb[RET_SHEET], wb[CM_SHEET]
    layout = read_layout(ph)
    if ticker in layout.tickers.values():
        raise PricesError(f"{ticker} already has a column ({_L(next(c for c, t in layout.tickers.items() if t == ticker))})")
    if col is None:
        free = [c for c in layout.gaps if c <= LAST_PH_COL_ORIGINAL]
        col = free[0] if free else max(max(layout.tickers) + 1, PH_FIRST_EXTRA_COL)
    if col < 2 or col > MAX_PH_COL:
        raise PricesError(f"column {col} is outside the Correlation Matrix's room (B..{_L(MAX_PH_COL)})")
    if ph.cell(row=HEADER_ROW, column=col).value not in (None, ""):
        raise PricesError(f"column {_L(col)} is not free")
    wk = fetch([ticker]) if write else None
    info = {"ticker": ticker, "column": _L(col), "reused_slot": col <= LAST_PH_COL_ORIGINAL}
    if not write:
        return info
    if ticker not in wk.columns:
        raise PricesError(f"the data provider returned nothing for {ticker}; nothing written")
    ref = ph.cell(row=HEADER_ROW, column=max(layout.tickers))
    _copy_style(ref, ph.cell(row=HEADER_ROW, column=col))
    ph.cell(row=HEADER_ROW, column=col).value = ticker
    src_col = max(layout.tickers)
    ph.column_dimensions[_L(col)].width = ph.column_dimensions[_L(src_col)].width or 6.7
    filled = 0
    for r, d in layout.date_rows:
        cell = ph.cell(row=r, column=col)
        _copy_style(ph.cell(row=r, column=src_col), cell)
        v = _val(wk, d, ticker)
        cell.value = v
        filled += v is not None
    info["weeks_filled"] = filled
    last_formula_row = _ret_last_formula_row(ret)
    _ensure_returns_column(ret, col, last_formula_row)
    info.update(_ensure_matrix(cm, ret, max(col, max(layout.tickers), LAST_PH_COL_ORIGINAL)))
    wb.save(path)
    return info
