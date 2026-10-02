"""Phase C of LANDRY_DATABASE_DESIGN.md, rolled out one tab at a time:
regenerate a workbook tab's data region from ``landry.db``.

Two tabs so far: the Journal (freeform text) and the Portfolio Drawdown Log
(three inputs plus five formula columns). What a generator owns, and what it
deliberately leaves alone:

  owns    the Table's data rows -- their values (and, for the Drawdown Log, their
          formulas), their cell formatting (so a hand-toggled wrap or font can't
          leak into a commit: regenerating restores it), their row heights --
          and, when the log outgrows its pre-formatted rows, everything keyed to
          the Table's last row: Table ref / filter, print area, conditional-format
          ranges, Schema Reference's row-range note, and other tabs' formulas that
          read a fixed range of it (Action Items reads the Drawdown Log's rows 3:42).
  leaves  the title row, header row, column widths, sheet view, page setup and
          footer, and every other tab. Other tabs reference these by position
          (Process Checklist: ``INDEX(Journal!$A:$A, MATCH(label, Journal!$B:$B,
          0))``), so each tab's layout is a contract this module must not disturb.

Journal order is the DB's id order, i.e. write order -- never date order:
future-dated recurring-event placeholders would otherwise bury the real recent
entries (see CLAUDE.md). The Drawdown Log is the opposite: DATE order, because
its running peak is a chronological chain, so a backfilled date lands in place.

The Drawdown Log writes the SAME formulas the tab already has (row-relative)
rather than static values: the DB stores only its inputs, nothing is
re-implemented in Python to drift from the sheet, and static values wait for the
whole-workbook cutover. The consequence is that the Part 7 bands (-10/-20/-30%),
their labels, cash floors and new-position rules now live in this file -- change
them here, not by hand in the sheet, or the next regeneration reverts the edit.
``drawdown.py`` is NOT the source: it words things differently and treats the
exact -10/-20/-30% boundaries the other way (``<`` where the sheet has ``>=``).

Journal row heights are derived layout, not data, so they are never stored:
data rows are written without a height and the mandatory recalc pass
(LibreOffice) fits them and flags them auto-height, which also keeps Excel's
wrap-off/wrap-on trick for a compact view working. The one exception is a note
that needs more than Excel's 409.5pt row ceiling: after the first pass those rows
(and only those, measured, not guessed) are pinned to the ceiling and recalculated
once more, because LibreOffice would otherwise write a height Excel has to clamp.

    python -m landry.generate journal|drawdown|all --workbook W.xlsx --db landry.db [--out OUT.xlsx] [--no-recalc]
    python -m landry.generate verify [--tab journal|drawdown] --workbook W.xlsx --db landry.db
"""

from __future__ import annotations

import argparse
import contextlib
import datetime
import os
import re
import sqlite3
import sys
from collections import Counter
from typing import Callable, Dict, List, NamedTuple, Optional, Sequence, Tuple

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries

from landry import models

HEADER_ROW = 2
FIRST_ROW = 3
MAX_ROW_PT = 409.5        # Excel's hard row-height ceiling

_EDGE = Side(style="thin", color="FF999999")
_BORDER = Border(left=_EDGE, right=_EDGE, top=_EDGE, bottom=_EDGE)


class GenerateError(RuntimeError):
    """The workbook or database isn't in a shape generation can trust."""


# ----------------------------------------------------------- shared helpers --

def _table_layout(ws, table_name: str, headers: Tuple[str, ...]):
    """Validate a tab's skeleton; return (Table, last Table row)."""
    if table_name not in ws.tables:
        raise GenerateError(f"'{ws.title}' has no table named {table_name}")
    table = ws.tables[table_name]         # bracket access: .items() yields ref strings
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    if (min_col, min_row, max_col) != (1, HEADER_ROW, len(headers)):
        raise GenerateError(
            f"{table_name} ref {table.ref} is not A{HEADER_ROW}:"
            f"{get_column_letter(len(headers))}<n>; other tabs reference this "
            f"layout by position, so it can't shift")
    found = tuple(ws.cell(row=HEADER_ROW, column=c).value
                  for c in range(1, len(headers) + 1))
    if found != headers:
        raise GenerateError(f"{ws.title} headers are {found}, expected {headers}")
    return table, max_row


def _put(cell, value) -> None:
    # Explicit attribute assignment: ws.cell(row, col, value=None) is a no-op on
    # a cell that already has content, which would leave stale rows behind.
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"        # a note that begins with '=' is text, not a formula


def _refuse_to_blank(ws, last_row: int, entries: list, db_table: str) -> None:
    """An empty DB table against a tab that has entries means the wrong
    database, not an instruction to erase them."""
    if not entries and any(ws.cell(row=r, column=1).value is not None
                           for r in range(FIRST_ROW, last_row + 1)):
        raise GenerateError(
            f"the database's {db_table} table is empty but '{ws.title}' has "
            f"entries -- wrong database? refusing to blank the log")


def _grow_table(ws, table, last_col: str, old_last: int, new_last: int) -> None:
    table.ref = f"A{HEADER_ROW}:{last_col}{new_last}"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref
    area = ws.print_area            # e.g. "'Journal'!$A$1:$C$302"
    if area and area.endswith(f"${old_last}"):
        ws.print_area = f"A1:{last_col}{new_last}"


def _update_schema_reference(wb, sheet: str, new_last: int) -> bool:
    """Keep Schema Reference's row-range note ('A2:C302') for ``sheet`` in
    step with its Table when it grows -- ``landry audit`` flags it as stale
    otherwise. Same 4-row block layout the audit walks: tab name in column A at
    row r, the Table row at r+2 with its range text in column C. Best effort: if
    the block isn't there or doesn't parse, leave it for the audit to report."""
    if "Schema Reference" not in wb.sheetnames:
        return False
    ws = wb["Schema Reference"]
    for r in range(4, ws.max_row + 1, 4):
        if ws.cell(row=r, column=1).value == sheet:
            cell = ws.cell(row=r + 2, column=3)
            m = re.search(r"\b[A-Z]{1,3}\d+:[A-Z]{1,3}(\d+)\b", str(cell.value or ""))
            if not m:
                return False
            cell.value = cell.value[:m.start(1)] + str(new_last) + cell.value[m.end(1):]
            return True
    return False


def _extend_conditional_formatting(ws, old_last: int, new_last: int) -> None:
    """Move any conditional-format range that ends at the old last row."""
    end = re.compile(r"^([A-Z]+\d+:[A-Z]+)(\d+)$")
    rebuilt = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        rng = str(cf.sqref)
        m = end.match(rng)
        if m and int(m.group(2)) == old_last:
            rng = f"{m.group(1)}{new_last}"
        for rule in cf.rules:
            rebuilt.add(rng, rule)
    ws.conditional_formatting = rebuilt


def _extend_cross_sheet_references(wb, sheet: str, old_last: int, new_last: int) -> int:
    """Other tabs' formulas that read a fixed range of ``sheet`` ending at its
    old last row (``'Sheet'!E3:E42``) must follow the Table, or they silently
    stop seeing the new rows -- the cross-tab-drift bug class. Returns how many
    cells were rewritten."""
    ref = re.compile(r"('" + re.escape(sheet) + r"'!)(\$?[A-Z]{1,3}\$?\d+):(\$?[A-Z]{1,3}\$?)(\d+)")

    def bump(m):
        row = str(new_last) if int(m.group(4)) == old_last else m.group(4)
        return f"{m.group(1)}{m.group(2)}:{m.group(3)}{row}"

    changed = 0
    for other in wb.worksheets:
        if other.title == sheet:
            continue
        for row in other.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and sheet in v:
                    new = ref.sub(bump, v)
                    if new != v:
                        c.value = new
                        changed += 1
    return changed


def _style_cell(cell, style) -> None:
    font, alignment, number_format, fill, border = style
    cell.font, cell.alignment, cell.number_format = font, alignment, number_format
    cell.fill, cell.border = fill, border


# ------------------------------------------------------------------ Journal --

JOURNAL_SHEET = "Journal"
JOURNAL_TABLE = "JournalTable"
JOURNAL_HEADERS = ("Date", "Ticker", "Notes")
BLANK_ROW_PT = 14.25      # height of the pre-formatted capacity rows below the last entry

# Canonical look of a Journal data row. Mirrors the live tab as of 2026-10-02
# (checked cell-by-cell by ``verify``) with one deliberate change: column B wraps
# on every row. It was on for 35 of 300 rows, mostly 12-53 -- where the long
# labels were -- so later rows (e.g. a 54-character ticker list) were clipped.
_NAVY = "FF002060"
_JOURNAL_FILL = PatternFill(fill_type="solid", fgColor="FFECF4FA", bgColor="FFEBF1DE")
# per column A, B, C: (font size, horizontal, vertical, wrap, number format)
_JOURNAL_COLUMNS = (
    (10, "center", "center", False, "mm/dd/yyyy"),
    (11, "center", "center", True, "General"),
    (10, "left", "top", True, "General"),
)
_JOURNAL_STYLES = tuple(
    (Font(name="Arial Narrow", size=size, color=_NAVY),
     Alignment(horizontal=h, vertical=v, wrap_text=wrap), fmt,
     _JOURNAL_FILL, _BORDER)
    for size, h, v, wrap, fmt in _JOURNAL_COLUMNS)


def _apply_journal_heights(ws, rows: int, last_row: int, pinned=()) -> None:
    """The Journal row-height policy. Entry rows get no stored height (the
    recalc pass fits them and flags them auto-height) unless pinned to Excel's
    ceiling; the pre-formatted capacity rows below them get the default."""
    for r in range(FIRST_ROW, last_row + 1):
        if r - FIRST_ROW >= rows:
            ws.row_dimensions[r].height = BLANK_ROW_PT
        else:
            ws.row_dimensions[r].height = MAX_ROW_PT if r in pinned else None


def generate_journal(conn: sqlite3.Connection, ws) -> dict:
    """Rewrite the Journal Table's data rows in ``ws`` from the database.

    Returns ``{rows, capacity, grew_to}``. Raises GenerateError rather than
    blank the log (see ``_refuse_to_blank``)."""
    table, last_row = _table_layout(ws, JOURNAL_TABLE, JOURNAL_HEADERS)
    entries = models.journal_rows(conn)
    _refuse_to_blank(ws, last_row, entries, "journal")

    old_last = last_row
    last_row = max(last_row, HEADER_ROW + len(entries))
    grew_to = last_row if last_row > old_last else None
    if grew_to:
        _grow_table(ws, table, "C", old_last, last_row)
        _update_schema_reference(ws.parent, JOURNAL_SHEET, last_row)

    for r in range(FIRST_ROW, last_row + 1):
        i = r - FIRST_ROW
        entry = entries[i] if i < len(entries) else None
        cells = [ws.cell(row=r, column=c) for c in (1, 2, 3)]
        if entry:
            _put(cells[0], datetime.datetime.fromisoformat(entry["date"]))
            _put(cells[1], entry["label"])
            _put(cells[2], entry["notes"])
        else:
            for cell in cells:
                _put(cell, None)
        for cell, style in zip(cells, _JOURNAL_STYLES):
            _style_cell(cell, style)
    _apply_journal_heights(ws, len(entries), last_row)
    return dict(rows=len(entries), capacity=last_row - HEADER_ROW, grew_to=grew_to)


def _rows_over_ceiling(path: str) -> List[int]:
    """Journal rows whose (LibreOffice-fitted) height exceeds Excel's ceiling."""
    ws = openpyxl.load_workbook(path)[JOURNAL_SHEET]
    return [r for r in range(FIRST_ROW, ws.max_row + 1)
            if (ws.row_dimensions[r].height or 0) > MAX_ROW_PT]


# ----------------------------------------------------------- Drawdown Log --

DRAWDOWN_SHEET = "Portfolio Drawdown Log"
DRAWDOWN_TABLE = "DrawdownLogTable"
DRAWDOWN_HEADERS = ("Date", "Portfolio Value ($)", "Running Peak ($)", "Drawdown %",
                    "Status", "Required Cash Floor", "New Position Initiation", "Notes")
DRAWDOWN_ROW_PT = 15.0


def _drawdown_formulas(r: int) -> Tuple[str, ...]:
    """Columns C-G for row ``r``, exactly as the live tab has them. Only the
    first row differs (no previous peak to chain from)."""
    peak = (f'=IF(B{r}="","",B{r})' if r == FIRST_ROW
            else f'=IF(B{r}="","",MAX(C{r - 1},B{r}))')
    return (
        peak,
        f'=IF(OR(B{r}="",C{r}=""),"",B{r}/C{r}-1)',
        f'=IF(D{r}="","",IF(D{r}>=-0.1,"Normal",IF(D{r}>=-0.2,"Elevated",'
        f'IF(D{r}>=-0.3,"Severe","Critical"))))',
        f'=IF(E{r}="","",IF(E{r}="Normal","5-15%",IF(E{r}="Elevated","15%",'
        f'IF(E{r}="Severe","20%","25%"))))',
        f'=IF(E{r}="","",IF(E{r}="Normal","Standard (score >=65)",'
        f'IF(E{r}="Elevated","Restricted (score >=80 only)","Suspended")))',
    )


# Canonical look of a Drawdown Log row, mirroring the live tab as of 2026-10-02.
# Inputs (A, B, H) are navy on pale blue; the formula columns are plain black.
_DRAWDOWN_FILL = PatternFill(fill_type="solid", fgColor="FFDDEBF7", bgColor="FFDDEBF7")
_NO_FILL = PatternFill(fill_type=None)
_BLACK = "FF000000"
# per column A..H: (bold, font colour, filled, horizontal, number format)
_DRAWDOWN_COLUMNS = (
    (False, _NAVY, True, "center", "mm/dd/yyyy"),
    (False, _NAVY, True, "center", "\\$#,##0"),
    (False, _BLACK, False, "center", "\\$#,##0"),
    (True, _BLACK, False, "center", "0.0%"),
    (True, _BLACK, False, "center", "General"),
    (False, _BLACK, False, "center", "General"),
    (False, _BLACK, False, "center", "General"),
    (False, _NAVY, True, "left", "General"),
)
_DRAWDOWN_STYLES = tuple(
    (Font(name="Arial Narrow", size=11, bold=bold, color=color),
     Alignment(horizontal=h, vertical="center", wrap_text=False), fmt,
     _DRAWDOWN_FILL if filled else _NO_FILL, _BORDER)
    for bold, color, filled, h, fmt in _DRAWDOWN_COLUMNS)


def generate_drawdown_log(conn: sqlite3.Connection, ws) -> dict:
    """Rewrite the Portfolio Drawdown Log Table from the database: inputs
    (date, value, notes) from the DB in date order, the five derived columns as
    the tab's own row-relative formulas in every pre-formatted row.

    Returns ``{rows, capacity, grew_to}``."""
    table, last_row = _table_layout(ws, DRAWDOWN_TABLE, DRAWDOWN_HEADERS)
    entries = models.drawdown_rows(conn)
    _refuse_to_blank(ws, last_row, entries, "drawdown_log")

    old_last = last_row
    last_row = max(last_row, HEADER_ROW + len(entries))
    grew_to = last_row if last_row > old_last else None
    if grew_to:
        wb = ws.parent
        _grow_table(ws, table, "H", old_last, last_row)
        _extend_conditional_formatting(ws, old_last, last_row)
        _extend_cross_sheet_references(wb, DRAWDOWN_SHEET, old_last, last_row)
        _update_schema_reference(wb, DRAWDOWN_SHEET, last_row)

    for r in range(FIRST_ROW, last_row + 1):
        i = r - FIRST_ROW
        entry = entries[i] if i < len(entries) else None
        cells = [ws.cell(row=r, column=c) for c in range(1, 9)]
        _put(cells[0], datetime.datetime.fromisoformat(entry["date"]) if entry else None)
        _put(cells[1], entry["portfolio_value"] if entry else None)
        _put(cells[7], entry["notes"] if entry else None)
        for cell, formula in zip(cells[2:7], _drawdown_formulas(r)):
            cell.value = formula            # a real formula, unlike _put's text guard
        for cell, style in zip(cells, _DRAWDOWN_STYLES):
            _style_cell(cell, style)
        ws.row_dimensions[r].height = DRAWDOWN_ROW_PT
    return dict(rows=len(entries), capacity=last_row - HEADER_ROW, grew_to=grew_to)


# ------------------------------------------------------------- workbook level --

class _Tab(NamedTuple):
    key: str
    sheet: str
    table: str
    ncols: int
    db_table: str
    generate: Callable[[sqlite3.Connection, object], dict]


_TABS: Dict[str, _Tab] = {t.key: t for t in (
    _Tab("journal", JOURNAL_SHEET, JOURNAL_TABLE, 3, "journal", generate_journal),
    _Tab("drawdown", DRAWDOWN_SHEET, DRAWDOWN_TABLE, 8, "drawdown_log", generate_drawdown_log),
)}


def _select_tabs(wb, tabs: Optional[Sequence[str]]) -> List[_Tab]:
    """``None`` means every generated tab the workbook actually has."""
    if tabs is None:
        return [t for t in _TABS.values() if t.sheet in wb.sheetnames]
    unknown = [k for k in tabs if k not in _TABS]
    if unknown:
        raise GenerateError(f"unknown tab(s) {unknown}; known: {sorted(_TABS)}")
    chosen = [_TABS[k] for k in tabs]
    missing = [t.sheet for t in chosen if t.sheet not in wb.sheetnames]
    if missing:
        raise GenerateError(f"workbook has no sheet(s) {missing}")
    return chosen


def _check_tables(conn: sqlite3.Connection, db_tables: Sequence[str], label: str) -> None:
    for name in db_tables:
        if not conn.execute("SELECT 1 FROM sqlite_master WHERE name=?", (name,)).fetchone():
            raise GenerateError(f"{label} has no {name} table -- not a landry database")


@contextlib.contextmanager
def _connection(db, db_tables: Sequence[str]):
    """``db`` is a path (opened here, closed after) or an open connection (left
    open: the ledger passes one holding an uncommitted write, so the regenerated
    tab sees the row before it is committed)."""
    if isinstance(db, sqlite3.Connection):
        _check_tables(db, db_tables, "the database")
        models.check_schema(db)
        yield db
        return
    if not os.path.exists(db):
        raise FileNotFoundError(
            f"{db} does not exist (sqlite would silently create an empty one)")
    conn = models.connect(db)
    try:
        _check_tables(conn, db_tables, db)
        models.check_schema(conn)
        yield conn
    finally:
        conn.close()


def generate_workbook(workbook_path: str, db, out_path: Optional[str] = None,
                      recalc: bool = True, tabs: Optional[Sequence[str]] = None) -> dict:
    """Regenerate ``tabs`` (default: every generated tab the workbook has) of
    ``workbook_path`` from ``db`` (a path or an open connection) and save to
    ``out_path`` (default: in place). ``recalc`` runs the mandatory LibreOffice pass afterwards -- openpyxl
    wipes every cached formula value on save, and that pass also fits the
    Journal's auto-height rows.

    Returns ``{"tabs": {key: {rows, capacity, grew_to}}, "recalc": ..., "capped_rows": [...]}``."""
    out_path = out_path or workbook_path
    wb = openpyxl.load_workbook(workbook_path)
    selected = _select_tabs(wb, tabs)
    with _connection(db, [t.db_table for t in selected]) as conn:
        result: dict = {"tabs": {t.key: t.generate(conn, wb[t.sheet]) for t in selected}}
    wb.save(out_path)
    result["capped_rows"] = []
    if recalc:
        from landry.xlsx_recalc import recalc as run_recalc
        result["recalc"] = run_recalc(out_path)
        journal = result["tabs"].get("journal")
        tall = _rows_over_ceiling(out_path) if journal else []
        if tall:
            wb = openpyxl.load_workbook(out_path)
            # Re-apply the policy rather than just re-save: openpyxl reads
            # LibreOffice's fitted heights back as explicit ones, which would
            # turn every auto-height row into a fixed one.
            _apply_journal_heights(wb[JOURNAL_SHEET], journal["rows"],
                                   HEADER_ROW + journal["capacity"], pinned=tall)
            wb.save(out_path)
            result["recalc"] = run_recalc(out_path)
        result["capped_rows"] = tall
    return result


# ------------------------------------------------------------------- verify --

def _rgb(color) -> Optional[str]:
    return color.rgb if color is not None and color.type == "rgb" else None


def _style_sig(cell) -> tuple:
    f, fl, b, a = cell.font, cell.fill, cell.border, cell.alignment
    return (f.name, f.sz, bool(f.b), bool(f.i), _rgb(f.color),
            fl.fill_type, _rgb(fl.fgColor),
            tuple((s.style, _rgb(s.color)) for s in (b.left, b.right, b.top, b.bottom)),
            a.horizontal, a.vertical, bool(a.wrap_text), cell.number_format)


_STYLE_FIELDS = ("font", "size", "bold", "italic", "font color", "fill", "fill color",
                 "borders", "h-align", "v-align", "wrap", "number format")


def _norm(value):
    """Compare numbers by value (an int 776000 and a float 776000.0 are the
    same cell) and everything else by type and value."""
    if isinstance(value, bool):
        return ("bool", value)
    if isinstance(value, (int, float)):
        return ("number", float(value))
    return (type(value).__name__, value)


def _conditional_formats(ws) -> list:
    return sorted((str(cf.sqref),
                   tuple((r.type, r.operator, tuple(r.formula or ())) for r in cf.rules))
                  for cf in ws.conditional_formatting)


def diff_tab(a, b, table_name: str, ncols: int) -> List[tuple]:
    """Differences between two sheets of the same generated tab as
    ``(kind, where, a, b)``: Table ref/filter, print area, conditional-format
    ranges, and every data cell's value (formulas compare as their text) and
    formatting. Row heights are excluded on purpose (derived layout)."""
    out = []
    ta, tb = a.tables[table_name], b.tables[table_name]
    if ta.ref != tb.ref:
        out.append(("table", "ref", ta.ref, tb.ref))
    fa = ta.autoFilter.ref if ta.autoFilter is not None else None
    fb = tb.autoFilter.ref if tb.autoFilter is not None else None
    if fa != fb:
        out.append(("table", "filter", fa, fb))
    if a.print_area != b.print_area:
        out.append(("print_area", "", a.print_area, b.print_area))
    cfa, cfb = _conditional_formats(a), _conditional_formats(b)
    if cfa != cfb:
        out.append(("conditional_formatting", "", cfa, cfb))
    for r in range(FIRST_ROW, max(a.max_row, b.max_row) + 1):
        for c in range(1, ncols + 1):
            ca, cb = a.cell(row=r, column=c), b.cell(row=r, column=c)
            if _norm(ca.value) != _norm(cb.value):
                out.append(("value", ca.coordinate, ca.value, cb.value))
            sa, sb = _style_sig(ca), _style_sig(cb)
            if sa != sb:
                for field, x, y in zip(_STYLE_FIELDS, sa, sb):
                    if x != y:
                        out.append(("style:" + field, ca.coordinate, x, y))
    return out


def diff_journal(a, b) -> List[tuple]:
    return diff_tab(a, b, JOURNAL_TABLE, 3)


def diff_drawdown_log(a, b) -> List[tuple]:
    return diff_tab(a, b, DRAWDOWN_TABLE, 8)


def verify_workbook(workbook_path: str, db,
                    tabs: Optional[Sequence[str]] = None) -> Dict[str, List[tuple]]:
    """Regenerate ``tabs`` in memory from the database (``db``: a path or an
    open connection) and diff each against
    the workbook as it stands: ``{tab: [diffs]}``. An empty list means
    generating would change nothing -- the DB and the tab agree on every value
    and every bit of formatting."""
    live = openpyxl.load_workbook(workbook_path)
    regenerated = openpyxl.load_workbook(workbook_path)
    selected = _select_tabs(live, tabs)
    with _connection(db, [t.db_table for t in selected]) as conn:
        for t in selected:
            t.generate(conn, regenerated[t.sheet])
    return {t.key: diff_tab(live[t.sheet], regenerated[t.sheet], t.table, t.ncols)
            for t in selected}


def verify_journal(workbook_path: str, db) -> List[tuple]:
    return verify_workbook(workbook_path, db, ("journal",))["journal"]


def verify_drawdown_log(workbook_path: str, db) -> List[tuple]:
    return verify_workbook(workbook_path, db, ("drawdown",))["drawdown"]


# ---------------------------------------------------------------------- CLI --

def _summarize(diffs: List[tuple], show: int = 6) -> str:
    if not diffs:
        return "identical: generating from the database would change nothing"
    lines = []
    for kind, n in Counter(d[0] for d in diffs).most_common():
        lines.append(f"  {kind}: {n}")
        for _, where, x, y in [d for d in diffs if d[0] == kind][:show]:
            lines.append(f"      {where or '-'}: {str(x)[:48]!r} -> {str(y)[:48]!r}")
    return "differences (workbook -> regenerated):\n" + "\n".join(lines)


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(
        description="Regenerate workbook tabs from landry.db (Phase C, one tab at a time)")
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name, help_ in (("journal", "rewrite the Journal tab from the database"),
                        ("drawdown", "rewrite the Portfolio Drawdown Log tab from the database"),
                        ("all", "rewrite every generated tab the workbook has")):
        p = sub.add_parser(name, help=help_)
        p.add_argument("--workbook", required=True)
        p.add_argument("--db", default=models.DEFAULT_DB_PATH)
        p.add_argument("--out", default=None, help="write here instead of in place")
        p.add_argument("--no-recalc", action="store_true",
                       help="skip the LibreOffice recalc (the workbook is then NOT safe to commit)")
    v = sub.add_parser("verify", help="diff tabs against what the database would generate")
    v.add_argument("--workbook", required=True)
    v.add_argument("--db", default=models.DEFAULT_DB_PATH)
    v.add_argument("--tab", choices=sorted(_TABS), default=None,
                   help="one tab (default: every generated tab the workbook has)")
    args = ap.parse_args(argv)

    if args.cmd == "verify":
        results = verify_workbook(args.workbook, args.db, (args.tab,) if args.tab else None)
        for key, diffs in results.items():
            print(f"{key}: {_summarize(diffs)}")
        return 1 if any(results.values()) else 0

    tabs = None if args.cmd == "all" else (args.cmd,)
    result = generate_workbook(args.workbook, args.db, args.out,
                               recalc=not args.no_recalc, tabs=tabs)
    for key, info in result["tabs"].items():
        print(f"{key}: {info['rows']} entries written, capacity {info['capacity']}"
              + (f", Table grew to row {info['grew_to']}" if info["grew_to"] else ""))
    if result["capped_rows"]:
        print(f"  {len(result['capped_rows'])} Journal entries exceed Excel's 409.5pt row "
              f"height and show their first lines in-cell: rows {result['capped_rows']}")
    if "recalc" in result:
        print(f"  recalc: {result['recalc'].get('status')}, "
              f"{result['recalc'].get('total_errors')} errors")
    return 0


if __name__ == "__main__":
    sys.exit(main())
