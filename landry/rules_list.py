"""The printable Hard Rules list (added 2026-10-05; Alan: "create an updated Rules list for printing").

``python -m landry rules-list [--pdf]`` writes ``docs/Landry System Hard Rules List v<VERSION>.docx`` (and, with
``--pdf``, a PDF beside it): every numbered Hard Rule of the rulebook, grouped by Part, in condensed wording, plus
a quick-reference page (decision bands, sizing, drawdown response, macro overlay, Hold Through, conflict order,
what to avoid, and since v1.05 the Stop-Review Ladder). Six Letter pages, Calibri.

WHY THE WORDING LIVES HERE AND IS CHECKED: the rule numbers are one running sequence across the whole rulebook, so
inserting a rule renumbers every later one (v1.04 inserted Rule 5 and shifted 5-50 to 6-51) -- the way every
earlier cheat sheet went stale (CLAUDE.md, 10/1-10/5). ``check()`` reads the rulebook docx itself (the numbered
paragraphs of its INDEX -- MASTER RULE REGISTER) and fails when the rulebook's version, its count of numbered rules,
or any number in a rule's wording no longer matches this list; ``tests/test_landry_rules_list.py`` runs it, and
``rules-list`` refuses to write a list that fails it. When it does: update PARTS below to the new text and
numbering, bump VERSION / REVISED, regenerate, re-print.

The wording is a condensation, not a copy: the text inside each Part of the rulebook governs. Every threshold is
the rulebook's own -- ``check()`` tests the numbers, not the prose, so read a changed rule's wording by eye.
"""

from __future__ import annotations

import datetime as dt
import glob
import os
import re
from collections import Counter
from typing import Dict, List, Optional

try:                                                       # python-docx: only this module needs it
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    HAVE_DOCX = True
except ImportError:                                        # pragma: no cover
    HAVE_DOCX = False

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


class RulesListError(RuntimeError):
    pass


def _need_docx() -> None:
    if not HAVE_DOCX:
        raise RulesListError("python-docx is not installed: pip install -r requirements.txt")


# --------------------------------------------------------------------------------------------------- the data --
# Each rule is (number, label-or-None, text); every number in the wording is the rulebook's own.

VERSION = "1.05"
REVISED = "2026-10-08"
INTRO_NOTE = ("Condensed from the Master Rule Register of {file}; every threshold is the rulebook’s own. Rule numbers "
              "are one running sequence: inserting a rule renumbers every later one (v1.04 inserted Rule 5, so the "
              "entry rules are now 6–14 and the sell triggers 22–35). v1.05 renumbered nothing: it added the Part 6 "
              "Stop-Review Ladder, an un-numbered Hard Rules block like Hold Through (see the quick reference), and put "
              "Rule 16's sector cap at market value on direct holdings. If this list and the text inside a Part ever "
              "differ, the Part governs.")
RENUMBERING_NOTE = ("From v1.04: no rule number changed. From v1.03: Rules 1–4 are unchanged; every v1.03 rule from 5 "
                    "onward is one number higher.")

PARTS = [
    {"title": "PART 3 — Composite Score & Confidence", "groups": [
        {"heading": None, "rules": [
            (1, "Tier 1 floor", "A Tier 1 weighted average below 3.0 cannot qualify as Buy or Strong Buy. "
                                "(Tier 1 weighted average = Σ Tier 1 score × weight ÷ 0.70.)"),
            (2, "Tier 1 score of 1", "Any Tier 1 indicator scored at 1 requires a mandatory qualitative review before any decision."),
            (3, "Two weak Tier 1s", "Two or more Tier 1 indicators scored at 2.0 or below automatically classify the stock as Avoid."),
            (4, "Low-confidence cap", "Two or more Tier 1 indicators tagged Low Confidence: no Strong Buy. Cap the classification "
                                      "at Buy until re-scored with better data."),
            (5, "Technical / momentum floor", "Relative Strength vs. SPY or Technical Trend scored at 1, at initial scoring or any "
                                              "later re-score, requires a mandatory qualitative review before a Buy or Strong Buy "
                                              "classification is relied upon, for a new entry or a position already held. Not an "
                                              "automatic reclassification: document the conclusion (reaffirmed, resized, or referred "
                                              "to Part 6) in the Journal."),
        ]}]},
    {"title": "PART 4 — Entry Rules", "groups": [
        {"heading": None, "rules": [
            (6, "Composite", "Composite Score of 65 or higher."),
            (7, "Tier 1 average", "Tier 1 weighted average of 3.0 or higher."),
            (8, "No Tier 1 at 1", "No Tier 1 indicator scored at 1."),
            (9, "Technical staging", "Price above the 200-week moving average, or reclaimed within the past 6 months: standard "
                                     "initial sizing. Otherwise the first purchase is limited to 50% of the normal initial size, "
                                     "with a 90-day technical review."),
            (10, "Monthly MACD", "Positive or turning positive: standard staging. Negative: initiate only at the reduced size in "
                                 "Rule 9. Technical weakness alone does not override a qualifying fundamental case."),
            (11, "Binary events", "No near-term binary event risk unless the thesis explicitly includes it."),
            (12, "Bear Case", "At entry, compute the implied 5-year annualized total return under a documented Bear Case "
                              "(non-cyclicals: cut the Base Case FCF growth assumption by at least 50% and use a conservative "
                              "terminal multiple; cyclicals: a documented trough-cycle assumption). It must be 0% or higher; if "
                              "negative, do not enter regardless of Composite Score."),
            (13, "Scenarios", "At entry and at each annual re-score, document Bull, Base and Bear, each with its own implied "
                              "5-year return and a Likely / Possible / Unlikely tag. Exactly one is Likely. No numeric probabilities."),
            (14, "Valuation ceiling", "No new nonfinancial operating-company position when P/FCF exceeds 50x, unless both (a) "
                                      "documented consensus FCF growth exceeds 30% annualized for at least the next two years and "
                                      "(b) the Bear Case clears the 0%-or-higher implied-return threshold. Use the Part 10 substitute where P/FCF is not "
                                      "economically meaningful."),
        ]}]},
    {"title": "PART 5 — Position Sizing", "groups": [
        {"heading": None, "rules": [
            (15, "Position cap", "No single position above 8% of portfolio at cost."),
            (16, "Sector cap", "No single sector above 25% of portfolio, measured at market value on direct holdings (ETFs are not looked through)."),
            (17, "Position floor", "At least 12 qualifying positions in a fully invested portfolio (12 is a floor, not a guarantee)."),
            (18, "Cash band", "Cash always between 5% and 15%, for dry powder and risk management."),
            (19, "No threshold lowering", "If fewer than 12 candidates qualify (Composite Score 65 or higher), do not lower the "
                                          "threshold to reach full investment. Hold the difference in cash."),
            (20, "Leverage cap", "Nonfinancial operating companies with Debt/FCF above 5x are capped at Composite 65 (no Strong "
                                 "Buy); the −1% sizing modifier applies above 4x. For financials and other businesses where Debt/FCF is not "
                                 "meaningful, use the Part 10 balance-sheet substitute."),
            (21, "Beta overlay", "Use 5-year weekly beta, 2-year as fallback. Reduce the maximum full position by 1% for beta "
                                 "1.40 to 1.59 and by 2% for beta 1.60 or higher; total beta adjustment capped at 2%. A sizing "
                                 "input only: investigate material disagreement across data sources, and assign Low Confidence "
                                 "when no reliable estimate exists."),
        ]}]},
    {"title": "PART 6 — Exit & Sell Discipline", "groups": [
        {"heading": "Sell triggers — fundamental (act within 30 days of confirmation)", "rules": [
            (22, None, "FCF Yield & Trend falls to 1 for two consecutive quarters."),
            (23, None, "Competitive Moat degrades from durable to absent."),
            (24, None, "ROIC falls below WACC for two consecutive years."),
            (25, None, "Management quality falls to 1 because of destructive capital allocation."),
            (26, None, "Revenue growth turns negative for two consecutive quarters without a credible cyclical explanation "
                       "(a documented cycle rationale applies here as in Part 4)."),
        ]},
        {"heading": "Sell triggers — valuation (trim or sell when price exceeds intrinsic value), then score-band rules", "rules": [
            (27, None, "Debt-to-FCF exceeds 6x with no credible deleveraging path."),
            (28, None, "P/FCF rises above 50x while FCF growth is below 15%."),
            (29, None, "Implied 5-year annualized return falls below 7%."),
            (30, None, "Position size grows above the maximum full position because of appreciation."),
            (31, "Strong Buy to Buy", "A re-score from Strong Buy to Buy requires trimming only to the Buy maximum full position "
                                      "within 30 days. A re-score to Watch List (50–64) does not force an immediate sale: additions stop and "
                                      "the 90-day review begins. Scores below 50 follow the rules immediately below."),
            (32, "Exit Review (35–49)", "Classify as Exit Review. Establish a dated 90-day remediation plan only when a credible, "
                                        "specifically documented path back to at least 50 exists; otherwise sell within 30 days."),
            (33, "Probationary Hold (50–64)", "No additions; a 90-day re-score is required. If it falls below 50, apply Rule 32. An "
                                              "approved Exit Review holding that has not returned to at least 50 by its 90-day "
                                              "deadline must be sold within 30 days; no extension is permitted."),
            (34, "Early Partial Trim", "Probationary Hold only. If the same Tier 1 or Tier 2 indicator(s) that caused entry are "
                                       "reconfirmed unimproved or worse before the 90-day re-score (a new filing, a scheduled "
                                       "re-score, or equivalent documented evidence), trim one-third of the share count held at "
                                       "entry. May recur up to twice more in the same 90-day window (maximum three trims; full "
                                       "liquidation in the worst case). Trims are not reversed by a later recovery; recovery to 65 or "
                                       "above ends further trims. Not triggered by price movement alone."),
            (35, None, "Composite Score below 35: sell within 30 days unless a more urgent fundamental trigger requires earlier action."),
        ]},
        {"heading": "Opportunity-cost replacement", "rules": [
            (36, "Replacement test", "A lower-ranked holding may be replaced only when all of these hold: (a) the new opportunity "
                                     "clears the Entry Checklist and Valuation Ceiling (Part 4); (b) the score gap is at least 15 "
                                     "points across two consecutive quarterly re-scores; (c) the holding is not protected by Hold "
                                     "Through; (d) available cash cannot fund the required initial size, or doing so would breach a "
                                     "hard cap; and (e) the comparison, tax effect and capital constraint are documented before "
                                     "execution."),
            (37, None, "Replacement may not be invoked more than once per holding in any rolling 12-month period."),
        ]}]},
    {"title": "PART 7 — Portfolio Risk Management", "groups": [
        {"heading": None, "rules": [
            (38, "Correlation cap", "Quarterly, calculate pairwise correlations on weekly total returns over the longest common "
                                    "period up to 36 months (24-month target, 12-month minimum). Flag any cluster of three or more "
                                    "holdings with pairwise correlation above 0.70 and document the shared economic risk driver. A "
                                    "new position that expands a flagged cluster starts at 50% of the normal initial size, and the "
                                    "cluster's aggregate exposure may not exceed 20% without written approval. A risk-review "
                                    "trigger, not an automatic forced sale."),
            (39, "Drawdown response", "Drawdown restrictions on new deployment apply automatically and are not overridden by "
                                      "individual conviction. Higher cash levels are reached through the Cash-Raising Waterfall, "
                                      "not by indiscriminate sale of fundamentally intact holdings. Normal deployment resumes when "
                                      "the regime-exit test is met; security-specific sizing restrictions remain until the "
                                      "holding is re-scored."),
        ]}]},
    {"title": "PART 8 — Tax-Aware Portfolio Management", "groups": [
        {"heading": None, "rules": [
            (40, "Mandated sells", "A sell mandated by a fundamental deterioration trigger (Part 6) executes within its required "
                                   "window regardless of holding period or tax lot."),
            (41, "Discretionary trims", "A discretionary trim (valuation-driven, sizing or sector-cap) may be delayed up to 30 days "
                                        "to let a lot reach long-term status, provided the delay does not violate any hard rule."),
            (42, "Lot selection", "When trimming a multi-lot position, use specific-lot identification and select the lot or "
                                  "combination expected to minimize after-tax cost. Highest basis is not always optimal: consider "
                                  "loss harvesting, short- and long-term rates, carryforwards and account-level constraints. "
                                  "Confirm the lot selection with the broker before the sale."),
            (43, "Loss harvesting", "Losses in Probationary Hold, Exit Review or underwater positions may be harvested before a "
                                    "mandatory sell deadline only when the replacement is not substantially identical and the "
                                    "transaction is coordinated across the relevant household accounts. Redeploy proceeds subject "
                                    "to the target cash allocation and all portfolio risk rules."),
            (44, "Tax never decides", "Tax considerations never independently trigger a trade; they govern only the timing and lot "
                                      "selection of a decision already justified elsewhere in the system."),
        ]}]},
    {"title": "PART 9 — Review Schedule, Documentation & Governance", "groups": [
        {"heading": None, "rules": [
            (45, "Quarterly", "Review all Watch List names and re-score any holding affected by material news."),
            (46, "Annually", "Complete a full re-score of every portfolio position and review every exit from the prior year."),
            (47, "After every exit", "Record the sell reason and the outcome versus the original thesis."),
            (48, "After a loss over 20%", "After any loss greater than 20% from cost, run a post-mortem and identify the missed "
                                          "warning signals."),
            (49, "Every 2 years", "Review the indicator weights; adjust only with documented evidence from multiple positions over time."),
        ]}]},
    {"title": "PART 10 — Sector Adaptations & Operating Definitions", "groups": [
        {"heading": None, "rules": [
            (50, "Financial leverage caps", "Composite Score capped at 65 (no Strong Buy) for banks with a CET1 ratio below 9%; "
                                            "insurers with an RBC ratio below 350%; brokers and asset managers with net leverage "
                                            "(excluding client matched-book financing) above 3x tangible equity, or Debt/EBITDA "
                                            "above 3x for fee-based asset managers. Reasoned defaults, to be reviewed and confirmed "
                                            "before being relied upon."),
            (51, "Financial valuation ceilings", "No new financial-company position when its multiple exceeds P/TBV of 2.5x (banks), "
                                                 "P/B of 2.0x or P/E of 20x (insurers), or P/E of 25x (brokers and asset managers), "
                                                 "unless both a documented ROE/growth justification and a tested sector-specific Bear "
                                                 "Case clear the same threshold required by the Bear-Case Implied Return Test "
                                                 "(Part 4). Reasoned defaults, to be reviewed and confirmed before being relied upon."),
        ]}]},
]

# ----------------------------------------------------------------- quick reference (not numbered rules) --

DECISION_BANDS = [   # score, decision, new candidate, existing holding  (Part 3)
    ("80–100", "STRONG BUY", "Initiate up to max initial size; build only after confirmation",
     "Core Hold / Add, subject to sizing and entry rules"),
    ("65–79", "BUY", "Initiate up to max initial size; build only after confirmation",
     "Hold; trim only if above Buy maximum or another rule requires action"),
    ("50–64", "WATCH LIST", "No new capital; monitor for improvement", "Probationary Hold; no additions; 90-day review plan"),
    ("35–49", "AVOID", "Do not initiate; revisit only after material improvement",
     "Exit Review; sell within 30 days unless a dated 90-day remediation plan is approved"),
    ("Below 35", "PASS", "Remove from active coverage", "Mandatory sell within 30 days"),
]

SIZING = [            # composite, max initial, max full  (Part 5)
    ("80–100 (Strong Buy)", "5% of portfolio", "8% of portfolio"),
    ("65–79 (Buy)", "3% of portfolio", "5% of portfolio"),
    ("Below 65", "0%", "0%"),
]
SIZING_MODIFIERS = [
    "Insider ownership above 5% with recent open-market purchases: +1% to max position",
    "Debt/FCF above 4x: −1% to max position (fragility discount)",
    "Single customer above 30% of revenue: −1% to max position",
    "Position in same sector as an existing holding above 5%: cap combined sector exposure at 25% of portfolio",
    "Macro headwind active (Part 7): −1% to all new positions in affected sectors",
    "Beta overlay: Rule 21",
]

DRAWDOWN = [          # level, status, required response  (Part 7)
    ("0–10%", "Normal", "Standard sizing and the 5–15% cash rules (Part 5)."),
    ("10–20%", "Elevated", "Raise the cash floor to 15% minimum. Suspend initiation of new positions scored below 80 "
                           "(Strong Buy only)."),
    ("20–30%", "Severe", "Raise the cash floor to 20% minimum. Suspend all new position initiation regardless of Composite "
                         "Score. Hold Through rules remain in force for fundamentally intact holdings."),
    ("30%+", "Critical", "Raise the cash floor to 25% minimum. Re-score the full portfolio within 2 weeks. Document whether "
                         "the decline reflects fundamentals, valuation compression or macro conditions before any further action."),
]
DRAWDOWN_NOTE = ("Drawdown is measured on aggregate portfolio total return, net of external deposits and withdrawals, from the "
                 "trailing high-water mark (daily closing values when available, otherwise weekly). A regime begins after five "
                 "consecutive trading days beyond a threshold and ends after ten consecutive trading days inside the less "
                 "restrictive band. A new high-water mark resets the measurement.")

MACRO = [             # condition, effect  (Part 7)
    ("Rising rate environment (Fed tightening cycle active)",
     "Reduce max position size by 1% for high-multiple growth stocks (P/FCF above 40x). Favor FCF yield and low-debt names."),
    ("Inverted yield curve (2s/10s inverted more than 3 months)",
     "Raise the cash floor to 10% minimum. No new positions in cyclicals or leveraged businesses. Favor defensive compounders."),
    ("Credit spreads widening rapidly (HY spreads above 600 bps)",
     "Pause all new entries. Reassess existing holdings with Debt/FCF above 3x. Hold cash."),
    ("Broad market downtrend (SPY below its 200-week MA)",
     "Reduce all new position sizes by 50%. Continue holding fundamentally intact positions. Accumulate Watch List "
     "candidates for entry when the trend recovers."),
    ("Sector-specific regulatory / geopolitical risk",
     "Reduce the sector cap from 25% to 15% for the affected sector. Re-evaluate moat scores for regulatory dependency."),
]
MACRO_NOTE = "Macro overlays are temporary modifiers: when the condition resolves, revert to standard sizing."

HOLD_THROUGH = [      # do NOT sell based on these alone  (Part 6)
    "A short-term price decline of 20–30% with no change in Tier 1/2 fundamentals: this is noise, unless a Stop-Review Ladder rung (below) has been reached.",
    "One quarter of earnings miss: evaluate the cause; sell only if it signals structural deterioration.",
    "Analyst downgrades: a lagging indicator; revisit Tier 1/2 signals instead.",
    "A market-wide selloff: if fundamentals are intact, often an opportunity to add, not sell.",
    "Narrative shift or negative press: evaluate facts, not sentiment.",
]

STOP_REVIEW = [       # Part 6, added in v1.05 -- an un-numbered Hard Rules block, like Hold Through
    "Rung 1: 20% or more below cost, with Relative Strength 1 or the price under its 200-day average: no additions (DCA included) "
    "and a documented re-underwrite within 10 trading days -- a fresh Tier 1 reread, the leading indicators, and the Rule 5 review "
    "concluded as reaffirmed, resized or referred; repeated every 90 days.",
    "Rung 2: 30% or more below cost on the same confirmation, with deterioration found: trim one third; a second confirmation trims another third.",
    "Rung 3: 40% or more below cost, Relative Strength 1 and 8 straight weeks under the 200-day average: Exit Review "
    "(sell within 30 days unless a documented 90-day remediation plan is approved).",
    "Never a price-only sale: a decline that reaches no rung, or a rung whose re-underwrite finds no deterioration, stays under Hold Through.",
]
STOP_REVIEW_NOTE = "Relative Strength 1 = behind SPY by more than 15% over six months; the 200-day average = the last 40 weekly closes."

CONFLICT_ORDER = [    # Part 10: higher rank always wins
    "Mandatory fundamental sell triggers",
    "Legal, regulatory, account and liquidity constraints (mandatory tax compliance, e.g. wash-sale rules, falls here)",
    "Portfolio hard caps and drawdown restrictions",
    "Entry eligibility and valuation ceilings",
    "Position sizing and staging",
    "Opportunity-cost replacement",
    "Tax optimization (discretionary lot selection and loss harvesting; never a mandatory compliance matter)",
]
CONFLICT_NOTE = "A lower-ranked rule may never override a higher-ranked rule. Document any unresolved conflict before trading."

AVOID = [             # Part 11: avoid, why
    ("P/E ratio as primary filter", "Use only as a secondary cross-check."),
    ("Analyst 12-month price targets", "A sentiment and estimate-dispersion check only."),
    ("Earnings surprises as primary signal", "Insufficient for a multi-year thesis."),
    ("Narrative momentum (“AI will change everything”)",
     "Require evidence in revenue, margins, cash generation and implied return."),
    ("Averaging down into deteriorating fundamentals",
     "Price decline is often opportunity; fundamental deterioration is a reason to sell. Never add when Tier 1 scores are falling."),
    ("Over-diversification below conviction threshold",
     "Hold only qualifying ideas and comply with portfolio risk limits."),
    ("Market timing as a primary strategy",
     "Use the macro-overlay to adjust staging and risk, not as an all-in / all-out forecast."),
]


# ------------------------------------------------------------------------------------------- the builder --

FONT = "Calibri"
NAVY = (0x1F, 0x38, 0x64)                              # RGB tuples: this module imports without python-docx
GREY = (0x59, 0x59, 0x59)
BODY_PT, TABLE_PT, REF_PT = 11, 10.5, 10
PART_FILL, GROUP_FILL, HEAD_FILL = "DCE6F1", "F2F2F2", "E7E6E6"
BORDER = "A6A6A6"
FULL = 10368                                           # usable width in DXA (7.2 in at 0.65 in margins)


# ----------------------------------------------------------------------------------------------- xml helpers --

def _font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def _shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _table_setup(table, widths):
    """Fixed layout, thin grey borders, tight cell margins, explicit column widths (DXA). The tblPr children are
    written in the order the schema wants (tblW, jc, tblBorders, tblLayout, tblCellMar, tblLook)."""
    tbl = table._tbl
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout", "w:tblBorders", "w:tblCellMar", "w:jc"):
        for el in tblpr.findall(qn(tag)):
            tblpr.remove(el)
    look = tblpr.find(qn("w:tblLook"))

    def put(el):
        if look is not None:
            look.addprevious(el)
        else:
            tblpr.append(el)

    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(sum(widths))); w.set(qn("w:type"), "dxa"); put(w)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); put(jc)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:space"), "0"); b.set(qn("w:color"), BORDER)
        borders.append(b)
    put(borders)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); put(lay)
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 20), ("left", 90), ("bottom", 20), ("right", 90)):
        m = OxmlElement(f"w:{side}"); m.set(qn("w:w"), str(v)); m.set(qn("w:type"), "dxa"); mar.append(m)
    put(mar)
    for gc, wd in zip(tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(wd))


def _set_widths(row, widths):
    for cell, wd in zip(row.cells, widths):
        tcpr = cell._element.get_or_add_tcPr()
        for el in tcpr.findall(qn("w:tcW")):
            tcpr.remove(el)
        tcw = OxmlElement("w:tcW"); tcw.set(qn("w:w"), str(wd)); tcw.set(qn("w:type"), "dxa")
        tcpr.insert(0, tcw)


def _cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def _para(cell, first=True):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
    return p


def _keep_next(row):
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.keep_with_next = True


def _span(row, n):
    cell = row.cells[0]
    for i in range(1, n):
        cell = cell.merge(row.cells[i])
    return cell


def _field(par, instr, size, color):
    """A complex field (PAGE / NUMPAGES) as five runs, each with the footer font, so the result is not left at the
    default size."""
    def put(el):
        r = par.add_run()
        _font(r, size=size, color=color)
        r._element.append(el)
    for kind in ("begin", "instr", "separate", "text", "end"):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = f" {instr} "
        elif kind == "text":
            el = OxmlElement("w:t"); el.text = "1"
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        put(el)


# ------------------------------------------------------------------------------------------------ builders --

def heading(doc, text, size=13, space_before=10, space_after=3, keep=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before); pf.space_after = Pt(space_after); pf.keep_with_next = keep
    _font(p.add_run(text), size=size, bold=True, color=NAVY)
    return p


def note(doc, text, size=9.5, italic=True, space_after=4, space_before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before); pf.space_after = Pt(space_after)
    _font(p.add_run(text), size=size, italic=italic, color=GREY)
    return p


def rules_table(doc, part, keep_together=True):
    widths = [620, FULL - 620]
    n_rows = 1 + sum((1 if g["heading"] else 0) + len(g["rules"]) for g in part["groups"])
    t = doc.add_table(rows=n_rows, cols=2)
    _table_setup(t, widths)
    i = 0
    row = t.rows[i]; _set_widths(row, widths)
    c = _span(row, 2); _shade(c, PART_FILL)
    _font(_para(c).add_run(part["title"]), size=12, bold=True, color=NAVY)
    _keep_next(row); _cant_split(row)
    first_rule_row = True
    for g in part["groups"]:
        if g["heading"]:
            i += 1
            row = t.rows[i]; _set_widths(row, widths)
            c = _span(row, 2); _shade(c, GROUP_FILL)
            _font(_para(c).add_run(g["heading"]), size=TABLE_PT - 1, bold=True, italic=True, color=NAVY)
            _keep_next(row); _cant_split(row)
        for n, label, text in g["rules"]:
            i += 1
            row = t.rows[i]; _set_widths(row, widths)
            _cant_split(row)
            a, b = row.cells
            p = _para(a); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _font(p.add_run(str(n)), size=TABLE_PT, bold=True, color=NAVY)
            p = _para(b)
            if label:
                _font(p.add_run(label + ". "), size=TABLE_PT, bold=True)
            _font(p.add_run(text), size=TABLE_PT)
            if keep_together:                                    # a short Part never splits across pages
                _keep_next(row)
    if keep_together:                                        # ... but its last row ends the chain
        for cell in t.rows[-1].cells:
            for pp in cell.paragraphs:
                pp.paragraph_format.keep_with_next = False
    return t


def grid_table(doc, headers, rows, widths, size=REF_PT, bold_first=False, shade_first_col=False, keep=True):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    _table_setup(t, widths)
    head = t.rows[0]; _set_widths(head, widths); _cant_split(head)
    head._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for cell, h in zip(head.cells, headers):
        _shade(cell, HEAD_FILL)
        _font(_para(cell).add_run(h), size=size, bold=True, color=NAVY)
    _keep_next(head)
    for r_i, data in enumerate(rows, start=1):
        row = t.rows[r_i]; _set_widths(row, widths); _cant_split(row)
        for c_i, (cell, val) in enumerate(zip(row.cells, data)):
            if shade_first_col and c_i == 0:
                _shade(cell, GROUP_FILL)
            _font(_para(cell).add_run(val), size=size, bold=(bold_first and c_i == 0) or (c_i == 1 and headers[1] in ("Decision", "Status")))
    if keep:                                                 # a short table never splits across pages
        for row in t.rows[:-1]:
            _keep_next(row)
    return t


def bullets(doc, items, size=REF_PT, numbered=False):
    for k, s in enumerate(items, start=1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.28); pf.first_line_indent = Inches(-0.2)
        pf.space_before = Pt(0); pf.space_after = Pt(1.5)
        _font(p.add_run((f"{k}.  " if numbered else "•  ") + s), size=size)


def _build(out, today):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.65)
    sec.top_margin, sec.bottom_margin = Inches(0.6), Inches(0.7)
    sec.footer_distance = Inches(0.4)
    st = doc.styles["Normal"]
    st.font.name = FONT; st.font.size = Pt(BODY_PT)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    zoom = doc.settings.element.find(qn("w:zoom"))          # python-docx's template leaves the required percent off
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    cp = doc.core_properties
    cp.title = f"Landry System Hard Rules List v{VERSION}"
    cp.author = "Landry System"
    cp.subject = "Printable list of the numbered Hard Rules"

    # title block
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    _font(p.add_run("THE LANDRY SYSTEM — HARD RULES LIST"), size=18, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    _font(p.add_run(f"Version {VERSION} (revised {REVISED})  |  {len(rules())} numbered Hard Rules, Parts 3–10  |  "
                    f"Printed {today.strftime('%B')} {today.day}, {today.year}"), size=10.5, color=GREY)
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "8"); b.set(qn("w:space"), "4"); b.set(qn("w:color"), "1F3864")
    bdr.append(b)
    ppr.insert_element_before(bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                              "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
                              "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                              "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                              "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                              "w:pPrChange")
    note(doc, INTRO_NOTE.format(file=f"LANDRY_SYSTEM_v{VERSION.replace('.', '-')}_FINAL.docx"), space_after=2)
    note(doc, RENUMBERING_NOTE, space_after=6)

    for part in PARTS:
        rules_table(doc, part, keep_together=not part["title"].startswith("PART 6"))
        sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0); sp.paragraph_format.space_before = Pt(0)
        _font(sp.add_run(""), size=4)

    # quick reference: starts wherever Part 10 ends
    heading(doc, "QUICK REFERENCE", size=15, space_before=8)
    note(doc, "Tables and lists the numbered rules lean on, taken from the Parts themselves (they are not numbered rules).",
         space_after=2)

    heading(doc, "Decision bands (Part 3)", size=11.5)
    grid_table(doc, ["Score", "Decision", "New candidate", "Existing holding"], DECISION_BANDS,
               [1050, 1350, 3600, 4368], bold_first=True)

    heading(doc, "Position sizing (Part 5, Rules 15–21)", size=11.5)
    grid_table(doc, ["Composite Score", "Max initial position", "Max full position"], SIZING, [3100, 3634, 3634],
               bold_first=True)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.keep_with_next = True
    _font(p.add_run("Sizing modifiers"), size=REF_PT, bold=True)
    bullets(doc, SIZING_MODIFIERS)

    heading(doc, "Portfolio drawdown response (Part 7, Rule 39)", size=11.5)
    grid_table(doc, ["Drawdown", "Status", "Required response"], DRAWDOWN, [1150, 1100, 8118], bold_first=True, keep=False)
    note(doc, DRAWDOWN_NOTE, size=9, space_after=2, space_before=3)

    heading(doc, "Macro overlay (Part 7)", size=11.5)
    grid_table(doc, ["Condition", "Effect on the system"], MACRO, [3400, 6968], bold_first=True)
    note(doc, MACRO_NOTE, size=9, space_after=2, space_before=3)

    heading(doc, "Hold Through: do NOT sell based on these alone (Part 6)", size=11.5)
    bullets(doc, HOLD_THROUGH)

    heading(doc, "Stop-Review Ladder: a falling price triggers a re-underwrite (Part 6, v1.05)", size=11.5)
    bullets(doc, STOP_REVIEW)
    note(doc, STOP_REVIEW_NOTE, size=9, space_after=2, space_before=2)

    heading(doc, "When rules conflict: higher rank wins (Part 10)", size=11.5)
    bullets(doc, CONFLICT_ORDER, numbered=True)
    note(doc, CONFLICT_NOTE, size=9, space_after=2, space_before=2)

    heading(doc, "What to avoid (Part 11)", size=11.5)
    grid_table(doc, ["Avoid", "Why"], AVOID, [3800, 6568], bold_first=True)

    # footer: title left, page x of y flush right (the Normal style carries no stray tab stops)
    fp = sec.footer.paragraphs[0]
    fp.style = doc.styles["Normal"]
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(7.2), WD_TAB_ALIGNMENT.RIGHT)
    _font(fp.add_run(f"Landry System v{VERSION} — Hard Rules List\tPage "), size=8.5, color=GREY)
    _field(fp, "PAGE", 8.5, GREY)
    _font(fp.add_run(" of "), size=8.5, color=GREY)
    _field(fp, "NUMPAGES", 8.5, GREY)
    doc.save(out)


# ---------------------------------------------------------------------------- reading and checking the rulebook --

def rulebook_path(repo_dir: Optional[str] = None) -> Optional[str]:
    """The newest LANDRY_SYSTEM_v<major>-<minor>_FINAL.docx in the repo root, or None."""
    found = glob.glob(os.path.join(repo_dir or REPO, "LANDRY_SYSTEM_v*_FINAL.docx"))

    def key(p: str):
        m = re.search(r"_v(\d+)-(\d+)_", os.path.basename(p))
        return (int(m.group(1)), int(m.group(2))) if m else (0, 0)
    return max(found, key=key) if found else None


def default_out(repo_dir: Optional[str] = None) -> str:
    return os.path.join(repo_dir or REPO, "docs", f"Landry System Hard Rules List v{VERSION}.docx")


def _numid(p) -> Optional[str]:
    ppr = p._p.pPr
    numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
    nid = numpr.find(qn("w:numId")) if numpr is not None else None
    return nid.get(qn("w:val")) if nid is not None else None


def rulebook_rules(path: str) -> Dict[int, str]:
    """{rule number: wording} from the rulebook's INDEX -- MASTER RULE REGISTER. The numbers are Word's own
    auto-numbering, so the n-th paragraph of the INDEX's main numbered list (the list with the most items there; the
    a)-e) sub-items of one rule are a different list) is Rule n; the paragraphs that follow it up to the next rule
    are part of its wording."""
    _need_docx()
    paras = Document(path).paragraphs
    starts = [i for i, p in enumerate(paras) if p.text.startswith("INDEX — MASTER RULE REGISTER")]
    if not starts:
        raise RulesListError("no 'INDEX — MASTER RULE REGISTER' heading in the rulebook")
    start = starts[-1]
    stop = next((i for i in range(start + 1, len(paras)) if paras[i].text.startswith("Part 11")), len(paras))
    region = paras[start:stop]
    ids = [_numid(p) for p in region]
    counts = Counter(i for i in ids if i)
    if not counts:
        raise RulesListError("the INDEX has no numbered paragraphs")
    main = counts.most_common(1)[0][0]
    rules: Dict[int, str] = {}
    n = 0
    for p, nid in zip(region, ids):
        if nid == main:
            n += 1
            rules[n] = p.text.strip()
        elif n and not p.style.name.startswith("Heading"):
            rules[n] += " " + p.text.strip()
    return rules


def rulebook_version(path: str) -> Optional[str]:
    _need_docx()
    for p in Document(path).paragraphs[:25]:
        m = re.match(r"^Version\s+(\d+\.\d+)\b", p.text.strip())
        if m:
            return m.group(1)
    return None


def _full_text(path: str) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def _norm(s: str) -> str:
    return (s.replace("−", "-").replace("–", "-").replace("—", "-").replace("×", "x").replace(",", "")
             .replace("“", '"').replace("”", '"').replace("’", "'"))


def _nums(s: str, drop_refs: bool = False) -> set:
    s = _norm(s)
    if drop_refs:                                           # "Rule 9", "Parts 3-10" cite other places, not thresholds
        s = re.sub(r"\b(?:Rules?|Parts?)\s+\d+(?:\s*(?:and|-)\s*\d+)?", " ", s)
    return set(re.findall(r"\d+(?:\.\d+)?", s))


def rules() -> Dict[int, tuple]:
    """{number: (label, text)} of this list, in order of appearance."""
    return {n: (label, text) for part in PARTS for g in part["groups"] for n, label, text in g["rules"]}


def _quick_reference_lines() -> List[str]:
    out: List[str] = []
    for row in DECISION_BANDS + SIZING + DRAWDOWN + MACRO + AVOID:
        out.extend(row)
    return out + SIZING_MODIFIERS + [DRAWDOWN_NOTE] + HOLD_THROUGH + STOP_REVIEW + [STOP_REVIEW_NOTE] + CONFLICT_ORDER


def check(path: Optional[str] = None) -> List[str]:
    """Problems that make this list wrong about the rulebook ([] when it matches): a different version, a different
    count or numbering of rules, a number in a rule's wording (or in the quick reference) that the rulebook does not
    contain. Numbers only -- the prose is a condensation to be read by eye when a rule changes."""
    _need_docx()
    path = path or rulebook_path()
    if not path or not os.path.exists(path):
        return ["no LANDRY_SYSTEM_v*_FINAL.docx found to check against"]
    problems: List[str] = []
    version = rulebook_version(path)
    if version != VERSION:
        problems.append(f"the rulebook is v{version}, this list is v{VERSION}")
    mine = rules()
    if sorted(mine) != list(range(1, len(mine) + 1)) or list(mine) != sorted(mine):
        problems.append("this list's rule numbers are not 1..N in order")
    theirs = rulebook_rules(path)
    if set(theirs) != set(mine):
        missing, extra = sorted(set(theirs) - set(mine)), sorted(set(mine) - set(theirs))
        problems.append(f"the rulebook numbers {len(theirs)} rules, this list {len(mine)}"
                        + (f"; rulebook only: {missing[:8]}" if missing else "")
                        + (f"; list only: {extra[:8]}" if extra else ""))
    for n in sorted(set(theirs) & set(mine)):
        stray = _nums(mine[n][1], drop_refs=True) - _nums(theirs[n])
        if stray:
            problems.append(f"Rule {n}: the list has {sorted(stray)} where the rulebook's rule does not")
    body = _nums(_full_text(path))
    for line in _quick_reference_lines():
        stray = _nums(line, drop_refs=True) - body
        if stray:
            problems.append(f"quick reference: {sorted(stray)} is not in the rulebook ({line[:60]!r})")
    return problems


# ----------------------------------------------------------------------------------------------------- writing --

def build(out_path: Optional[str] = None, today: Optional[dt.date] = None) -> str:
    """Write the printable list and return its path. Does not check it against the rulebook (``check()`` does)."""
    _need_docx()
    out = out_path or default_out()
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    _build(out, today or dt.date.today())
    return out


def to_pdf(docx_path: str, timeout: int = 180) -> Optional[str]:
    """Render ``docx_path`` to a PDF beside it with LibreOffice; the path, or None when soffice is missing."""
    import subprocess
    from landry.xlsx_recalc import soffice_path
    exe = soffice_path()
    if not exe:
        return None
    outdir = os.path.dirname(os.path.abspath(docx_path))
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path], check=True,
                   capture_output=True, timeout=timeout)
    pdf = os.path.splitext(os.path.abspath(docx_path))[0] + ".pdf"
    return pdf if os.path.exists(pdf) else None
