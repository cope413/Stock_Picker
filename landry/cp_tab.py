"""Current Positions: add a position row without openpyxl's ``insert_rows`` (added 2026-10-07).

The 9/30/26 GE / NFLX insertion used ``insert_rows``, which moves cell text and nothing else: the note's merge stayed behind
inside the grown Table, row heights stayed on their old rows, and every formula that named a row (block sums, the
Cash / Cash Equivalents list, the Excluding-Cash lines) kept pointing at the old ones -- Excel answered with a repair
prompt and the cash line read $396K instead of $469K for three days (CLAUDE.md, "Row surgery"). Every account block here
is a run of position rows, a SUBTOTAL row and an "Excluding Cash" row, and the Chase block, the combined totals and the
note sit below them, so adding one position to the Fidelity block shifts everything beneath it.

``insert_position`` does it the way Excel would: the rows from the insertion point down move down one (cells with their
styles, row heights, merges), every same-sheet A1 reference is rewritten by ``shift_formula`` (rows at or below the
insertion point move; a block-sum range ending on the row above it grows to take the new row in), the Table range grows,
and the new row's formulas are copied from the block's last row. References to other sheets and structured references
(``CurrentPositionsTable[[#This Row],...]``) are left alone; no other tab reads this one by A1 address (the audit checks).

It still leaves two things to the caller: the by-ticker repair of any formula that named a row by position (``fix_excluding_cash``
does the JT ULTRA one) and an open-in-Excel check, because LibreOffice and openpyxl accept files Excel rejects."""

from __future__ import annotations

import copy
import re
from typing import Dict, Optional, Tuple

from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils import get_column_letter

SHEET = "Current Positions"
TABLE = "CurrentPositionsTable"
FIRST_ROW = 3                      # first data row of the first block (row 2 is the Table header)

_REF = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)(?::(\$?)([A-Za-z]{1,3})(\$?)(\d+))?$")


class CurrentPositionsError(RuntimeError):
    """Raised instead of guessing: the sheet no longer has the structure this module was written against."""


def shift_formula(formula, at: int, block_first: int = FIRST_ROW):
    """The formula after a row is inserted at ``at`` on the same sheet.

    A cell or range edge on row ``at`` or below moves down one. A range whose last row is ``at - 1`` and whose first row
    is ``block_first`` or below grows to include the new row (the block-sum ranges: G3:G23 -> G3:G24). References that
    carry a sheet name or a table name are other sheets' or structured references and are returned unchanged."""
    if not (isinstance(formula, str) and formula.startswith("=")):
        return formula
    tok = Tokenizer(formula)
    for t in tok.items:
        if t.type != Token.OPERAND or t.subtype != Token.RANGE or "!" in t.value or "[" in t.value:
            continue
        m = _REF.match(t.value)
        if not m:
            continue
        a1, c1, a2, r1, b1, c2, b2, r2 = m.groups()
        r1 = int(r1)
        if r2 is None:                                                   # a single cell
            t.value = f"{a1}{c1}{a2}{r1 + 1 if r1 >= at else r1}"
            continue
        r2 = int(r2)
        new1 = r1 + 1 if r1 >= at else r1
        new2 = r2 + 1 if r2 >= at else r2
        if r2 == at - 1 and r1 >= block_first:                            # a block-sum range ending just above the new row
            new2 = at
        t.value = f"{a1}{c1}{a2}{new1}:{b1}{c2}{b2}{new2}"
    return tok.render()


def retarget(formula, from_row: int, to_row: int):
    """The formula with its same-sheet single-cell references to ``from_row`` pointed at ``to_row`` (a row's own cells:
    ``E23*F23`` -> ``E24*F24``, ``$B23`` -> ``$B24``). Ranges and other sheets' references are left alone."""
    if not (isinstance(formula, str) and formula.startswith("=")):
        return formula
    tok = Tokenizer(formula)
    for t in tok.items:
        if t.type != Token.OPERAND or t.subtype != Token.RANGE or "!" in t.value or "[" in t.value:
            continue
        m = _REF.match(t.value)
        if m and m.group(8) is None and int(m.group(4)) == from_row:
            a1, c1, a2, _, *_ = m.groups()
            t.value = f"{a1}{c1}{a2}{to_row}"
    return tok.render()


def block_bounds(ws, account: str) -> Tuple[int, int, int]:
    """(first data row, last data row, SUBTOTAL row) of one account's block, found by the account name in column A and the
    'Positions' / '... SUBTOTAL' row that closes it."""
    first = last = sub = None
    for r in range(FIRST_ROW, ws.max_row + 1):
        a, c = ws.cell(r, 1).value, ws.cell(r, 3).value
        if a == account and ws.cell(r, 2).value:
            first = first or r
            last = r
        elif a == "Positions" and isinstance(c, str) and c.startswith(account) and "SUBTOTAL" in c:
            sub = r
            break
    if first is None or sub is None:
        raise CurrentPositionsError(f"cannot find the {account!r} block (data rows {first}-{last}, subtotal {sub})")
    if sub != last + 1:
        raise CurrentPositionsError(f"the {account!r} subtotal is on row {sub} but its last position is on row {last}: "
                                    "there are rows in between this module does not know how to place a new position among")
    return first, last, sub


def find_row(ws, account: str, ticker: str) -> Optional[int]:
    for r in range(FIRST_ROW, ws.max_row + 1):
        if ws.cell(r, 1).value == account and ws.cell(r, 2).value == ticker:
            return r
    return None


def _insert_row(ws, at: int, block_first: int) -> None:
    """Make room: rows ``at`` and below move down one (cells with styles, heights, merges), formulas rewritten, Table grown."""
    straddling = [str(m) for m in ws.merged_cells.ranges if m.min_row < at <= m.max_row]
    if straddling:
        raise CurrentPositionsError(f"a merged range spans the insertion point: {straddling}")
    for rng in ws.conditional_formatting:
        for cr in str(rng.sqref).split():
            m = _REF.match(cr)
            if m and m.group(8) is not None and int(m.group(8)) < 1_000_000 and int(m.group(8)) >= at:
                raise CurrentPositionsError(f"conditional formatting {cr} reaches below the insertion point; extend it by hand")
    if ws.data_validations.dataValidation:
        raise CurrentPositionsError("this sheet has data validations; this module does not move them")
    below = [str(m) for m in ws.merged_cells.ranges if m.min_row >= at]
    for m in below:
        ws.unmerge_cells(m)
    old_max = ws.max_row
    heights = {r: ws.row_dimensions[r].height for r in range(at, old_max + 1)}
    ws.move_range(f"A{at}:{get_column_letter(ws.max_column)}{old_max}", rows=1, cols=0, translate=False)
    for r in range(old_max, at - 1, -1):                               # heights follow their rows; the new row gets the usual 15
        ws.row_dimensions[r + 1].height = heights[r]
    ws.row_dimensions[at].height = 15.0
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = shift_formula(c.value, at, block_first)
    for m in below:
        lo, hi = m.split(":")
        mm = _REF.match(f"{lo}:{hi}")
        ws.merge_cells(f"{mm.group(2)}{int(mm.group(4)) + 1}:{mm.group(6)}{int(mm.group(8)) + 1}")
    tbl = ws.tables[TABLE]
    lo, hi = tbl.ref.split(":")
    mm = _REF.match(f"{lo}:{hi}")
    if int(mm.group(4)) < at <= int(mm.group(8)):
        tbl.ref = f"{mm.group(2)}{mm.group(4)}:{mm.group(6)}{int(mm.group(8)) + 1}"


def insert_position(wb, account: str, ticker: str, description: str, asset_class: str, quantity: float, cost_basis: float,
                    price: float, note: Optional[str] = None) -> int:
    """Add one position at the end of ``account``'s block, above its SUBTOTAL row; returns the new row.

    The new row copies the block's last row (formulas retargeted to its own row, styles, the Monitor price lookup with
    ``price`` as its fallback number); the quantity, cost basis and text are typed. ``cost_basis`` is the cost of ALL the
    shares (Current Positions' column H), not per share."""
    ws = wb[SHEET]
    if find_row(ws, account, ticker):
        raise CurrentPositionsError(f"{ticker} already has a row in {account}")
    first, last, sub = block_bounds(ws, account)
    template = last
    _insert_row(ws, sub, first)
    new = sub                                                           # the template row is untouched above it
    for col in range(1, ws.max_column + 1):
        src, dst = ws.cell(template, col), ws.cell(new, col)
        dst._style = copy.copy(src._style)
        v = src.value
        if isinstance(v, str) and v.startswith("="):
            dst.value = retarget(v, template, new)
    f_old = ws.cell(template, 6).value
    m = re.search(r'="",([0-9.]+),INDEX', f_old or "")
    if not m:
        raise CurrentPositionsError(f"the Price formula on row {template} has an unexpected shape: {f_old!r}")
    ws.cell(new, 6).value = retarget(f_old, template, new).replace(f'="",{m.group(1)},INDEX', f'="",{price:g},INDEX', 1)
    ws.cell(new, 1).value = account
    ws.cell(new, 2).value = ticker
    ws.cell(new, 3).value = description
    ws.cell(new, 4).value = asset_class
    ws.cell(new, 5).value = quantity
    ws.cell(new, 8).value = cost_basis
    ws.cell(new, 13).value = note
    return new


def fix_excluding_cash(ws, account: str, cash_tickers: Tuple[str, ...]) -> Dict[str, str]:
    """Re-derive the block's "Excluding Cash" unrealized-% line BY TICKER: ``=I<sub>/(H<sub>-H<cash1>-H<cash2>...)``.

    The JT ULTRA line read ``=I24/(H24-H14-H13)`` after the 9/8/26 Visa insertion: H13 became Visa's cost basis, not
    FZFXX's, so the line subtracted a stock instead of the second cash fund (found 10/7/26 doing this surgery)."""
    first, last, sub = block_bounds(ws, account)
    rows = []
    for t in cash_tickers:
        r = find_row(ws, account, t)
        if r is None:
            raise CurrentPositionsError(f"{t} not found in {account}")
        rows.append(r)
    line = sub + 1
    if ws.cell(line, 9).value != "Excluding Cash":
        raise CurrentPositionsError(f"row {line} is not the Excluding-Cash line ({ws.cell(line, 9).value!r})")
    before = ws.cell(line, 10).value
    ws.cell(line, 10).value = f"=I{sub}/(H{sub}" + "".join(f"-H{r}" for r in rows) + ")"
    return {"was": before, "now": ws.cell(line, 10).value}
