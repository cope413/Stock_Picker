"""``landry accounts``: bring Current Positions' cash rows up to the brokerage accounts (Alan, 2026-10-10: "read the accts").

The System's two accounts -- Fidelity JT ULTRA (...3711) and Chase Self-Directed (...3693) -- move between statements:
dividends land, transfers go out, a tranche settles. Until now each catch-up was hand arithmetic over a pasted CSV. This
module does the same work from two inputs and shows its inputs beside every figure:

* a SNAPSHOT of what the account shows now (a small JSON file: the cash balance of each money fund, cash not yet swept
  into one, the account total, optionally the share count of each position), written from the brokerage's own page or
  typed from a download; and
* the TRANSACTIONS since the last booking (the Chase and Fidelity CSV exports, or a normalized list inside the snapshot).

It then (1) walks the cash -- the cash already booked plus every new transaction must equal the cash the account shows,
to the cent, or the difference is reported as unexplained and nothing is written; (2) books the observed cash, plus any
transfer Alan has announced that the broker does not show yet (a PENDING flow, labelled as such in the row's note); (3)
keeps the ledger of external deposits and withdrawals in ``landry_accounts.json``, which is what the Portfolio Drawdown
Log's value is net of (Part 7): Drawdown value = portfolio total + net external withdrawals; and (4) lists what it does
NOT do -- a buy, a sale or a reinvestment changes a position row and its Performance Tracking lot, which stay with
``cp_tab.insert_position`` / ``perf_tab``; a money fund crossing zero needs its lot opened or closed in the same pass.

Every figure written from here is an ESTIMATE until the monthly statement reconciles it (statements are the month-end
record; downloads and site reads are for sizing and interim values), and each row's note says so.

Snapshot file::

    {"as_of": "2026-10-10", "source": "site read",
     "fidelity": {"total": 400337.60, "cash": {"FZDXX": 669.69, "FZFXX": 0}, "unswept": 43.43,
                  "positions": {"VEEV": 100},
                  "open_orders": ["10/8 Sell 40 CRWD limit $285.60 (good 'til canceled)"]},
     "chase":    {"total": 298607.58, "cash": {"QACDS": 427.26, "VMFXX": 27966.40}},
     "transactions": [{"date": "2026-10-09", "account": "chase", "kind": "external", "symbol": "",
                       "quantity": 0, "amount": -20000.00, "description": "BANKLINK ACH PUSH"}]}

``open_orders`` lists every order the account shows as open or pending (one line each; an empty list means "looked, none").
They are never booked -- an open order has not filled -- but every report lists them, and the last read is kept so the
weekly and month-end summaries can include them (``landry accounts orders``).

``kind`` is one of ``external`` (money entering or leaving the account), ``income`` (dividends, interest, foreign tax,
fees), ``trade`` (buy / sell / reinvestment of a position), ``internal`` (sweeps and money-fund moves: cash to cash) or
``other`` (unrecognized: counted in the walk and reported).
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import os
from dataclasses import asdict, dataclass, field
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STATE_FILE = os.path.join(REPO, "landry_accounts.json")
SHEET = "Current Positions"
CENT = 0.005
KINDS = ("external", "income", "trade", "internal", "other")
PENDING_MATCH_DAYS = 10       # a pending transfer is the posted one if the same amount posts within this many days


@dataclass(frozen=True)
class Account:
    key: str
    sheet_name: str                   # column A of Current Positions
    number: str                       # as the broker's export writes it
    label: str
    cash_funds: Tuple[str, ...]
    core: str                         # where deposits and unswept cash land by default


ACCOUNTS: Dict[str, Account] = {
    "fidelity": Account("fidelity", "JT ULTRA (Fidelity)", "X64063711", "Fidelity JT ULTRA ...3711",
                        ("FZDXX", "FZFXX"), "FZFXX"),
    "chase": Account("chase", "Self-Directed (Chase)", "...3693", "Chase Self-Directed ...3693",
                     ("QACDS", "VMFXX"), "QACDS"),
}


class AccountsError(RuntimeError):
    pass


@dataclass(frozen=True)
class Txn:
    date: dt.date
    account: str
    kind: str
    symbol: str
    quantity: float
    amount: float
    description: str

    def key(self) -> str:
        return f"{self.date.isoformat()}|{self.kind}|{self.symbol}|{self.quantity:.5f}|{self.amount:.2f}|{self.description[:60]}"


def _num(v) -> float:
    s = str(v if v is not None else "").replace(",", "").replace("$", "").strip()
    if s in ("", "--"):
        return 0.0
    return float(s)


def _date(v) -> dt.date:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise AccountsError(f"cannot read the date {v!r}")


# ------------------------------------------------------------------ parsers --

_CHASE_KIND = {"BNK": "external", "DBS": "internal", "WDL": "internal", "DIVIDEND": "income", "INTEREST": "income",
               "FEE": "income", "TAX": "income"}


def parse_chase_csv(path: str) -> List[Txn]:
    """Chase's transactions export for the Self-Directed account (one account per file)."""
    acct = ACCOUNTS["chase"]
    out: List[Txn] = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            if not (row.get("Trade Date") or "").strip():
                continue
            number = (row.get("Account Number") or "").strip()
            if number and number[-4:] != acct.number[-4:]:
                continue
            typ = (row.get("Type") or "").strip().upper()
            symbol = (row.get("Ticker") or "").strip().upper()
            kind = _CHASE_KIND.get(typ)
            if kind is None:
                if typ in ("BUY", "SELL", "REINVEST"):
                    kind = "internal" if symbol in acct.cash_funds else "trade"
                else:
                    kind = "other"
            out.append(Txn(_date(row["Trade Date"]), "chase", kind, symbol, _num(row.get("Quantity")),
                           _num(row.get("Amount USD")), " ".join((row.get("Description") or "").split())))
    return out


_FIDELITY_EXTERNAL = ("ELECTRONIC FUNDS TRANSFER", "DIRECT DEBIT", "DIRECT DEPOSIT", "DEBIT CARD", "CHECK", "BILL PAYMENT",
                      "TRANSFERRED", "WIRE", "CASH ADVANCE", "DEPOSIT")
_FIDELITY_INCOME = ("DIVIDEND RECEIVED", "INTEREST", "FOREIGN TAX PAID", "FEE", "LONG-TERM CAP GAIN", "SHORT-TERM CAP GAIN",
                    "RETURN OF CAPITAL", "MARGIN INTEREST")
_FIDELITY_INTERNAL = ("REDEMPTION FROM CORE ACCOUNT", "PURCHASE INTO CORE ACCOUNT", "JOURNALED JNL VS A/C TYPES")
_FIDELITY_TRADE = ("YOU BOUGHT", "YOU SOLD", "REINVESTMENT")


def _fidelity_kind(action: str, symbol: str, acct: Account) -> str:
    a = action.upper()
    if a.startswith(_FIDELITY_INTERNAL):
        return "internal"
    if a.startswith(_FIDELITY_TRADE):
        return "internal" if symbol in acct.cash_funds else "trade"
    if a.startswith(_FIDELITY_INCOME):
        return "income"
    if a.startswith(_FIDELITY_EXTERNAL):
        return "external"
    return "other"


def parse_fidelity_csv(path: str) -> List[Txn]:
    """Fidelity's Accounts History export. It lists every household account; only the System's is read."""
    acct = ACCOUNTS["fidelity"]
    out: List[Txn] = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        lines = f.read().splitlines()
    start = next((i for i, ln in enumerate(lines) if ln.startswith("Run Date,")), None)
    if start is None:
        raise AccountsError(f"{os.path.basename(path)}: no 'Run Date' header -- not a Fidelity Accounts History export")
    for row in csv.DictReader(lines[start:]):
        if (row.get("Account Number") or "").strip() != acct.number:
            continue
        action = " ".join((row.get("Action") or "").split())
        symbol = (row.get("Symbol") or "").strip().upper()
        out.append(Txn(_date(row["Run Date"]), "fidelity", _fidelity_kind(action, symbol, acct), symbol,
                       _num(row.get("Quantity")), _num(row.get("Amount ($)")), action))
    return out


def _snapshot_txns(snap: dict) -> List[Txn]:
    out = []
    for t in snap.get("transactions") or []:
        if t.get("account") not in ACCOUNTS:
            raise AccountsError(f"snapshot transaction for an unknown account: {t!r}")
        if t.get("kind") not in KINDS:
            raise AccountsError(f"snapshot transaction kind must be one of {KINDS}: {t!r}")
        out.append(Txn(_date(t["date"]), t["account"], t["kind"], str(t.get("symbol") or "").upper(),
                       _num(t.get("quantity")), _num(t.get("amount")), " ".join(str(t.get("description") or "").split())))
    return out


# -------------------------------------------------------------------- state --

def load_state(path: str = STATE_FILE) -> dict:
    if not os.path.exists(path):
        raise AccountsError(f"{os.path.basename(path)} is missing: it holds the external-flow ledger and what has been booked")
    with open(path) as f:
        return json.load(f)


def save_state(state: dict, path: str = STATE_FILE) -> None:
    tmp = path + ".tmp"
    with open(tmp, "w") as f:
        json.dump(state, f, indent=1)
        f.write("\n")
    os.replace(tmp, path)


def net_external_withdrawals(state: dict) -> float:
    """What has left the two accounts, net of what entered, on the ledger's basis (posted and pending alike)."""
    return round(-sum(f["amount"] for f in state["flows"]), 2)


def pending_for(state: dict, account: str) -> float:
    return round(sum(f["amount"] for f in state["flows"] if f["account"] == account and f.get("status") == "pending"), 2)


def add_pending(state: dict, account: str, amount: float, date: dt.date, note: str) -> dict:
    if account not in ACCOUNTS:
        raise AccountsError(f"unknown account {account!r} (have: {', '.join(ACCOUNTS)})")
    if abs(amount) < CENT:
        raise AccountsError("a pending flow needs a non-zero amount (positive = into the account)")
    flow = {"date": date.isoformat(), "account": account, "amount": round(amount, 2), "status": "pending",
            "note": " ".join(note.split())}
    state["flows"].append(flow)
    return flow


def remove_pending(state: dict, account: str, amount: float) -> dict:
    for f in state["flows"]:
        if f["account"] == account and f.get("status") == "pending" and abs(f["amount"] - amount) < CENT:
            state["flows"].remove(f)
            return f
    raise AccountsError(f"no pending flow of {amount:,.2f} on {account}")


def _unbooked(state: dict, account: str, txns: Iterable[Txn]) -> List[Txn]:
    """Transactions on or after the account's booked-through date that an earlier booking did not already count."""
    st = state["accounts"][account]
    through = _date(st["booked_through"])
    seen = list(st.get("booked_keys") or [])
    out = []
    for t in sorted((t for t in txns if t.account == account and t.date >= through), key=lambda t: t.date):
        k = t.key()
        if k in seen:
            seen.remove(k)          # identical twins on one day: each booked key hides one of them
        else:
            out.append(t)
    return out


# --------------------------------------------------------------- reconcile --

@dataclass
class AccountResult:
    account: str
    booked_cash: Dict[str, float]            # Current Positions now, per money fund
    booked_pending: float                    # pending flows inside that figure
    observed_cash: Dict[str, float]
    unswept: float
    new_txns: List[Txn]
    walk_expected: Optional[float]           # None: no transactions supplied
    unexplained: Optional[float]
    pending: float                           # pending flows to carry in the new figure
    target_cash: Dict[str, float]            # what the rows should read
    posted_flows: List[Txn] = field(default_factory=list)
    cleared_pending: List[dict] = field(default_factory=list)
    trades: List[Txn] = field(default_factory=list)
    others: List[Txn] = field(default_factory=list)
    crossings: List[str] = field(default_factory=list)
    quantity_diffs: List[Tuple[str, float, float]] = field(default_factory=list)
    observed_total: Optional[float] = None
    projected_total: Optional[float] = None
    open_orders: Optional[List[str]] = None  # None: the snapshot did not say

    @property
    def attention(self) -> List[str]:
        a = self.label
        out = []
        if self.unexplained is not None and abs(self.unexplained) >= CENT:
            out.append(f"{a}: cash walk is off by ${self.unexplained:,.2f} (the account shows more than the transactions "
                       f"explain if positive) -- nothing booked for this account")
        for t in self.trades:
            out.append(f"{a}: {t.date:%m/%d} {t.description[:70]} ({t.symbol} {t.quantity:+g}, ${t.amount:,.2f}) changes a "
                       f"position row and its Performance Tracking lot -- NOT booked by this command")
        for t in self.others:
            out.append(f"{a}: {t.date:%m/%d} unrecognized transaction '{t.description[:70]}' ${t.amount:,.2f} -- counted "
                       f"in the cash walk; check whether it is an external flow")
        for c in self.crossings:
            out.append(f"{a}: {c}")
        for tk, have, seen in self.quantity_diffs:
            out.append(f"{a}: {tk} is {have:g} shares on Current Positions but {seen:g} at the broker")
        return out

    @property
    def label(self) -> str:
        return ACCOUNTS[self.account].label

    @property
    def writable(self) -> bool:
        return self.unexplained is None or abs(self.unexplained) < CENT


def _rows(ws, acct: Account) -> Dict[str, int]:
    return {str(ws.cell(r, 2).value).upper(): r for r in range(3, ws.max_row + 1)
            if ws.cell(r, 1).value == acct.sheet_name and ws.cell(r, 2).value}


def _pending_fund(acct: Account, booked: Dict[str, float], observed: Dict[str, float]) -> str:
    """The fund that carries pending and unswept cash: the core fund when it is in use, else the fund holding the cash
    (a zero row has a closed Performance Tracking lot, so nothing unconfirmed is parked on one)."""
    if abs(booked.get(acct.core, 0.0)) >= CENT or abs(observed.get(acct.core, 0.0)) >= CENT:
        return acct.core
    return max(acct.cash_funds, key=lambda f: (observed.get(f, 0.0), booked.get(f, 0.0)))


def reconcile(ws_formulas, ws_values, state: dict, snap: dict, txns: Sequence[Txn],
              covered: Optional[Iterable[str]] = None) -> List[AccountResult]:
    """Compare each account in the snapshot with Current Positions. Pure: reads two views of the sheet (formulas for the
    typed quantities, cached values for market values) and writes nothing."""
    as_of = _date(snap["as_of"])
    covered = set(ACCOUNTS if covered is None else covered)      # accounts whose transactions were supplied
    results = []
    for key, acct in ACCOUNTS.items():
        obs = snap.get(key)
        if not obs:
            continue
        rows = _rows(ws_formulas, acct)
        missing = [f for f in acct.cash_funds if f not in rows]
        if missing:
            raise AccountsError(f"{acct.sheet_name}: no Current Positions row for {', '.join(missing)}")
        unknown = [f for f in (obs.get("cash") or {}) if f.upper() not in acct.cash_funds]
        if unknown:
            raise AccountsError(f"{acct.label}: {', '.join(unknown)} is not one of its money funds {acct.cash_funds}")
        booked = {f: _num(ws_formulas.cell(rows[f], 5).value) for f in acct.cash_funds}
        observed = {f: 0.0 for f in acct.cash_funds}
        observed.update({f.upper(): round(_num(v), 2) for f, v in (obs.get("cash") or {}).items()})
        unswept = round(_num(obs.get("unswept")), 2)
        booked_pending = round(float(state["accounts"][key].get("booked_pending") or 0.0), 2)

        new = _unbooked(state, key, txns)
        posted = [t for t in new if t.kind == "external"]
        cleared, still_pending = [], []
        unmatched = list(posted)
        for f in state["flows"]:
            if f["account"] != key or f.get("status") != "pending":
                continue
            hit = next((t for t in unmatched if abs(t.amount - f["amount"]) < CENT
                        and 0 <= (t.date - _date(f["date"])).days <= PENDING_MATCH_DAYS), None)
            if hit:
                unmatched.remove(hit)
                cleared.append(f)
            else:
                still_pending.append(f)
        pending = round(sum(f["amount"] for f in still_pending), 2)

        walk = unexplained = None
        if key in covered:
            moved = sum(t.amount for t in new if t.kind != "internal")
            walk = round(sum(booked.values()) - booked_pending + moved, 2)
            unexplained = round(sum(observed.values()) + unswept - walk, 2)

        fund = _pending_fund(acct, booked, observed)
        target = dict(observed)
        target[fund] = round(target[fund] + unswept + pending, 2)

        res = AccountResult(key, booked, booked_pending, observed, unswept, new, walk, unexplained, pending, target,
                            posted_flows=posted, cleared_pending=cleared,
                            trades=[t for t in new if t.kind == "trade"], others=[t for t in new if t.kind == "other"])
        for f in acct.cash_funds:
            if abs(booked[f]) < CENT <= abs(target[f]):
                res.crossings.append(f"{f} goes from zero to ${target[f]:,.2f}: open its Performance Tracking lot in the same pass")
            elif abs(target[f]) < CENT <= abs(booked[f]):
                res.crossings.append(f"{f} goes from ${booked[f]:,.2f} to zero: close its Performance Tracking lot in the same pass")
        for tk, qty in (obs.get("positions") or {}).items():
            tk = tk.upper()
            if tk in acct.cash_funds:
                continue
            have = _num(ws_formulas.cell(rows[tk], 5).value) if tk in rows else 0.0
            if abs(have - _num(qty)) > 0.0005:
                res.quantity_diffs.append((tk, have, _num(qty)))
        if obs.get("open_orders") is not None:
            res.open_orders = [" ".join(str(o).split()) for o in obs["open_orders"]]
        if obs.get("total") is not None:
            res.observed_total = round(_num(obs["total"]), 2)
            others = sum(_num(ws_values.cell(r, 7).value) for tk, r in rows.items() if tk not in acct.cash_funds)
            res.projected_total = round(others + sum(target.values()), 2)
        results.append(res)
    if not results:
        raise AccountsError("the snapshot names neither account (expected a 'fidelity' and/or 'chase' block)")
    return results


def _note(res: AccountResult, fund: str, as_of: dt.date, source: str, state: dict) -> str:
    acct = ACCOUNTS[res.account]
    d = f"{as_of.month}/{as_of.day}/{as_of:%y}"
    parts = [f"{d}: ${res.target_cash[fund]:,.2f}, estimate ({source}, {d}) until the monthly statement reconciles it."]
    extra = round(res.target_cash[fund] - res.observed_cash[fund], 2)
    if abs(extra) >= CENT:
        bits = [f"the account shows ${res.observed_cash[fund]:,.2f}"]
        if abs(res.unswept) >= CENT:
            bits.append(f"${res.unswept:,.2f} of cash not yet swept in")
        for f in state["flows"]:
            if f["account"] == res.account and f.get("status") == "pending":
                fd = _date(f["date"])
                bits.append(f"a PENDING transfer of ${f['amount']:,.2f} ({fd.month}/{fd.day}, announced by Alan, not yet "
                            f"posted at {acct.label.split()[0]})")
        parts.append("Made of " + " + ".join(bits) + ".")
    parts.append("Booked by `landry accounts`.")
    return " ".join(parts)


def apply(path: str, snap: dict, txns: Sequence[Txn], state_path: str = STATE_FILE, *,
          covered: Optional[Iterable[str]] = None, recalc: bool = True) -> Tuple[List[AccountResult], dict]:
    """Write the cash rows (quantity, cost basis, note) for every account whose cash walk closes, record posted external
    flows, retire pending ones that posted, and move each account's booked-through marker. One save, one recalc."""
    import openpyxl

    from landry import ledger, xlsx_recalc

    if ledger._excel_has_open(path):
        raise AccountsError("Excel has the workbook open; close it and retry")
    state = load_state(state_path)
    as_of = _date(snap["as_of"])
    source = str(snap.get("source") or "download")
    wbv = openpyxl.load_workbook(path, data_only=True)
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[SHEET]
        covered = set(ACCOUNTS if covered is None else covered)
        results = reconcile(ws, wbv[SHEET], state, snap, txns, covered)
        wrote = False
        for res in results:
            if not res.writable:
                continue
            acct = ACCOUNTS[res.account]
            st = state["accounts"][res.account]
            for f in res.cleared_pending:
                state["flows"].remove(f)
            for t in res.posted_flows:
                state["flows"].append({"date": t.date.isoformat(), "account": res.account, "amount": round(t.amount, 2),
                                       "status": "posted", "note": t.description[:80]})
            rows = _rows(ws, acct)
            for fund in acct.cash_funds:
                r = rows[fund]
                changed = abs(res.target_cash[fund] - res.booked_cash[fund]) >= CENT
                ws.cell(r, 5).value = res.target_cash[fund]
                ws.cell(r, 8).value = res.target_cash[fund]
                if changed or abs(res.target_cash[fund] - res.observed_cash[fund]) >= CENT:
                    ws.cell(r, 13).value = _note(res, fund, as_of, source, state)
                wrote = wrote or changed
            st["booked_pending"] = res.pending
            if res.account in covered:
                st["booked_through"] = as_of.isoformat()
                st["booked_keys"] = [t.key() for t in txns if t.account == res.account and t.date >= as_of]
            st["last_read"] = {"as_of": as_of.isoformat(), "source": source, "total": res.observed_total,
                               "cash": res.observed_cash, "unswept": res.unswept, "open_orders": res.open_orders}
        if wrote:
            wb.save(path)
    finally:
        wb.close()
        wbv.close()
    save_state(state, state_path)
    if wrote and recalc:
        xlsx_recalc.recalc(path)
    return results, state


# ------------------------------------------------------------------- report --

def format_report(results: Sequence[AccountResult], state: dict, snap: dict, *, applied: bool,
                  portfolio_total: Optional[float] = None) -> str:
    as_of = _date(snap["as_of"])
    out = [f"Accounts -- {snap.get('source') or 'download'} as of {as_of:%Y-%m-%d} ({'BOOKED' if applied else 'check only, nothing written'})"]
    attention: List[str] = []
    for res in results:
        acct = ACCOUNTS[res.account]
        out.append("")
        out.append(acct.label)
        for f in acct.cash_funds:
            out.append(f"  {f:<6} booked {res.booked_cash[f]:>12,.2f}   account shows {res.observed_cash[f]:>12,.2f}   "
                       f"-> {res.target_cash[f]:>12,.2f}")
        if abs(res.unswept) >= CENT:
            out.append(f"  + {res.unswept:,.2f} cash not yet swept into a fund")
        if abs(res.pending) >= CENT:
            out.append(f"  + {res.pending:,.2f} PENDING transfer(s) announced, not yet posted (carried as an estimate)")
        for f in res.cleared_pending:
            out.append(f"  pending {f['amount']:,.2f} of {f['date']} has POSTED: now in the account's own balance")
        if res.walk_expected is None:
            out.append("  cash walk: no transactions supplied -- balance booked as shown, external flows NOT checked")
        else:
            n = len([t for t in res.new_txns if t.kind != "internal"])
            out.append(f"  cash walk: booked {sum(res.booked_cash.values()) - res.booked_pending:,.2f} + {n} new transaction(s) "
                       f"{sum(t.amount for t in res.new_txns if t.kind != 'internal'):+,.2f} = {res.walk_expected:,.2f}; "
                       f"account shows {sum(res.observed_cash.values()) + res.unswept:,.2f}; "
                       f"unexplained {res.unexplained:+,.2f}")
            for t in res.new_txns:
                if t.kind != "internal":
                    out.append(f"    {t.date:%m/%d} {t.kind:<8} {t.amount:>12,.2f}  {(t.symbol + ' ' if t.symbol else '')}{t.description[:60]}")
        if res.observed_total is not None:
            carried = res.pending
            out.append(f"  account total: broker {res.observed_total:,.2f}"
                       + (f" + pending {carried:,.2f} = {res.observed_total + carried:,.2f}" if abs(carried) >= CENT else "")
                       + f"; workbook {res.projected_total:,.2f} "
                       f"(difference {res.projected_total - res.observed_total - carried:+,.2f}, prices as of Market Data)")
        if res.open_orders is None:
            out.append("  open orders: not read")
        elif not res.open_orders:
            out.append("  open orders: none")
        else:
            out.append(f"  open orders ({len(res.open_orders)}, not filled, nothing booked):")
            out += [f"    {o}" for o in res.open_orders]
        attention += res.attention
    addback = net_external_withdrawals(state)
    if not applied:               # what the ledger will read once these are booked
        for res in results:
            if res.writable:
                addback = round(addback - sum(t.amount for t in res.posted_flows)
                                + sum(f["amount"] for f in res.cleared_pending), 2)
    out.append("")
    out.append(f"External flows (ledger basis: {state.get('basis', 'see landry_accounts.json')}): "
               f"net withdrawals {addback:,.2f}" + (" (includes pending)" if any(f.get('status') == 'pending' for f in state['flows']) else ""))
    if portfolio_total is not None:
        out.append(f"Drawdown Log value = portfolio total {portfolio_total:,.2f} + {addback:,.2f} = {portfolio_total + addback:,.2f} "
                   f"(prices as of the last Market Data refresh; log it with `landry drawdown add` after a close)")
    if attention:
        out.append("")
        out.append("NEEDS ATTENTION")
        out += ["  " + a for a in attention]
    return "\n".join(out)


def format_flows(state: dict) -> str:
    out = [f"External flows -- {state.get('basis', '')}".rstrip(" -")]
    for f in sorted(state["flows"], key=lambda f: f["date"]):
        out.append(f"  {f['date']}  {ACCOUNTS[f['account']].label:<28} {f['amount']:>12,.2f}  {f.get('status', 'posted'):<8} {f.get('note', '')}")
    out.append(f"  net external withdrawals: {net_external_withdrawals(state):,.2f}")
    return "\n".join(out)


def format_orders(state: dict) -> str:
    """Open / pending orders as of each account's last booked read."""
    out = ["Open orders (as of each account's last read; not filled, nothing booked)"]
    for key, acct in ACCOUNTS.items():
        last = state["accounts"][key].get("last_read") or {}
        orders = last.get("open_orders")
        head = f"  {acct.label} ({last.get('as_of', 'never read')}):"
        if orders is None:
            out.append(head + " not read")
        elif not orders:
            out.append(head + " none")
        else:
            out.append(head)
            out += [f"    {o}" for o in orders]
    return "\n".join(out)


def load_inputs(snapshot: str, chase_csv: Optional[str], fidelity_csv: Optional[str]) -> Tuple[dict, List[Txn], set]:
    """(snapshot, transactions, the accounts whose transactions were supplied). A snapshot's own ``transactions`` list
    covers the accounts in ``transactions_cover`` (default: every account the snapshot has a block for)."""
    with open(snapshot) as f:
        snap = json.load(f)
    if "as_of" not in snap:
        raise AccountsError("the snapshot needs an 'as_of' date")
    txns = _snapshot_txns(snap)
    if chase_csv:
        txns += parse_chase_csv(chase_csv)
    if fidelity_csv:
        txns += parse_fidelity_csv(fidelity_csv)
    covered = set()
    if "transactions" in snap:
        covered |= set(snap.get("transactions_cover") or [k for k in ACCOUNTS if snap.get(k)])
    if chase_csv:
        covered.add("chase")
    if fidelity_csv:
        covered.add("fidelity")
    return snap, txns, covered
