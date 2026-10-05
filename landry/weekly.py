"""The weekly Friday-close routine, as one command (added 2026-10-04).

``python -m landry weekly`` -- run by the Saturday scheduled task, or by hand after any Friday close:

1. Price History: append every completed Friday that is not in the tab yet (``prices.append_weeks``,
   which also extends Returns (Calc) one row ahead and moves the footer and the charts);
2. Market Data and the Monitor tab's earnings dates: refresh in place (``market.refresh``), then keep Current
   Positions' fallback prices equal to Market Data's (``market.sync_fallbacks``) and write Performance Tracking's
   SPY benchmark (``market.sync_performance``; the lots' own prices there are Market Data lookups);
3. ONE LibreOffice recalc for both -- a workbook saved by openpyxl is not safe to commit without it;
4. read Rule 38's status back from the Correlation Matrix, run ``landry audit`` and the Journal / Drawdown Log
   drift guard (``landry db status``) -- an Excel hand-edit of either generated tab shows up here weekly.

It never commits and never touches another tab. The CLI layer refuses to start while Excel has the
workbook open. A problem in one step does not stop the other (each only ever saves a finished
result), but any problem makes the exit status non-zero, so a scheduled run cannot report success
over it. The functions here return data; ``format_report`` is the only thing that words it.
"""

from __future__ import annotations

import datetime as dt
import os
from typing import Callable, Dict, List, Optional

from landry import market, prices

_FIELD_LABEL = {"price": "price", "volume": "volume", "market_cap_m": "market cap", "pe": "P/E",
                "wk52_low": "52-wk low", "wk52_high": "52-wk high", "dividend_yield": "yield"}
_MAX_LINES = 40

# Audit findings about the Scoring tab's provenance (landry_scores.json backing) are a known, slowly-shrinking
# backlog (CLAUDE.md, "Standing gap"). This routine never touches the Scoring tab, so they cannot be caused
# by it: they are reported, but they do not make a run "need attention" -- otherwise every run would.
STANDING_AUDIT = ("scoring_verification",)


# --------------------------------------------------------------------- run --

def _weekend_reference(path: str, now: Optional[dt.datetime]) -> Optional[Dict[str, float]]:
    """Newest Price History closes -- but only when a quote should equal them: the run is on a weekend
    and the newest row is the Friday that has just closed. Any other day the quote has moved on."""
    from zoneinfo import ZoneInfo
    et_zone = ZoneInfo("America/New_York")
    n = now or dt.datetime.now(et_zone)
    et = n.astimezone(et_zone) if n.tzinfo else n.replace(tzinfo=et_zone)
    if et.weekday() not in (5, 6):
        return None
    try:
        last, closes = prices.latest_closes(path)
    except prices.PricesError:
        return None
    return closes if last == prices.last_completed_friday(now) else None


def run(path: str, *, write: bool = True, now: Optional[dt.datetime] = None, allow_big_moves: bool = False,
        do_prices: bool = True, do_market: bool = True, positions_only: bool = False,
        fetch_weeks: Optional[Callable] = None, snapshot: Optional[Callable] = None,
        earnings: Optional[Callable] = None, sleep: Optional[Callable] = None) -> dict:
    """Steps 1 and 2. Returns {'prices', 'market', 'problems', 'warnings', 'wrote', 'dry_run'}.
    ``positions_only`` skips the network: it only brings Current Positions' fallback prices into line with
    the Market Data tab as it stands."""
    rep: dict = {"prices": None, "market": None, "problems": [], "warnings": [], "wrote": False,
                 "dry_run": not write, "positions_only": positions_only}

    if do_prices:
        extra = {"fetch": fetch_weeks} if fetch_weeks else {}
        try:
            p = prices.append_weeks(path, write=write, force=allow_big_moves, now=now, **extra)
        except prices.PricesError as e:
            rep["problems"].append(f"Price History: {e}")
        else:
            rep["prices"] = p
            rep["warnings"] += [f"Price History: {w}" for w in p.get("warnings", [])]
            rep["wrote"] = rep["wrote"] or (write and bool(p["appended"]))

    if do_market:
        reference = None
        if write or (rep["prices"] and rep["prices"]["up_to_date"]):
            reference = _weekend_reference(path, now)
        extra = {"sleep": sleep} if sleep else {}
        try:
            m = market.refresh(path, snapshot=snapshot, earnings=earnings, write=write, reference=reference,
                               today=now.date() if now else None, do_market=not positions_only,
                               do_earnings=not positions_only, **extra)
        except market.MarketError as e:
            rep["problems"].append(f"Market Data: {e}")
        else:
            rep["market"] = m
            md = m["market"]
            failed = md["failed"]
            if md["rows"] and not md["refreshed"]:
                rep["problems"].append(f"Market Data: no quote came back for any of {md['rows']} tickers "
                                       f"(network or provider down?) -- nothing refreshed")
            elif len(failed) > max(2, md["rows"] // 4):
                rep["problems"].append(f"Market Data: no quote for {len(failed)} of {md['rows']} tickers: "
                                       f"{', '.join(failed)}")
            elif failed:
                rep["warnings"].append(f"Market Data: no quote for {', '.join(failed)} -- their rows keep last "
                                       f"week's values")
            rep["warnings"] += [f"Market Data: {w}" for w in md["warnings"]]
            pos = m["positions"]
            rep["warnings"] += [f"Current Positions: no Market Data price for {t} (row {r}) -- its fallback price stays stale"
                                for r, t in pos["no_price"]]
            rep["warnings"] += [f"Current Positions: the price formula in row {r} ({t}) is not in the expected shape -- left alone"
                                for r, t in pos["odd"]]
            perf = m["performance"]
            rep["warnings"] += [f"Performance Tracking: no Market Data price for {t} -- its held lot shows no Current Price"
                                for t in perf["no_price"]]
            if perf.get("spy_failed"):
                rep["warnings"].append("Performance Tracking: no SPY quote came back -- the benchmark keeps its last value")
            rep["warnings"] += [f"Performance Tracking: {w}" for w in perf["warnings"]]
            rep["wrote"] = rep["wrote"] or m["wrote"]
    return rep


# ------------------------------------------------------------------ verify --

def rule38_status(path: str) -> Optional[dict]:
    """The Correlation Matrix's own tally: positions over the 0.70 cap (J2) and the status text (J3)."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            ws = wb[prices.CM_SHEET]
            return {"over_cap": ws["J2"].value, "status": ws["J3"].value}
        finally:
            wb.close()
    except Exception:
        return None


def database_status(path: str, db_path: Optional[str] = None) -> dict:
    """The Journal / Drawdown Log drift guard (``landry db status``) as data. ``ok`` is True (the tabs and the
    database agree), False (they differ, or the check could not run) or None (no database yet: a fresh checkout
    builds it on its first write, so that is a note, not a failure)."""
    from landry import ledger, models
    db = db_path or models.DEFAULT_DB_PATH
    if not os.path.exists(db):
        return {"ok": None, "note": f"no database at {os.path.basename(db)} yet (`landry db pull` builds it from the workbook)"}
    try:
        report = ledger.status(path, db)
    except Exception as e:
        return {"ok": False, "error": str(e)}
    tabs = {k: {"rows": v["db_rows"], "value_diffs": len(v["value_diffs"]), "format_diffs": len(v["format_diffs"])}
            for k, v in report.items()}
    return {"ok": not any(t["value_diffs"] for t in tabs.values()), "tabs": tabs}


def _recalc(path: str) -> dict:
    from landry.xlsx_recalc import recalc
    return recalc(path)


def _audit(path: str, repo_dir: Optional[str]):
    from landry.audit import run_all
    return run_all(path, repo_dir=repo_dir)


def verify(path: str, rep: dict, *, recalc_fn: Optional[Callable] = None, audit_fn: Optional[Callable] = None,
           repo_dir: Optional[str] = None, recalc_now: bool = True, db_fn: Optional[Callable] = None,
           db_path: Optional[str] = None) -> dict:
    """Step 3 and 4: one recalc if anything was written, then Rule 38, the database drift guard and the
    audit. Mutates and returns ``rep``."""
    if rep["wrote"] and not recalc_now:
        rep["warnings"].append("recalc skipped (--no-recalc): the workbook is NOT safe to commit until it is run")
    if rep["wrote"] and recalc_now:
        res = (recalc_fn or _recalc)(path)
        rep["recalc"] = res
        if res.get("error"):
            rep["problems"].append(f"recalc failed: {res['error']} -- the workbook is NOT safe to commit")
        elif res.get("total_errors"):
            rep["problems"].append(f"recalc found {res['total_errors']} formula error(s): "
                                   f"{', '.join(sorted(res.get('error_summary', {})))}")
        if res.get("clamped_rows"):
            rep["problems"].append(f"recalc capped {len(res['clamped_rows'])} row height(s) at Excel's 409.5pt maximum")
    rep["rule38"] = rule38_status(path)
    rep["db"] = (db_fn or database_status)(path, db_path)
    if rep["db"].get("ok") is False:
        detail = rep["db"].get("error") or ", ".join(f"{k}: {t['value_diffs']} cell(s) differ" for k, t in rep["db"]["tabs"].items() if t["value_diffs"])
        rep["problems"].append(f"database: the Journal / Drawdown Log tab and landry.db disagree ({detail}) -- `python -m landry db status`; "
                               f"`db pull` if the workbook is right, `db regenerate` if the database is")
    try:
        checks = (audit_fn or _audit)(path, repo_dir)
    except Exception as e:                      # the audit must never be the reason a run says nothing
        rep["audit"] = None
        rep["problems"].append(f"audit could not run: {e}")
        return rep
    failed = [(c.name, c.detail) for c in checks if not c.ok and not c.name.startswith(STANDING_AUDIT)]
    standing = [(c.name, c.detail) for c in checks if not c.ok and c.name.startswith(STANDING_AUDIT)]
    rep["audit"] = {"checks": len(checks), "failed": failed, "standing": standing}
    if failed:
        rep["problems"].append(f"audit: {len(failed)} of {len(checks)} checks failed ({', '.join(n for n, _ in failed)})")
    return rep


# ------------------------------------------------------------------ report --

def _cap(lines: List[str], indent: str, limit: int = _MAX_LINES) -> List[str]:
    out = [indent + s for s in lines[:limit]]
    if len(lines) > limit:
        out.append(f"{indent}... and {len(lines) - limit} more")
    return out


def _positions_lines(pos: dict, rep: dict) -> List[str]:
    """The Current Positions fallback-price line, grouped by ticker (a ticker held in two accounts has two rows)."""
    if not pos["rows"]:
        return []
    by_ticker: Dict[str, list] = {}
    for t, _r, old, new in pos["updated"]:
        by_ticker.setdefault(t, []).append((old, new))
    verb = "would update" if rep.get("dry_run") else "updated"
    out = [f"Positions       fallback prices (Current Positions col F): {len(pos['updated'])} {verb} "
           f"across {len(by_ticker)} tickers, {pos['unchanged']} already current"]
    out += _cap([f"{t} {olds[0][0]:g} -> {olds[0][1]:g}" + (f" ({len(olds)} rows)" if len(olds) > 1 else "")
                 for t, olds in by_ticker.items()], "                ", 12)
    return out


def _performance_lines(perf: Optional[dict], rep: dict) -> List[str]:
    """The Performance Tracking line: the SPY benchmark (the only cells the routine writes there) and how many held
    lots price from Market Data."""
    if not perf or not perf.get("present"):
        return []

    def usd(v):
        return f"${v:,.2f}" if isinstance(v, (int, float)) else "(blank)"

    def day(d):
        return f"{d:%Y-%m-%d}" if isinstance(d, dt.date) else "(blank)"
    held = f"{perf['held']} held lots price from Market Data"
    if rep.get("positions_only"):
        spy = "SPY not refreshed (positions only)"
    elif perf.get("spy_failed"):
        spy = "SPY NOT refreshed (no quote)"
    elif perf["changed"]:
        spy = (f"SPY {'would be' if rep.get('dry_run') else 'now'} {usd(perf['spy_new'])} as of {day(perf['as_of_new'])} "
               f"(was {usd(perf['spy_old'])} as of {day(perf['as_of_old'])})")
    else:
        spy = f"SPY {usd(perf['spy_new'])} as of {day(perf['as_of_new'])}, unchanged"
    return [f"Performance     {spy}; {held}"]


def format_report(rep: dict, *, workbook: str = "", when: Optional[dt.datetime] = None) -> str:
    when = when or dt.datetime.now().astimezone()
    head = f"Weekly routine -- {when:%A %Y-%m-%d %H:%M %Z}" + (f" ({workbook})" if workbook else "")
    lines = [head]
    if rep.get("dry_run"):
        lines.append("DRY RUN: nothing was written, recalculated or saved.")
    lines.append("")

    p = rep.get("prices")
    if p is None:
        lines.append("Price History   not run" if not any(s.startswith("Price History") for s in rep["problems"])
                     else "Price History   FAILED (see below)")
    elif p["up_to_date"]:
        lines.append(f"Price History   up to date through {p['last_date_before']} (no completed Friday is missing)")
    else:
        verb = "would append" if rep.get("dry_run") else "appended"
        lines.append(f"Price History   {verb} {', '.join(p['appended'])}")

    m = rep.get("market")
    if m is None:
        lines.append("Market Data     not run" if not any(s.startswith("Market Data") for s in rep["problems"])
                     else "Market Data     FAILED (see below)")
    elif rep.get("positions_only"):
        lines.append("Market Data     not refreshed (positions only: prices are read from the tab as it stands)")
        lines += _positions_lines(m["positions"], rep)
        lines += _performance_lines(m.get("performance"), rep)
    else:
        md, em = m["market"], m["earnings"]
        verb = "would change" if rep.get("dry_run") else "changed"
        lines.append(f"Market Data     {md['refreshed']} of {md['rows']} tickers refreshed; "
                     f"{verb} {md['cells_changed']} cells in {md['rows_changed']} rows")
        if md["cleared"]:
            lines.append("                cleared (the provider reports nothing paid in the last 12 months): "
                         + "; ".join(f"{t} {_FIELD_LABEL.get(f, f)}" for t, f in md["cleared"]))
        if md["kept"]:
            lines.append("                kept as they were (no figure from the provider): "
                         + "; ".join(f"{t} {_FIELD_LABEL.get(f, f)}" for t, f in md["kept"]))
        kept = [f"{t}: {why}" for t, why in em["kept"]]
        lines.append(f"Earnings dates  {em['rows']} tickers: {len(em['updated'])} "
                     f"{'to update' if rep.get('dry_run') else 'updated'}, {em['unchanged']} unchanged, "
                     f"{len(kept)} kept, {len(em['no_date'])} with no date from the provider")
        lines += _cap([f"{t} {o or '(blank)'} -> {n}" for t, o, n in em["updated"]], "                ")
        if kept:
            lines += _cap([f"kept {k}" for k in kept], "                ")
        if em["no_date"]:
            lines.append("                no date: " + ", ".join(em["no_date"]))
        if em["upcoming"]:
            lines.append(f"                held names reporting within {market.UPCOMING_DAYS} days: "
                         + ", ".join(f"{t} {d:%m/%d}" for t, d in em["upcoming"]))
        lines += _positions_lines(m["positions"], rep)
        lines += _performance_lines(m.get("performance"), rep)

    r = rep.get("recalc")
    if r is not None:
        lines.append(f"Recalc          {r.get('status') or 'FAILED'}, {r.get('total_errors', '?')} formula errors")
    r38 = rep.get("rule38")
    if r38 is not None:
        lines.append(f"Rule 38         {r38['over_cap']} position(s) over the 0.70 cap -- {r38['status']}")
    d = rep.get("db")
    if d is not None:
        if d.get("ok") is None:
            lines.append(f"Database        {d['note']}")
        elif d.get("error"):
            lines.append(f"Database        DRIFT CHECK FAILED: {d['error']}")
        else:
            body = ", ".join(f"{k} {t['rows']} entries" for k, t in d["tabs"].items())
            fmt = sum(t["format_diffs"] for t in d["tabs"].values())
            lines.append(f"Database        {body}: " + ("in sync" if d["ok"] else "DISAGREES with the workbook")
                         + (f" ({fmt} formatting difference(s) the next Journal / Drawdown write resets)" if fmt else ""))
    a = rep.get("audit")
    if a is not None:
        lines.append(f"Audit           {a['checks']} checks, {len(a['failed'])} failed"
                     + (f"; {len(a['standing'])} standing Scoring-provenance finding(s), not touched by this "
                        f"routine ({', '.join(n.split(':', 1)[-1] for n, _ in a['standing'])})" if a.get("standing") else ""))
        lines += _cap([f"[FAIL] {n}: {d}" for n, d in a["failed"]], "                ")

    if rep["warnings"]:
        lines += ["", "Warnings"]
        lines += _cap(rep["warnings"], "  - ", 20)
    if rep["problems"]:
        lines += ["", "NEEDS ATTENTION"]
        lines += _cap(rep["problems"], "  - ")
    lines += ["", "Nothing is committed." if not rep.get("dry_run") else "Nothing was changed."]
    if rep["wrote"] and not rep.get("dry_run"):
        lines[-1] += " Review the diff, then tell Claude to commit."
    return "\n".join(lines)
